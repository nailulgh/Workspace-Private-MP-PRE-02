# -*- coding: utf-8 -*-
"""
Generator Script for LKM 3: MANAJEMEN INTEGRASI PROYEK
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def create_element(name):
    return OxmlElement(name)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = create_element('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def set_table_borders(table, color="000000", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    tblBorders = create_element('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = create_element(f'w:{border_name}')
        border.set(qn('w:val'), val)
        border.set(qn('w:sz'), sz)
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    tblPr.append(tblBorders)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = create_element('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = create_element(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(create_element('w:tblHeader'))

def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(create_element('w:cantSplit'))

def set_col_widths(table, widths_cm):
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body_p(doc, text, bold_prefix=None, space_after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        r_pre.font.color.rgb = RGBColor(0, 0, 0)
        
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(12)
    r_body.font.bold = False
    r_body.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_num_p(doc, num_str, bold_prefix, text, space_after=4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    
    r_num = p.add_run(num_str + " ")
    r_num.font.name = "Times New Roman"
    r_num.font.size = Pt(12)
    r_num.font.bold = True
    
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        
    r_body = p.add_run(text)
    r_body.font.name = "Times New Roman"
    r_body.font.size = Pt(12)
    r_body.font.bold = False
    return p

def create_table(doc, col_widths, headers, data, align_cols=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    set_table_borders(table)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    hdr_cells = table.rows[0].cells
    make_row_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_shading(hdr_cells[i], "E8EEF5")
        set_cell_margins(hdr_cells[i])
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)
            
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        prevent_row_split(row)
        for c_idx, val in enumerate(row_data):
            cell = row.cells[c_idx]
            cell.text = str(val)
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            if align_cols and c_idx < len(align_cols):
                p.alignment = align_cols[c_idx]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)
                r.font.bold = False
                r.font.color.rgb = RGBColor(0, 0, 0)
                
    set_col_widths(table, col_widths)
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)
    p_sp.paragraph_format.line_spacing = 1.15
    return table

def add_cover(doc):
    # Paragraph 0: Title
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_after = Pt(5.8)
    p0.paragraph_format.line_spacing = 1.0
    r0 = p0.add_run("LKM 3: MANAJEMEN INTEGRASI PROYEK")
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(12)
    r0.font.bold = True
    
    # Paragraph 1: Subtitle
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(10)
    p1.paragraph_format.line_spacing = 1.0
    r1 = p1.add_run("Mata Kuliah: Teori Manajemen Proyek")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    
    # Paragraph 2: Empty
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(10)
    
    # Paragraph 3: Dosen Pengampu Header
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(5.8)
    p3.paragraph_format.line_spacing = 1.5
    r3 = p3.add_run("Dosen Pengampu:")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(12)
    r3.font.bold = True
    
    # Paragraph 4: Dosen Pengampu Name
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_after = Pt(5.8)
    p4.paragraph_format.line_spacing = 1.5
    r4 = p4.add_run("Dr. Muhammad Ainul Yaqin, S.Si, M.Kom")
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(12)
    
    # Paragraph 5: Space
    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_after = Pt(5.8)
    p5.paragraph_format.line_spacing = 1.5
    r5 = p5.add_run("                ")
    
    # Paragraph 6: Space
    p6 = doc.add_paragraph()
    p6.paragraph_format.space_after = Pt(5.5)
    p6.paragraph_format.line_spacing = 1.5
    
    # Paragraph 7: Logo
    p7 = doc.add_paragraph()
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p7.paragraph_format.space_after = Pt(5.5)
    p7.paragraph_format.line_spacing = 1.5
    if os.path.exists("uin_logo.png"):
        r7 = p7.add_run()
        r7.add_picture("uin_logo.png", width=Cm(4.85))
        
    # Paragraph 8: Space
    p8 = doc.add_paragraph()
    p8.paragraph_format.space_after = Pt(5.8)
    p8.paragraph_format.line_spacing = 1.5
    
    # Paragraph 9: Oleh
    p9 = doc.add_paragraph()
    p9.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p9.paragraph_format.space_after = Pt(5.8)
    p9.paragraph_format.line_spacing = 1.5
    r9 = p9.add_run("Oleh :")
    r9.font.name = "Times New Roman"
    r9.font.size = Pt(12)
    r9.font.bold = True
    
    # Paragraph 10: Nama Mahasiswa
    p10 = doc.add_paragraph()
    p10.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p10.paragraph_format.space_after = Pt(5.8)
    p10.paragraph_format.line_spacing = 1.5
    r10 = p10.add_run("Muhammad Nailul Ghufron Majid")
    r10.font.name = "Times New Roman"
    r10.font.size = Pt(12)
    
    # Paragraph 11: NIM
    p11 = doc.add_paragraph()
    p11.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p11.paragraph_format.space_after = Pt(5.8)
    p11.paragraph_format.line_spacing = 1.5
    r11 = p11.add_run("240605110160")
    r11.font.name = "Times New Roman"
    r11.font.size = Pt(12)
    
    # Paragraphs 12-14: Space
    for _ in range(3):
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(5.8)
        p_sp.paragraph_format.line_spacing = 1.5
        
    # Paragraph 15-19: Institusi
    inst_lines = [
        "PRODI TEKNIK INFORMATIKA",
        "FAKULTAS SAINS DAN TEKNOLOGI",
        "UNIVERSITAS ISLAM NEGERI MAULANA MALIK IBRAHIM",
        "MALANG",
        "2026"
    ]
    for idx, line in enumerate(inst_lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.5
        r = p.add_run(line)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = True
        
    doc.add_page_break()

print("Cover module written.")
