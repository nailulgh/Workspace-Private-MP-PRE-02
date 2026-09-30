# -*- coding: utf-8 -*-
"""
Section B Builder: Workshop & Tugas Kelompok (Checkpoint #1)
Project Charter PRE-02: Sistem Early Warning Prediksi Probabilitas Keterlambatan Proyek
"""

import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_b(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    create_table = helpers['create_table']
    
    add_heading_1(doc, "B. KEGIATAN WORKSHOP DAN TUGAS KELOMPOK (CHECKPOINT #1)")
    
    # Header Project Charter
    add_heading_2(doc, "DOKUMEN PROJECT CHARTER (PIAGAM PROYEK) - CHECKPOINT #1")
    add_body_p(doc, 
        "Berikut ini adalah Piagam Proyek (Project Charter) yang disusun secara terstruktur oleh Tim Proyek PRE-02 "
        "sebagai dokumen otorisasi resmi pelaksanaan proyek riset-terapan. Dokumen ini disusun dengan mengacu secara ketat "
        "pada standar PMBOK Guide Edisi Ke-6 (PMI, 2018, Bab 4.1) serta memuat ke-11 elemen wajib sesuai instruksi penugasan:")

    # Metadata Proyek Table
    headers_meta = ["Atribut Dokumen", "Keterangan Identitas Proyek"]
    data_meta = [
        ["Kode / Identitas Proyek", "PRE-02 / FNE-2026 (Prediksi Probabilitas Keterlambatan Proyek Multi Sumber Daya Terbatas)"],
        ["Tanggal Efektif Pengesahan", "27 September 2026 (Pertemuan Ke-3 Perkuliahan)"],
        ["Versi Dokumen", "Versi 1.0 (Baseline Resmi Disahkan)"],
        ["Klasifikasi Akses", "Dokumen Tata Kelola Internal Konsorsium FNE & Akademik UIN Malang"]
    ]
    create_table(doc, [4.2, 10.6], headers_meta, data_meta, [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    # 1. Judul Proyek
    add_heading_3(doc, "1. Judul Proyek")
    add_body_p(doc, 
        "Pengembangan Sistem Early Warning Prediksi Probabilitas Keterlambatan Proyek Berbasis Machine Learning "
        "pada Portofolio Multi-Modul dengan Keterbatasan Sumber Daya (Studi Kasus: Megaproyek Super ERP Farm Nation Enterprise 2026).",
        bold_prefix="Judul Resmi Proyek: ")
    add_body_p(doc, 
        "\"PRE-02: Predictive Early Warning System for Project Delay Probability\".",
        bold_prefix="Nama Singkat (Project Codename): ")

    # 2. Latar Belakang Proyek
    add_heading_3(doc, "2. Latar Belakang Proyek (Business Case & Context)")
    add_body_p(doc, 
        "Megaproyek Super ERP Farm Nation Enterprise (FNE) 2026 merupakan inisiatif strategis transformasi digital pertanian terpadu "
        "yang mencakup 8 (delapan) sub-proyek berskala korporat, membentang dari modul fondasi arsitektur (P-FNE-01), operasional lapangan (P-FNE-02), "
        "hingga kecerdasan buatan otonom (P-FNE-07 dan P-FNE-08). Karakteristik megaproyek ini memiliki kompleksitas arsitektur microservices "
        "yang sangat tinggi serta rantai dependensi antar-modul yang ketat (predecessor-successor dependencies).")
    add_body_p(doc, 
        "Namun, hasil audit pemantauan fase inisiasi menunjukkan bahwa proyek menghadapi kendala kritis berupa kelangkaan tenaga ahli "
        "pengembang perangkat lunak senior (shared developer bottleneck). Keterbatasan ini menyebabkan pemanfaatan developer melebihi kapasitas normal "
        "(resource utilization rate mencapai 1.10 hingga 1.25), sehingga memicu penumpukan antrean task dan keterlambatan penyelesaian modul-modul fondasi. "
        "Karena dependensi sistem yang bersifat sekuensial berjenjang (Tier-0 hingga Tier-4), keterlambatan pada satu task di modul Core Platform (P-FNE-01) "
        "secara otomatis memicu efek domino keterlambatan (cascading schedule delays) pada modul-modul hilir seperti modul Finance maupun Autonomous Enterprise.")
    add_body_p(doc, 
        "Di sisi lain, praktik pemantauan risiko manajemen proyek konvensional saat ini hanya mengandalkan label kualitatif statis (seperti risiko "
        "Low, Medium, atau High) yang bersifat subjektif dan tidak memberikan kejelasan ambang batas intervensi manajerial. Manajer Proyek tidak "
        "mengetahui secara pasti berapa peluang kuantitatif suatu task akan meleset dari deadline dan tindakan korektif apa yang harus diprioritaskan. "
        "Oleh karena itu, proyek PRE-02 ini diinisiasi secara mendesak untuk membangun model supervised learning yang menghasilkan estimasi probabilitas "
        "numerik kontinu P(Delay) bernilai antara 0.00 hingga 1.00 yang terkalibrasi secara presisi, sehingga berfungsi sebagai Sistem Peringatan Dini "
        "(Early Warning System) objektif bagi Manajer Proyek dalam menyelamatkan jadwal dan anggaran proyek sebelum kegagalan terjadi.")

    # 3. Tujuan Proyek (SMART)
    add_heading_3(doc, "3. Tujuan Proyek (SMART Objectives)")
    add_body_p(doc, 
        "Tujuan pelaksanaan proyek PRE-02 dirumuskan secara tegas menggunakan kerangka kerja SMART (Specific, Measurable, Achievable, Relevant, Time-bound):")
    add_num_p(doc, "a.", "Spesifik (Specific):",
        "Membangun dan melatih pipeline model prediktif berbasis machine learning yang memanfaatkan 7 fitur data monitoring proyek PMBOK "
        "(SPI, utilisasi sumber daya, skor risiko, jumlah dependensi pendahulu, jumlah change request, rencana durasi, dan rencana jam kerja) "
        "untuk menghasilkan estimasi nilai probabilitas keterlambatan numerik kontinu P(Delay).")
    add_num_p(doc, "b.", "Terukur (Measurable):",
        "Mencapai performa diskriminasi model klasifikasi dengan nilai Area Under ROC Curve (ROC-AUC) minimal sebesar 0.85, tingkat kesalahan probabilitas "
        "Brier Score maksimal 0.15, serta menghasilkan kurva kalibrasi (Reliability Diagram) yang teruji mendekati garis ideal.")
    add_num_p(doc, "c.", "Dapat Dicapai (Achievable):",
        "Eksperimen dirancang secara realistis dengan memanfaatkan dataset empiris 26 core module tasks dari dokumen PMP resmi ke-8 sub-proyek Super ERP FNE "
        "dan menerapkan 4 algoritma teruji (Logistic Regression, Random Forest, Gradient Boosting, dan Multi-Layer Perceptron Neural Network).")
    add_num_p(doc, "d.", "Relevan (Relevant):",
        "Mengonversi probabilitas prediksi ke dalam visualisasi dashboard Early Warning System tiga zona (Hijau: P < 0.35; Kuning: 0.35 <= P < 0.65; "
        "dan Merah: P >= 0.65) guna memandu keputusan korektif prioritas re-alokasi developer oleh Manajer Proyek FNE.")
    add_num_p(doc, "e.", "Terikat Waktu (Time-bound):",
        "Menyelesaikan seluruh tahapan riset, eksperimen komputasi, penyusunan paper ilmiah format IMRAD, dan presentasi luaran tuntas dalam rentang "
        "waktu 8 (delapan) pekan perkuliahan semester ganjil tahun akademik 2026.")

    # 4. Ruang Lingkup Awal (Scope Boundary)
    add_heading_3(doc, "4. Ruang Lingkup Awal (Project Scope: In-Scope vs Out-of-Scope)")
    add_body_p(doc, 
        "Untuk memastikan kejelasan batas pekerjaan dan mencegah timbulnya scope creep selama pelaksanaan proyek, batasan ruang lingkup awal "
        "didefinisikan secara tegas dalam Tabel 3 berikut:")
        
    headers_scope = ["Kategori Ruang Lingkup", "Rincian Batasan Pekerjaan Proyek PRE-02"]
    data_scope = [
        [
            "Pekerjaan Termasuk (In-Scope)",
            "1. Ekstraksi, pembersihan, dan verifikasi data 26 core task dari 8 dokumen PMP dan KAK Super ERP FNE.\n"
            "2. Pembentukan kamus data dengan 7 variabel bebas monitoring dan 1 variabel target biner keterlambatan aktual.\n"
            "3. Pelatihan, validasi silang (cross-validation), dan penyetelan hyperparameter pada 4 algoritma Machine Learning.\n"
            "4. Analisis kalibrasi probabilitas menggunakan Brier Score, Log Loss, dan kurva kalibrasi (Reliability Diagram).\n"
            "5. Analisis feature importance untuk mengidentifikasi variabel yang paling berpengaruh terhadap keterlambatan.\n"
            "6. Perancangan skema ambang batas intervensi (Threshold Analysis) untuk zona peringatan dini (Hijau, Kuning, Merah).\n"
            "7. Penyusunan naskah paper ilmiah akademik lengkap (IMRAD) serta slide presentasi sidang progres mingguan."
        ],
        [
            "Pekerjaan Tidak Termasuk (Out-of-Scope)",
            "1. Pengkodean fungsional logika bisnis aplikasi ERP dari ke-8 modul Super ERP Farm Nation Enterprise.\n"
            "2. Pengadaan infrastruktur server produksi fisik on-premise atau instalasi kabel jaringan di perkebunan.\n"
            "3. Pemasangan dan integrasi perangkat keras fisik sensor IoT, drone, atau mesin pertanian di lapangan perkebunan.\n"
            "4. Pengambilan keputusan manajerial final terkait pemecatan atau penggantian vendor pihak ketiga di luar simulasi model.\n"
            "5. Pembuatan sistem aplikasi mobile native di luar lingkup prototipe dashboard visualisasi analitik."
        ]
    ]
    create_table(doc, [4.5, 10.3], headers_scope, data_scope, [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY])

    # 5. Manfaat Proyek
    add_heading_3(doc, "5. Manfaat Proyek (Project Benefits & Value Proposition)")
    add_body_p(doc, 
        "Keberhasilan penyelesaian proyek PRE-02 akan memberikan manfaat strategis yang terbagi ke dalam tiga dimensi utama:")
    add_num_p(doc, "a.", "Manfaat Operasional bagi Manajer Proyek (Operational Benefits):",
        "Memberikan kapabilitas pemantauan prediktif (predictive foresight) sehingga Manajer Proyek dapat mendeteksi task berisiko tinggi 2 hingga 3 pekan "
        "sebelum tenggat waktu berakhir. Hal ini memungkinkan tindakan pemulihan jadwal (seperti leveling sumber daya atau pergeseran dependensi) "
        "dilakukan secara proaktif dan presisi.")
    add_num_p(doc, "b.", "Manfaat Finansial bagi Konsorsium Korporat (Financial Benefits):",
        "Mereduksi potensi denda penalti keterlambatan proyek korporat senilai ratusan juta rupiah yang diakibatkan oleh terhambatnya peluncuran modul "
        "kritis serta meminimalkan biaya lembur (crash cost) yang tidak efektif akibat penanganan krisis yang terlambat.")
    add_num_p(doc, "c.", "Manfaat Ilmiah dan Akademik (Scientific & Academic Benefits):",
        "Menghasilkan kontribusi ilmiah berupa kerangka kerja monitoring proyek probabilistik terkalibrasi pada lingkungan multi-proyek bersumber daya "
        "terbatas, yang mengisi research gap pada literatur manajemen rekayasa perangkat lunak dan siap dipublikasikan pada jurnal bereputasi.")

    # 6. Stakeholder Utama
    add_heading_3(doc, "6. Stakeholder Utama (Key Stakeholders Analysis)")
    add_body_p(doc, 
        "Identifikasi pihak-pihak berkepentingan utama yang memiliki dampak langsung maupun tidak langsung terhadap pelaksanaan proyek PRE-02 "
        "disajikan dalam Tabel 4 di bawah ini:")
        
    headers_stk = ["Nama / Peran Stakeholder", "Klasifikasi Organisasi", "Pengaruh / Kepentingan", "Tanggung Jawab & Ekspektasi Utama"]
    data_stk = [
        [
            "Dr. Muhammad Ainul Yaqin, S.Si, M.Kom\n(Dosen Pengampu / Sponsor Akademik)",
            "Program Studi Teknik Informatika, UIN Malang",
            "High Power /\nHigh Interest",
            "Memberikan bimbingan metodologis, menetapkan standar luaran akademis, menguji ketepatan analisis, dan memberikan pengesahan formal atas deliverables."
        ],
        [
            "Direktur PMO Farm Nation Enterprise\n(Sponsor Eksekutif Korporat)",
            "Konsorsium Super ERP FNE 2026",
            "High Power /\nHigh Interest",
            "Menyediakan mandat operasional, mengotorisasi akses dataset historis 8 sub-proyek, dan mengevaluasi kelayakan adopsi model ke dashboard FNE."
        ],
        [
            "Muhammad Nailul Ghufron Majid\n(Project Manager & Lead Data/ML)",
            "Tim Riset Mahasiswa PRE-02",
            "High Power /\nHigh Interest",
            "Memimpin integrasi seluruh proses proyek, mengelola jadwal 8 pekan, mengeksekusi pipeline machine learning, kalibrasi, dan memastikan kepatuhan standar."
        ],
        [
            "Mukti Baskara\n(Lead Author & Qualitative Specialist)",
            "Tim Riset Mahasiswa PRE-02",
            "Medium Power /\nHigh Interest",
            "Bertanggung jawab atas sintesis studi literatur, formulasi research gap, narasi naskah paper ilmiah IMRAD, dan penyusunan materi presentasi progres."
        ],
        [
            "Rafiq\n(System Analyst & Testing Specialist)",
            "Tim Riset Mahasiswa PRE-02",
            "Medium Power /\nHigh Interest",
            "Menganalisis dependensi arsitektur antar-modul, melakukan pengujian metrik evaluasi (ROC-AUC, Brier Score), dan memvalidasi keakuratan kamus data."
        ],
        [
            "Para Manajer Proyek Modul FNE\n(P-FNE-01 s.d P-FNE-08)",
            "Tim Pengembang Konsorsium FNE",
            "Medium Power /\nHigh Interest",
            "Bertindak sebagai pengguna akhir (end-users) sistem peringatan dini, menyediakan klarifikasi profil risiko task, dan mengevaluasi kepraktisan dashboard."
        ]
    ]
    create_table(doc, [3.8, 3.2, 2.5, 5.3], headers_stk, data_stk, 
                 [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY])

    # 7. Sponsor Proyek
    add_heading_3(doc, "7. Sponsor Proyek (Project Sponsor)")
    add_body_p(doc, 
        "Pihak yang mengotorisasi keberadaan proyek, menyediakan dukungan pendanaan fasilitas komputasi, serta memberikan mandat hukum "
        "pelaksanaan proyek PRE-02 adalah:", bold_prefix="Entitas Pengesah: ")
    add_num_p(doc, "1.", "Sponsor Eksekutif Korporat:",
        "Direktur Project Management Office (PMO) Konsorsium Super ERP Farm Nation Enterprise 2026, selaku pemilik mandat portofolio proyek.")
    add_num_p(doc, "2.", "Sponsor dan Penguji Akademik:",
        "Dr. Muhammad Ainul Yaqin, S.Si, M.Kom, selaku Dosen Pengampu Mata Kuliah Teori Manajemen Proyek, Fakultas Sains dan Teknologi, "
        "Universitas Islam Negeri Maulana Malik Ibrahim Malang.")

    # 8. Manajer Proyek & Wewenang
    add_heading_3(doc, "8. Manajer Proyek dan Otoritas yang Diberikan (Project Manager & Authority)")
    add_body_p(doc, 
        "Nama yang ditunjuk secara resmi untuk memimpin, mengoordinasikan, dan mengendalikan proyek riset PRE-02 dari awal hingga akhir adalah "
        "Muhammad Nailul Ghufron Majid (NIM: 240605110160).", bold_prefix="Penunjukan Resmi: ")
    add_body_p(doc, 
        "Sesuai mandat Project Charter ini, Manajer Proyek diberikan tingkat wewenang formal sebagai berikut:", bold_prefix="Wewenang Manajer Proyek: ")
    add_bullet_p(doc, "Otoritas Pengalokasian Tugas dan Kepemimpinan Tim:", 
        "Berwenang penuh membagi tanggung jawab kerja harian, menetapkan target milestone mingguan, dan mengevaluasi produktivitas anggota tim.")
    add_bullet_p(doc, "Otoritas Pengambilan Keputusan Teknis:", 
        "Berwenang menetapkan arsitektur algoritma machine learning, library pemrograman Python, prosedur kalibrasi probabilitas, dan struktur dataset.")
    add_bullet_p(doc, "Otoritas Pengelolaan Anggaran Operasional:", 
        "Berwenang menyetujui komitmen pembelanjaan biaya komputasi cloud, server storage, dan operasional riset dalam batas plafon anggaran yang disetujui sponsor.")
    add_bullet_p(doc, "Otoritas Penjadwalan dan Pengendalian Perubahan:", 
        "Berwenang menolak usulan perubahan teknis yang berisiko merusak jalur kritis jadwal 8 pekan perkuliahan dan membawa usulan perubahan strategis ke hadapan Sponsor.")

    # 9. Anggaran Awal Proyek
    add_heading_3(doc, "9. Anggaran Awal Proyek (Preliminary Cost Estimate Baseline)")
    add_body_p(doc, 
        "Estimasi total biaya awal yang disetujui untuk mendukung eksekusi komputasi, penyediaan perangkat, pengolahan data, dan luaran proyek PRE-02 "
        "adalah sebesar Rp 45.000.000,- (Empat Puluh Lima Juta Rupiah). Rincian alokasi anggaran awal disajikan dalam Tabel 5 berikut:")
        
    headers_cost = ["No", "Kategori Pembelanjaan Biaya Proyek", "Rincian Item Pengeluaran", "Estimasi Biaya (Rp)"]
    data_cost = [
        [
            "1", 
            "Infrastruktur Cloud GPU & Server Training", 
            "Sewa komputasi cloud (AWS EC2 GPU Instance / Google Cloud Vertex AI) untuk training model, grid-search hyperparameter, dan komputasi kalibrasi.",
            "12.000.000"
        ],
        [
            "2", 
            "Perangkat Keras Pengembang & Storage", 
            "Alokasi sewa dedicated local workstation untuk pemodelan data berkecepatan tinggi serta media penyimpanan terenkripsi repositori data PMP.",
            "15.000.000"
        ],
        [
            "3", 
            "Lisensi Perangkat Lunak & Repositori", 
            "Lisensi Git Enterprise, tools visualisasi analitik data, API sandbox environment, dan perangkat lunak manajemen referensi ilmiah.",
            "5.000.000"
        ],
        [
            "4", 
            "Ekstraksi & Sanitasi Data Historis", 
            "Biaya penyiapan data, audit validasi kamus data dari 8 dokumen PMP FNE, dan kompensasi tim verifikator kualitas data monitoring.",
            "6.000.000"
        ],
        [
            "5", 
            "Operasional Tim, FGD & Publikasi", 
            "Konsumsi koordinasi mingguan, biaya penyelenggaraan Focus Group Discussion (FGD) validasi pakar, pendaftaran hak cipta, dan submit artikel jurnal.",
            "7.000.000"
        ],
        [
            "", 
            "TOTAL ESTIMASI ANGGARAN AWAL (COST BASELINE)", 
            "Plafon anggaran terpadu proyek PRE-02 yang diotorisasi Sponsor",
            "45.000.000"
        ]
    ]
    create_table(doc, [1.0, 4.5, 6.3, 3.0], headers_cost, data_cost, 
                 [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT])

    # 10. Jadwal Awal dan Milestone Kunci
    add_heading_3(doc, "10. Jadwal Awal dan Milestone Kunci (Milestone Schedule)")
    add_body_p(doc, 
        "Jadwal pelaksanaan proyek PRE-02 dirancang terintegrasi dengan siklus 8 (delapan) pertemuan perkuliahan manajemen proyek. "
        "Setiap milestone mewakili titik gerbang evaluasi (stage gate) yang wajib dipenuhi:")
        
    headers_ms = ["Milestone", "Pertemuan / Waktu", "Fokus Aktivitas Utama", "Deliverables Kunci yang Dihasilkan"]
    data_ms = [
        [
            "M1: Project Kickoff & Scoping",
            "Pertemuan 1\n(Pekan 1)",
            "Orientasi pemahaman umum judul penelitian, identifikasi latar belakang, penetapan variabel X dan Y, serta pembagian peran tim.",
            "Slide Presentasi Progress 1 dan Notulensi Pemahaman Konseptual Topik."
        ],
        [
            "M2: Literature & Research Gap",
            "Pertemuan 2\n(Pekan 2)",
            "Penelusuran paper ilmiah bereputasi, sintesis teori EVM dan ML, serta formulasi research gap probabilitas terkalibrasi.",
            "Matriks Literatur (8 paper inti) dan Draft Awal Tinjauan Pustaka (Bab II)."
        ],
        [
            "M3: Data Preparation & Charter",
            "Pertemuan 3\n(Pekan 3)",
            "Ekstraksi 26 task dari 8 PMP FNE, pembersihan data, pengesahan Project Charter, dan deskripsi karakteristik dataset.",
            "Dataset Final `dataset_pre02_fne.csv` dan Dokumen Project Charter (Checkpoint #1)."
        ],
        [
            "M4: Experiment Design Pipeline",
            "Pertemuan 4\n(Pekan 4)",
            "Penyusunan flowchart metodologi eksperimen ML, skrip data preprocessing, pembagian Train/Test split, dan penyiapan 4 model.",
            "Flowchart Eksperimen, Skrip Pipeline Python, dan Draft Metodologi (Bab III)."
        ],
        [
            "M5: Experiment Execution",
            "Pertemuan 5\n(Pekan 5)",
            "Eksekusi komputasi machine learning, pelatihan Logistic Regression, Random Forest, Gradient Boosting, dan MLP, serta log hasil.",
            "Dataset Hasil Prediksi Probabilitas dan Log Training Metrik Evaluasi Awal."
        ],
        [
            "M6: Model Evaluation & Calibration",
            "Pertemuan 6\n(Pekan 6)",
            "Analisis kinerja diskriminasi (ROC-AUC, Precision, Recall), uji kalibrasi Brier Score, dan plotting Reliability Diagram.",
            "Tabel Metrik Performa Lengkap, Grafik ROC Curve, dan Calibration Curve."
        ],
        [
            "M7: EWS Threshold & Discussion",
            "Pertemuan 7\n(Pekan 7)",
            "Threshold analysis untuk menentukan ambang batas EWS 3 zona (Hijau/Kuning/Merah), interpretasi feature importance, dan implikasi PM.",
            "Draft Bagian Pembahasan (Bab IV) dan Panduan Kebijakan Early Warning System."
        ],
        [
            "M8: Paper Finalization & Closure",
            "Pertemuan 8\n(Pekan 8)",
            "Finalisasi penulisan paper ilmiah lengkap format IMRAD, penyempurnaan abstrak, slide presentasi akhir sidang, dan pengarsipan.",
            "Naskah Paper Ilmiah Final, Slide Presentasi Sidang, dan Final Project Report."
        ]
    ]
    create_table(doc, [3.2, 2.2, 5.0, 4.4], headers_ms, data_ms, 
                 [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.JUSTIFY])

    # 11. Kriteria Keberhasilan Proyek
    add_heading_3(doc, "11. Kriteria Keberhasilan Proyek (Project Success Criteria)")
    add_body_p(doc, 
        "Keberhasilan proyek PRE-02 diukur berdasarkan pemenuhan target ganda: kriteria kinerja teknis rekayasa machine learning "
        "dan kriteria keberhasilan manajerial proyek:")
        
    headers_crit = ["Dimensi Penilaian", "Parameter Tolok Ukur", "Target Keberhasilan yang Wajib Dicapai"]
    data_crit = [
        [
            "Kinerja Teknis Machine Learning (Predictive Performance)",
            "Kemampuan Diskriminasi Model (ROC-AUC)",
            "Nilai ROC-AUC pada data uji (test set) mencapai minimal 0.85 (kategori sangat baik)."
        ],
        [
            "Kinerja Teknis Machine Learning (Predictive Performance)",
            "Akurasi Kalibrasi Probabilitas (Brier Score)",
            "Brier Score bernilai <= 0.15 dan kurva kalibrasi berhimpit mendekati garis diagonal ideal."
        ],
        [
            "Kinerja Teknis Machine Learning (Predictive Performance)",
            "Keberfungsian Early Warning System (EWS)",
            "Mampu mengelompokkan 26 task modul FNE ke dalam 3 zona risiko secara objektif dan logis."
        ],
        [
            "Kinerja Manajerial & Tata Kelola (Project Management)",
            "Kepatuhan Jadwal Jalur Kritis (Schedule Variance)",
            "Seluruh 8 milestone diselesaikan tepat waktu sesuai agenda perkuliahan (SV >= 0, SPI >= 1.0)."
        ],
        [
            "Kinerja Manajerial & Tata Kelola (Project Management)",
            "Kepatuhan Anggaran (Cost Variance)",
            "Pengeluaran aktual tidak melampaui plafon anggaran Rp 45.000.000,- (CV >= 0, CPI >= 1.0)."
        ],
        [
            "Kinerja Manajerial & Tata Kelola (Project Management)",
            "Kelayakan Luaran Akademik (Deliverables Quality)",
            "Paper ilmiah memenuhi standar struktur IMRAD dan disetujui Dosen Pengampu untuk disubmit ke jurnal."
        ]
    ]
    create_table(doc, [3.5, 4.3, 7.0], headers_crit, data_crit, 
                 [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.JUSTIFY])

    # Lembar Pengesahan
    add_heading_3(doc, "LEMBAR PENGESAHAN DAN OTORISASI FORMAL PROJECT CHARTER")
    add_body_p(doc, 
        "Dengan menandatangani dokumen ini, pihak Sponsor secara resmi mengesahkan eksistensi proyek PRE-02 dan memberikan mandat "
        "penuh kepada Manajer Proyek untuk melaksanakan kegiatan sesuai ruang lingkup, jadwal, dan alokasi sumber daya yang disepakati.")
        
    headers_sign = ["Peran Pengesah", "Nama Pejabat / Tanggal", "Tanda Tangan Otorisasi"]
    data_sign = [
        [
            "Project Manager (Penerima Mandat)",
            "Muhammad Nailul Ghufron Majid\nNIM. 240605110160\nTanggal: 27 September 2026",
            "\n\n( ...................................................... )\n"
        ],
        [
            "Project Sponsor & Dosen Pengampu",
            "Dr. Muhammad Ainul Yaqin, S.Si, M.Kom\nNIP. 197801012005011002\nTanggal: 27 September 2026",
            "\n\n( ...................................................... )\n"
        ],
        [
            "Executive Corporate Sponsor",
            "Direktur Project Management Office\nKonsorsium Super ERP FNE 2026\nTanggal: 27 September 2026",
            "\n\n( ...................................................... )\n"
        ]
    ]
    create_table(doc, [4.5, 5.3, 5.0], headers_sign, data_sign, 
                 [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER])

print("Section B Builder written successfully.")
