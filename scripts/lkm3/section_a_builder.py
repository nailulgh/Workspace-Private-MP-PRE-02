# -*- coding: utf-8 -*-
"""
Section A Builder: Tugas Tertulis Individu
Pertemuan 3: Manajemen Integrasi Proyek
"""

import docx
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_a(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    create_table = helpers['create_table']
    
    add_heading_1(doc, "A. TUGAS TERTULIS INDIVIDU")
    
    # Soal 1
    add_heading_2(doc, "1. Apa yang dimaksud dengan manajemen integrasi proyek dan mengapa area pengetahuan ini dianggap paling fundamental?")
    add_body_p(doc, 
        "Manajemen Integrasi Proyek (Project Integration Management) adalah area pengetahuan yang mencakup proses-proses dan "
        "aktivitas-aktivitas yang diperlukan untuk mengidentifikasi, mendefinisikan, mengombinasikan, menyatukan, dan mengoordinasikan "
        "berbagai proses serta aktivitas manajemen proyek dalam seluruh Kelompok Proses Manajemen Proyek (Project Management Process Groups), "
        "mulai dari Inisiasi, Perencanaan, Pelaksanaan, Pemantauan dan Pengendalian, hingga Penutupan (PMI, 2018, hal. 69). "
        "Dalam pengertian esensial, manajemen integrasi berfokus pada pengambilan keputusan strategis terkait alokasi sumber daya, "
        "penyeimbangan berbagai sasaran dan batasan proyek yang saling bersaing, serta pengkoordinasian ketergantungan antar-area pengetahuan "
        "manajemen proyek secara menyeluruh dan berkesinambungan.")
    
    add_body_p(doc, 
        "Area pengetahuan Manajemen Integrasi Proyek diposisikan sebagai area yang paling fundamental dalam disiplin manajemen proyek "
        "karena empat alasan krusial berikut:")
    
    add_num_p(doc, "1.", "Fungsi Konsolidator dan Perekat Sistemik (The Unifying Glue):",
        "Dari sepuluh Area Pengetahuan yang didefinisikan dalam PMBOK Guide, sembilan area lainnya (Ruang Lingkup, Jadwal, Biaya, Kualitas, "
        "Sumber Daya, Komunikasi, Risiko, Pengadaan, dan Pemangku Kepentingan) bersifat spesifik dan terfokus pada domain teknis parsial. "
        "Manajemen integrasi berfungsi sebagai perekat yang menyatukan seluruh elemen tersebut menjadi satu kesatuan dokumen dan eksekusi "
        "yang utuh. Mengambil analogi sebuah orkestra simfoni, sembilan area pengetahuan lain adalah musisi dengan instrumennya masing-masing, "
        "sedangkan manajemen integrasi adalah peran konduktor yang memastikan seluruh instrumen berharmonisasi menghasilkan karya musik yang indah.")
        
    add_num_p(doc, "2.", "Tanggung Jawab Eksklusif Manajer Proyek yang Tidak Dapat Didelegasikan:",
        "PMBOK Guide secara eksplisit menyatakan bahwa manajemen integrasi adalah satu-satunya area pengetahuan yang menjadi tanggung jawab "
        "mutlak dan eksklusif dari Manajer Proyek (PMI, 2018, hal. 72). Peran-peran teknis spesifik dapat didelegasikan—misalnya penjadwalan "
        "kepada project scheduler, mitigasi risiko kepada risk manager, atau estimasi biaya kepada cost estimator—namun akuntabilitas "
        "dalam menyatukan seluruh subsistem tersebut dan melihat gambaran besar proyek (helicopter view) sepenuhnya bertumpu pada manajer proyek.")
        
    add_num_p(doc, "3.", "Pengendali Trade-Off dan Keseimbangan Batasan Proyek (Triple Constraints Balancing):",
        "Dalam realitas pelaksanaan proyek, perubahan pada satu parameter batasan dipastikan menimbulkan efek domino terhadap batasan lain. "
        "Sebagai contoh, apabila sponsor meminta percepatan jadwal penyelesaian (schedule compression), tindakan tersebut menuntut penambahan "
        "biaya lembur/crash cost (cost impact), meningkatkan utilisasi developer hingga berisiko burnout (resource impact), serta memicu "
        "potensi cacat logika program (quality and risk impact). Manajemen integrasi menyediakan kerangka kerja rasional bagi manajer proyek "
        "untuk menimbang, menegosiasikan, dan mengendalikan kompromi (trade-offs) tersebut secara objektif.")
        
    add_num_p(doc, "4.", "Pengawal Siklus Hidup Proyek dari Hulu ke Hilir (End-to-End Governance):",
        "Manajemen integrasi mengawal proyek dari fase paling awal—yaitu saat proyek masih berupa kebutuhan bisnis abstrak yang diresmikan "
        "melalui Project Charter—lalu diterjemahkan ke dalam Project Management Plan yang terperinci, diarahkan eksekusinya, dipantau variansnya, "
        "dikendalikan perubahannya melalui dewan formal, hingga penutupan resmi dan pengarsipan aset pengetahuan organisasi. Ketiadaan integrasi "
        "akan mengakibatkan proyek berjalan terfragmentasi, menimbulkan silo operasional, dan berujung pada kegagalan proyek secara sistemik.")

    # Soal 2
    add_heading_2(doc, "2. Sebutkan dan jelaskan secara singkat ketujuh proses dalam manajemen integrasi proyek!")
    add_body_p(doc, 
        "Berdasarkan standar Project Management Body of Knowledge (PMBOK Guide) Edisi Ke-6 (PMI, 2018, hal. 70–71), Manajemen Integrasi Proyek "
        "tersusun atas 7 (tujuh) proses terpadu yang terdistribusi ke dalam 5 (lima) Kelompok Proses Manajemen Proyek. "
        "Ringkasan ketujuh proses tersebut disajikan secara sistematis dalam Tabel 1 di bawah ini:")
        
    headers_t1 = ["No", "Nama Proses Integrasi", "Process Group", "Fokus dan Penjelasan Esensial", "Keluaran Utama (Key Outputs)"]
    data_t1 = [
        [
            "1", 
            "Develop Project Charter", 
            "Initiating", 
            "Proses mengembangkan dokumen formal yang mencatat otorisasi resmi keberadaan proyek dan memberikan wewenang kepada manajer proyek untuk mengalokasikan sumber daya organisasi.",
            "• Project Charter\n• Assumption Log"
        ],
        [
            "2", 
            "Develop Project Management Plan", 
            "Planning", 
            "Proses mendefinisikan, menyiapkan, dan mengoordinasikan seluruh rencana subsistem (subsidiary plans) serta baselines kinerja ke dalam satu rencana induk terpadu (Project Management Plan).",
            "• Project Management Plan (lengkap dengan baselines dan subsidiary plans)"
        ],
        [
            "3", 
            "Direct and Manage Project Work", 
            "Executing", 
            "Proses memimpin, mengarahkan, dan mengeksekusi aktivitas yang telah direncanakan dalam Project Management Plan serta mengimplementasikan approved change requests guna memproduksi deliverables.",
            "• Deliverables\n• Work Performance Data\n• Issue Log\n• Change Requests"
        ],
        [
            "4", 
            "Manage Project Knowledge", 
            "Executing", 
            "Proses memanfaatkan pengetahuan yang ada (existing knowledge) dan menciptakan pengetahuan baru (tacit & explicit) untuk mencapai sasaran proyek serta mendukung pembelajaran berkelanjutan organisasi.",
            "• Lessons Learned Register\n• Project Management Plan Updates\n• OPAs Updates"
        ],
        [
            "5", 
            "Monitor and Control Project Work", 
            "Monitoring & Controlling", 
            "Proses melacak, meninjau, dan melaporkan kemajuan serta kinerja proyek secara holistik terhadap target baseline guna mengidentifikasi deviasi dan merekomendasikan tindakan korektif/preventif.",
            "• Work Performance Reports\n• Change Requests\n• PMP & Project Docs Updates"
        ],
        [
            "6", 
            "Perform Integrated Change Control", 
            "Monitoring & Controlling", 
            "Proses meninjau seluruh usulan permintaan perubahan (change requests), mengevaluasi dampaknya secara komprehensif, menyetujui atau menolak perubahan melalui Change Control Board (CCB), serta mengelola baseline.",
            "• Approved Change Requests\n• Change Log\n• PMP & Project Docs Updates"
        ],
        [
            "7", 
            "Close Project or Phase", 
            "Closing", 
            "Proses memfinalisasi dan menyelesaikan seluruh kegiatan di seluruh process groups untuk menutup proyek, fase, atau kontrak secara administratif dan formal, termasuk serah terima deliverables akhir.",
            "• Final Product/Service Transition\n• Final Project Report\n• OPAs Updates (Arsip)"
        ]
    ]
    create_table(doc, [1.0, 3.2, 2.2, 5.4, 3.0], headers_t1, data_t1, 
                 [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.LEFT])

    # Soal 3
    add_heading_2(doc, "3. Apa yang dimaksud dengan Project Charter? Siapa yang menerbitkannya dan apa fungsinya bagi proyek?")
    add_body_p(doc, 
        "Project Charter (Piagam Proyek) adalah dokumen resmi yang diterbitkan oleh pemrakarsa proyek (project initiator) atau sponsor "
        "yang secara formal mengotorisasi keberadaan suatu proyek dan memberikan wewenang hukum serta manajerial kepada manajer proyek "
        "untuk menerapkan sumber daya organisasi dalam kegiatan-kegiatan proyek (PMI, 2018, hal. 75; Kerzner, 2017). "
        "Project Charter mendokumentasikan pemahaman tingkat tinggi (high-level understanding) mengenai tujuan proyek, sasaran terukur, "
        "kriteria keberhasilan, batasan awal, asumsi strategis, risiko utama, dan komitmen anggaran.")
        
    add_body_p(doc, 
        "Pihak yang Menerbitkan Project Charter:", bold_prefix="Pihak yang Berwenang: ")
    add_body_p(doc, 
        "Project Charter tidak diterbitkan oleh Manajer Proyek. Dokumen ini wajib diterbitkan dan disahkan oleh entitas di luar tim proyek "
        "yang memiliki wewenang pendanaan dan tata kelola organisasi, seperti: Sponsor Proyek (Project Sponsor), Komite Pengarah Portofolio "
        "(Portfolio/Program Steering Committee), Kantor Manajemen Proyek Korporat (Corporate PMO), atau jajaran Manajemen Eksekutif. "
        "Manajer proyek dapat berperan aktif dalam membantu perumusan draf awal charter melalui wawancara dan fasilitasi, namun tanda tangan "
        "pengesahan dan legitimasi kuasa mutlak berasal dari pihak sponsor.")
        
    add_body_p(doc, 
        "Fungsi Fundamental Project Charter bagi Keberhasilan Proyek:", bold_prefix="Fungsi Dokumen: ")
    add_num_p(doc, "a.", "Landasan Otoritas dan Mandat Formal Manajer Proyek:",
        "Memberikan legitimasi legal bagi Manajer Proyek untuk memimpin tim kerja, mengeluarkan instruksi kerja, memesan sumber daya bersama "
        "(shared resources seperti shared developers atau infrastruktur komputasi), serta menandatangani komitmen pembelanjaan modal proyek.")
    add_num_p(doc, "b.", "Kontrak Kemitraan Awal (Internal Agreement / Compact):",
        "Bertindak sebagai 'kontrak kesepakatan' tingkat tinggi antara sponsor penyedia dana dengan tim pelaksana proyek mengenai deliverables "
        "utama apa yang disepakati akan diserahkan dan batas toleransi kegagalan proyek.")
    add_num_p(doc, "c.", "Penyelarasan Strategis terhadap Nilai Bisnis (Strategic Alignment):",
        "Menghubungkan inisiatif proyek dengan dokumen kasus bisnis (Business Case) dan sasaran strategis jangka panjang korporat, menjamin "
        "bahwa proyek tidak dieksekusi atas dasar asumsi subjektif yang sia-sia melainkan memiliki justifikasi return on investment yang nyata.")
    add_num_p(doc, "d.", "Penetapan Batasan Awal untuk Mencegah Scope Creep Dini:",
        "Mengunci batasan makro mengenai apa yang masuk ke dalam cakupan proyek (in-scope) dan apa yang berada di luar cakupan proyek (out-of-scope), "
        "sehingga membentengi tim dari perluasan ekspektasi pemangku kepentingan yang tidak terkendali sebelum perencanaan rinci dimulai.")
    add_num_p(doc, "e.", "Titik Tolak Primer Perencanaan Rinci (Input to Develop PMP):",
        "Menjadi input paling mendasar dan acuan batas dalam mengembangkan rencana manajemen proyek, menyusun kamus WBS, mengumpulkan kebutuhan "
        "teknis terperinci, dan menyusun baseline jadwal terintegrasi.")

    # Soal 4
    add_heading_2(doc, "4. Sebutkan minimal delapan komponen utama yang terdapat dalam Project Management Plan!")
    add_body_p(doc, 
        "Project Management Plan adalah dokumen komprehensif yang mengintegrasikan seluruh rencana subsistem dari setiap area pengetahuan "
        "ke dalam satu kesatuan rencana induk yang utuh dan koheren (PMI, 2018, hal. 86–89; Larson & Gray, 2021). "
        "Berdasarkan standar PMBOK Guide Edisi Ke-6, komponen utama dalam Project Management Plan terbagi menjadi tiga kategori besar: "
        "Subsidiary Management Plans, Performance Measurement Baselines, dan Komponen Tata Kelola Tambahan. "
        "Berikut adalah 12 (dua belas) komponen utama yang menyusun rencana manajemen proyek:")
        
    add_num_p(doc, "1.", "Scope Management Plan:",
        "Menetapkan tata cara mendefinisikan, mengembangkan, memantau, memvalidasi, dan mengendalikan ruang lingkup proyek secara terstruktur.")
    add_num_p(doc, "2.", "Requirements Management Plan:",
        "Menguraikan bagaimana kebutuhan fungsional dan non-fungsional pemangku kepentingan akan dianalisis, didokumentasikan, dan dilacak status pemenuhannya (traceability).")
    add_num_p(doc, "3.", "Schedule Management Plan:",
        "Menetapkan metodologi penjadwalan, perangkat lunak yang digunakan, kriteria penentuan durasi aktivitas, tingkat presisi jadwal, serta mekanisme pemantauan varians jadwal.")
    add_num_p(doc, "4.", "Cost Management Plan:",
        "Menjelaskan bagaimana estimasi biaya disusun, struktur akun penganggaran (budgeting), alokasi cadangan kontingensi, serta teknik pengendalian biaya seperti Earned Value Management (EVM).")
    add_num_p(doc, "5.", "Quality Management Plan:",
        "Menetapkan standar kualitas yang wajib dipenuhi oleh deliverables proyek, metodologi pengujian (quality control), dan proses penjaminan kepatuhan prosedur (quality assurance).")
    add_num_p(doc, "6.", "Resource Management Plan:",
        "Memberikan panduan tentang bagaimana sumber daya manusia (tim proyek) dan sumber daya fisik (perangkat keras, server, lisensi) diidentifikasi, diperoleh, dialokasikan, dikelola, dan dilepaskan secara tertib.")
    add_num_p(doc, "7.", "Communications Management Plan:",
        "Menetapkan kebutuhan informasi pemangku kepentingan, format laporan status mingguan/bulanan, saluran komunikasi formal, frekuensi distribusi, dan alur eskalasi masalah.")
    add_num_p(doc, "8.", "Risk Management Plan:",
        "Mendefinisikan pendekatan identifikasi risiko, matriks probabilitas dan dampak, kriteria toleransi risiko organisasi, alokasi risk owner, serta strategi mitigasi dan kontingensi.")
    add_num_p(doc, "9.", "Procurement Management Plan:",
        "Menjelaskan strategi pengadaan barang/jasa dari pihak ketiga, pemilihan jenis kontrak (fixed price / time & material), proses tender, serta administrasi evaluasi vendor.")
    add_num_p(doc, "10.", "Stakeholder Engagement Plan:",
        "Menjabarkan strategi dan taktik komunikasi untuk melibatkan pemangku kepentingan secara efektif berdasarkan tingkat kepentingan dan pengaruh mereka guna meminimalisasi resistensi.")
    add_num_p(doc, "11.", "Tiga Performance Measurement Baselines (PMB):",
        "Tolok ukur terpadu yang disetujui untuk mengukur kinerja aktual proyek, mencakup: (a) Scope Baseline (Project Scope Statement, WBS, dan WBS Dictionary); "
        "(b) Schedule Baseline (jadwal resmi dengan tanggal mulai, selesai, dan milestone); serta (c) Cost Baseline (anggaran bertahap waktu di luar management reserve).")
    add_num_p(doc, "12.", "Change Management Plan & Configuration Management Plan:",
        "Menetapkan prosedur formal tata kelola pengajuan, evaluasi dampak, dan persetujuan perubahan proyek melalui CCB, serta mekanisme pelacakan versi konfigurasi artefak sistem.")

    # Soal 5
    add_heading_2(doc, "5. Apa perbedaan antara Project Charter dan Project Management Plan dalam hal tujuan, isi, dan waktu pembuatannya?")
    add_body_p(doc, 
        "Meskipun keduanya merupakan dokumen kunci dalam tata kelola manajemen proyek, Project Charter dan Project Management Plan "
        "memiliki perbedaan fundamental yang mencakup dimensi filosofis, tingkat kedalaman teknis, otoritas pengesahan, dan posisi dalam "
        "siklus hidup proyek (PMI, 2018; Kerzner, 2017). "
        "Matriks perbandingan komparatif komprehensif antara kedua dokumen ini disajikan secara mendalam pada Tabel 2 berikut:")
        
    headers_t2 = ["Dimensi Pembanding", "Project Charter (Piagam Proyek)", "Project Management Plan (Rencana Manajemen Proyek)"]
    data_t2 = [
        [
            "Tujuan Pokok (Primary Purpose)",
            "Mengotorisasi eksistensi proyek secara resmi, menyelaraskan proyek dengan strategi bisnis korporat, dan memberikan mandat formal kepada Manajer Proyek.",
            "Menjadi pedoman operasional dan peta navigasi komprehensif bagi tim mengenai bagaimana proyek dieksekusi, dipantau, dikendalikan, dan ditutup."
        ],
        [
            "Waktu Pembuatan (Timing in Life Cycle)",
            "Disusun pada Fase Inisiasi (Initiating Process Group), sebelum proyek disetujui secara formal dan sebelum perencanaan rinci dimulai.",
            "Disusun pada Fase Perencanaan (Planning Process Group) setelah Charter disahkan, serta dikembangkan secara bertahap (progressive elaboration)."
        ],
        [
            "Pihak Penyusun & Pengesah",
            "Diterbitkan dan disahkan oleh Sponsor Proyek / Manajemen Eksekutif / PMO. Manajer Proyek dapat membantu menyusun draf.",
            "Disusun secara kolaboratif oleh Manajer Proyek bersama tim pengembang dan para ahli fungsional, lalu disetujui oleh Sponsor dan Key Stakeholders."
        ],
        [
            "Tingkat Kerincian (Level of Detail)",
            "Tingkat tinggi (High-level / Macro). Bersifat ikhtisar strategis tanpa detail operasional mikro.",
            "Sangat terperinci (Detailed / Micro). Memuat dekomposisi pekerjaan mikro (WBS), estimasi kuantitatif hari/jam, dan matriks penugasan teknis."
        ],
        [
            "Isi Utama Dokumen",
            "Tujuan strategis, business case summary, high-level scope, milestone kasar, plafon estimasi anggaran awal, risiko makro, dan batas wewenang PM.",
            "10 subsidiary management plans, 3 baselines kinerja resmi (Scope, Schedule, Cost), metrik kualitas, rencana respons risiko, dan change control procedure."
        ],
        [
            "Fleksibilitas & Dinamika Perubahan",
            "Sangat statis. Jarang mengalami perubahan kecuali terjadi redefinisi bisnis fundamental yang disepakati langsung oleh sponsor.",
            "Dokumen hidup (Living document). Dinamis dan terus diperbarui secara formal melalui proses Perform Integrated Change Control sepanjang siklus proyek."
        ],
        [
            "Ukuran & Kompleksitas Dokumen",
            "Ringkas dan padat, umumnya berkisar antara 2 hingga 5 halaman dokumen formal.",
            "Komprehensif, multi-bagian, dan tebal, sering kali mencapai puluhan hingga ratusan halaman bergantung pada skala proyek."
        ]
    ]
    create_table(doc, [3.2, 5.8, 5.8], headers_t2, data_t2, 
                 [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.JUSTIFY])

    # Soal 6
    add_heading_2(doc, "6. Mengapa proses Manage Project Knowledge penting dilakukan, dan apa manfaatnya bagi organisasi?")
    add_body_p(doc, 
        "Proses Manage Project Knowledge adalah proses memanfaatkan pengetahuan yang ada (baik tacit maupun explicit) serta menciptakan "
        "pengetahuan baru guna mencapai tujuan proyek saat ini dan memperkaya kapabilitas pembelajaran organisasi di masa depan (PMI, 2018, hal. 98–104). "
        "Proses ini diperkenalkan secara khusus sebagai proses mandiri dalam PMBOK Guide Edisi Ke-6 untuk menegaskan bahwa proyek bukan hanya "
        "menghasilkan luaran fisik/digital, melainkan juga wadah penghasil intelektual dan pengalaman yang sangat bernilai.")
        
    add_body_p(doc, 
        "Urgensi Fundamental Mengapa Proses Ini Wajib Dilakukan:", bold_prefix="Urgensi Pelaksanaan: ")
    add_num_p(doc, "1.", "Mencegah Terjadinya Fenomena 'Amnesia Korporat' (Corporate Amnesia):",
        "Setiap proyek menghadapi rintangan teknis dan dinamika manajerial yang unik. Apabila pengetahuan tersebut tidak dikelola, "
        "seluruh wawasan, pemecahan masalah kritis, dan solusi arsitektur akan lenyap begitu proyek selesai atau saat personel tim berpindah unit/perusahaan (turnover).")
    add_num_p(doc, "2.", "Menjembatani Pengetahuan Eksplisit (Explicit) dan Pengetahuan Tersirat (Tacit Knowledge):",
        "Pengetahuan eksplisit (dokumen spesifikasi, kode program, konfigurasi server) sangat mudah diarsipkan dalam repositori data. "
        "Namun, pengetahuan tacit (kepekaan teknis, intuisi pemecahan bug rumit, teknik negosiasi dengan mandor lapangan) berada di dalam benak manusia. "
        "Manage Project Knowledge menciptakan suasana interaksi sosial yang kondusif (seperti knowledge-sharing sessions, peer reviews, dan komunitas praktik) "
        "agar pengetahuan tacit dapat ditransformasikan dan diserap oleh anggota tim lainnya.")
    add_num_p(doc, "3.", "Mengeliminasi Pemborosan Biaya Akibat 'Menemukan Kembali Roda' (Reinventing the Wheel):",
        "Tanpa transfer pengetahuan, tim pada sub-proyek baru akan membuang ratusan jam kerja dan biaya komputasi untuk memecahkan kendala "
        "yang sebenarnya telah berhasil diselesaikan oleh tim proyek terdahulu.")
        
    add_body_p(doc, 
        "Manfaat Strategis Konkret bagi Organisasi:", bold_prefix="Manfaat Organisasional: ")
    add_bullet_p(doc, "Akselerasi Kurva Pembelajaran (Steeper Learning Curve):", 
        "Mempersingkat masa adaptasi (onboarding) developer atau peneliti baru, sehingga produktivitas tim dapat dicapai lebih cepat.")
    add_bullet_p(doc, "Peningkatan Tingkat Kematangan Tata Kelola (Organizational Maturity):", 
        "Memperkaya Organizational Process Assets (OPAs) dengan repositori lessons learned yang teruji, template dokumen yang disempurnakan, dan basis data metrik kinerja historis.")
    add_bullet_p(doc, "Reduksi Risiko Berulang dan Efisiensi Anggaran:", 
        "Menghindarkan organisasi dari kegagalan teknis serupa yang memicu penalti kontrak, serta meningkatkan daya saing korporat dalam mengelola portofolio proyek kompleks.")

    # Soal 7
    add_heading_2(doc, "7. Mengapa setiap permintaan perubahan dalam proyek harus melalui proses Perform Integrated Change Control?")
    add_body_p(doc, 
        "Dalam eksekusi proyek perangkat lunak dan rekayasa kompleks, perubahan kebutuhan adalah suatu keniscayaan yang tidak dapat dihindari. "
        "Namun, membiarkan perubahan terjadi secara bebas tanpa kendali merupakan penyebab paling dominan atas kegagalan proyek IT di seluruh dunia. "
        "Proses Perform Integrated Change Control (PICC) adalah gerbang pengendalian formal yang bertugas meninjau seluruh permintaan perubahan "
        "(change requests), menyetujui atau menolaknya, serta mengelola implementasi perubahan pada deliverables, baselines, dan dokumen proyek (PMI, 2018, hal. 113–120).")
        
    add_body_p(doc, 
        "Alasan Krusial Mengapa Setiap Perubahan Wajib Melalui Proses PICC:", bold_prefix="Alasan Fundamental: ")
    add_num_p(doc, "1.", "Evaluasi Dampak Holistik Lintas Batasan (Comprehensive Multi-Constraint Impact Analysis):",
        "Proyek adalah suatu ekosistem yang terikat erat. Penambahan sebuah fitur perangkat lunak tidak hanya membutuhkan penambahan baris kode, "
        "tetapi juga berdampak pada: perubahan skema basis data, lonjakan waktu kompilasi, penambahan beban uji kualitas QA, pemanjangan durasi jalur kritis (schedule), "
        "lonjakan biaya lembur (cost), hingga kemunculan risiko celah keamanan baru. PICC mewajibkan dilakukannya analisis dampak menyeluruh sebelum keputusan diambil.")
    add_num_p(doc, "2.", "Pencegahan Mutlak terhadap Sindrom Scope Creep:",
        "Scope creep adalah fenomena perluasan ruang lingkup proyek secara perlahan tanpa persetujuan formal, tanpa kompensasi penambahan waktu (jadwal), "
        "dan tanpa alokasi anggaran tambahan. Tanpa PICC, developer cenderung menyetujui permintaan fitur tambahan dari pengguna secara lisan di koridor "
        "atau pesan instan (informal agreements), yang berujung pada pembengkakan beban kerja dan kegagalan memenuhi tenggat waktu.")
    add_num_p(doc, "3.", "Menjaga Keabsahan dan Integritas Performance Measurement Baseline (PMB):",
        "Evaluasi kinerja proyek menggunakan teknik Earned Value Management (seperti indeks SPI dan CPI) hanya valid jika baseline rencana yang menjadi tolok ukur "
        "memiliki integritas tinggi. Jika baseline dapat diubah semena-mena oleh siapa pun, manajemen akan kehilangan kompas evaluasi yang objektif.")
    add_num_p(doc, "4.", "Tata Kelola Transparan Melalui Dewan Pengendali Perubahan (Change Control Board / CCB):",
        "PICC menetapkan struktur wewenang yang tegas di mana keputusan perubahan strategis didelegasikan kepada Change Control Board (CCB) "
        "yang beranggotakan sponsor, perwakilan pengguna kunci, manajer teknis, dan manajer proyek. Seluruh riwayat pengajuan, analisis, alasan penolakan, "
        "atau rincian persetujuan dicatat secara transparan di dalam Change Log.")

    # Soal 8
    add_heading_2(doc, "8. Apa saja kegiatan yang dilakukan dalam proses Close Project or Phase?")
    add_body_p(doc, 
        "Proses Close Project or Phase adalah proses memfinalisasi seluruh kegiatan di seluruh Kelompok Proses Manajemen Proyek untuk menyelesaikan "
        "proyek atau fase proyek secara administratif dan formal (PMI, 2018, hal. 121–127; Larson & Gray, 2021). "
        "Penutupan yang tertib memastikan bahwa proyek tidak berakhir secara menggantung (abrupt termination), seluruh kewajiban kontraktual telah tuntas, "
        "dan organisasi memetik manfaat maksimal dari hasil investasi. Delapan kegiatan utama yang wajib dilaksanakan dalam proses ini meliputi:")
        
    add_num_p(doc, "1.", "Verifikasi dan Penerimaan Formal Deliverables Akhir (Final Acceptance Sign-Off):",
        "Memastikan bahwa seluruh produk, modul perangkat lunak, atau luaran riset yang telah divalidasi mutu teknisnya secara formal disetujui "
        "dan ditandatangani berita acara penerimaannya (User Acceptance Testing sign-off) oleh sponsor proyek atau klien.")
    add_num_p(doc, "2.", "Serah Terima dan Transisi Produk Akhir (Product Handover & Transition):",
        "Menyerahkan hasil akhir proyek beserta seluruh dokumentasi arsitektur, manual operasional sistem, hak akses repositori, dan lisensi "
        "kepada unit operasional (seperti tim IT Support, tim pemeliharaan, atau unit bisnis perkebunan pengguna sistem).")
    add_num_p(doc, "3.", "Audit Administratif dan Penutupan Rekening Biaya Proyek (Financial & Cost Closure):",
        "Melakukan rekonsiliasi keuangan menyeluruh, menyelesaikan pembayaran faktur tertunda kepada vendor, melakukan audit pengeluaran terhadap "
        "cost baseline, dan menutup akun pembukuan proyek agar tidak ada lagi pembebanan biaya liar.")
    add_num_p(doc, "4.", "Penyelesaian dan Penutupan Kontrak Pengadaan (Procurement & Contract Closure):",
        "Memastikan semua klaim kontraktual dengan penyedia layanan pihak ketiga (misal vendor cloud GPU, konsultan eksternal) telah dipenuhi, "
        "mengevaluasi kinerja vendor, serta menerbitkan surat penyelesaian kontrak formal.")
    add_num_p(doc, "5.", "Dokumentasi dan Pengarsipan Lessons Learned (Post-Project Review / Post-Mortem):",
        "Menyelenggarakan pertemuan evaluasi akhir bersama tim proyek untuk mendiskusikan apa yang telah berhasil dengan baik, kendala yang dihadapi, "
        "kesalahan estimasi yang terjadi, dan rekomendasi perbaikan untuk dicatatkan ke dalam Lessons Learned Register.")
    add_num_p(doc, "6.", "Pembaruan dan Pengarsipan Aset Proses Organisasi (Archiving Historical Records):",
        "Mengarsipkan seluruh dokumen perencanaan, baseline jadwal aktual, diagram arsitektur, log perubahan, dan dataset ke dalam sistem repositori "
        "terpusat korporat yang aman sebagai referensi berharga bagi proyek-proyek di masa depan.")
    add_num_p(doc, "7.", "Pelepasan dan Realokasi Sumber Daya Proyek (Resource Release & Demobilization):",
        "Secara resmi melepaskan fasilitas fisik, menghentikan sewa cloud computing instance yang tidak lagi digunakan, serta mengembalikan anggota "
        "tim proyek (shared software developers) ke departemen fungsional masing-masing atau ke proyek baru.")
    add_num_p(doc, "8.", "Penyusunan dan Distribusi Laporan Kinerja Akhir Proyek (Final Project Report):",
        "Menyusun laporan komprehensif yang merangkum pencapaian ruang lingkup, realisasi jadwal (variance SPI), realisasi anggaran (variance CPI), "
        "derajat pemenuhan kriteria keberhasilan bisnis, dan ringkasan risiko yang berhasil dikelola, lalu mendistribusikannya kepada para pemangku kepentingan.")

print("Section A Builder written successfully.")
