# -*- coding: utf-8 -*-
"""
Section D Builder: Refleksi Mandiri
Pertemuan 8: Manajemen Kualitas Proyek
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_d(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "D. REFLEKSI MANDIRI")

    add_body_p(doc,
        "Bagian ini memuat refleksi kritis dan introspeksi personal saya terhadap penguasaan materi Manajemen Kualitas Proyek "
        "yang dipelajari pada Pertemuan 8, serta bagaimana prinsip-prinsip penjaminan mutu diaplikasikan dalam konteks riset "
        "proyek kelompok PRE-02 (Sistem Prediksi Probabilitas Keterlambatan Proyek Berbasis Machine Learning sebagai Early Warning System):")

    # =========================================================================
    # REFLEKSI 1
    # =========================================================================
    add_heading_2(doc, "1. Menurut Anda, mengapa manajemen kualitas sering dianggap sebagai investasi, bukan biaya, dalam proyek? Berikan analisis dengan mengacu pada konsep Cost of Quality!")
    add_body_p(doc,
        "Menurut analisis saya, pandangan tradisional yang menganggap kualitas semata-mata sebagai beban pengeluaran (cost burden) "
        "terjadi karena paradigma sempit yang hanya melihat uang keluar di fase awal tanpa memperhitungkan potensi kehancuran finansial "
        "di fase akhir. Mengacu pada kerangka Cost of Quality (COQ), manajemen kualitas adalah bentuk investasi strategis dengan tingkat "
        "pengembalian (Return on Investment / ROI) yang sangat tinggi.")
    add_body_p(doc,
        "Ketika kita mengalokasikan anggaran untuk Biaya Pencegahan (Prevention Cost)—seperti merancang arsitektur kode yang bersih, "
        "mengadakan pelatihan tim, dan mengonfigurasi otomatisasi pengujian—kita pada hakikatnya sedang 'membeli asuransi' untuk mengeliminasi "
        "Biaya Kegagalan (Failure Cost). Berdasarkan prinsip The 1-10-100 Rule, dana sebesar Rp 1 juta yang diinvestasikan pada pencegahan "
        "mampu menyelamatkan organisasi dari kerugian Rp 100 juta akibat kegagalan eksternal di lingkungan produksi (seperti downtime sistem, "
        "kehilangan kepercayaan klien, atau denda pelanggaran SLA). Oleh karena itu, pengeluaran kualitas adalah investasi pelindung nilai proyek.")

    # =========================================================================
    # REFLEKSI 2
    # =========================================================================
    add_heading_2(doc, "2. Dari tujuh alat kualitas dasar yang dipelajari, alat mana yang menurut Anda paling mudah dipahami dan paling sulit dipahami? Jelaskan alasan Anda!")
    add_body_p(doc,
        "Evaluasi komparatif pemahaman saya terhadap Tujuh Alat Kualitas Dasar (Seven Basic Quality Tools):")
    add_bullet_p(doc, "Paling Mudah Dipahami – Check Sheet (Lembar Periksa) dan Pareto Diagram: Check sheet sangat intuitif karena berupa formulir tabulasi data kejadian yang konkret. Sementara itu, Diagram Pareto memberikan visualisasi hierarki yang sangat lugas dengan menerapkan prinsip 80/20, sehingga tim dapat langsung melihat secara visual mana masalah yang krusial (vital few) dan mana yang sekunder tanpa membutuhkan komputasi statistik yang rumit.")
    add_bullet_p(doc, "Paling Sulit Dipahami – Control Chart (Bagan Kendali): Alat ini menuntut pemahaman mendalam tentang teori probabilitas dan statistik inferensial, khususnya dalam menentukan Batas Kendali Atas (UCL) dan Batas Kendali Bawah (LCL) sebesar 3 standar deviasi (±3σ). Tantangan terbesarnya adalah membedakan secara tepat antara \"Variasi Alami / Common Cause Variation\" (yang memang inheren dalam sistem dan tidak perlu diover-reaksi) dengan \"Variasi Khusus / Special Cause Variation\" (yang merupakan anomali kritis yang menuntut intervensi langsung manajer proyek).")

    # =========================================================================
    # REFLEKSI 3
    # =========================================================================
    add_heading_2(doc, "3. Bagaimana Anda akan menggunakan Fishbone Diagram dan Pareto Chart secara bersamaan untuk menyelesaikan masalah kualitas dalam proyek? Jelaskan urutan penggunaan kedua alat tersebut!")
    add_body_p(doc,
        "Dalam metodologi pemecahan masalah kualitas (Problem Solving Workflow), Fishbone Diagram dan Pareto Chart "
        "bekerja secara komplementer dalam suatu alur kerja dua arah yang terintegrasi:")
    add_num_p(doc, "Langkah 1 – Menggunakan Pareto Chart untuk Menentukan PRIORITAS (Fokus Masalah): Ketika proyek menghadapi puluhan jenis keluhan atau cacat produk, Pareto Chart dijalankan terlebih dahulu untuk menyaring dan mengidentifikasi 20% jenis cacat yang menyumbang 80% dampak negatif. Ini mencegah tim membuang energi menganalisis hal-hal kecil yang tidak signifikan.", 1)
    add_num_p(doc, "Langkah 2 – Menggunakan Fishbone Diagram untuk Menganalisis AKAR PENYEBAB: Setelah cacat prioritas nomor satu berhasil diisolasi melalui Pareto (misalnya Cacat A), barulah Fishbone Diagram digambar untuk membedah masalah tunggal tersebut ke dalam kategori 5M (Man, Machine, Method, Material, Measurement) guna menemukan root cause paling mendasar.", 2)
    add_num_p(doc, "Langkah 3 – Validasi Ulang dengan Pareto Chart Kedua (Opsi Lanjutan): Setelah akar-akar penyebab dipetakan, frekuensi kemunculan masing-masing akar penyebab dapat kembali ditabulasi ke dalam Diagram Pareto kedua untuk memilih akar penyebab mana yang paling sering memicu terjadinya cacat tersebut.", 3)

    # =========================================================================
    # REFLEKSI 4
    # =========================================================================
    add_heading_2(doc, "4. Menurut Anda, apakah kualitas selalu berarti biaya yang lebih tinggi? Jelaskan dengan contoh dari proyek kelompok Anda!")
    add_body_p(doc,
        "Tidak. Mengidentikkan kualitas tinggi dengan biaya yang selalu lebih mahal adalah mitos yang keliru. "
        "Hal ini terjadi karena kerancuan antara konsep Kualitas (Quality) dengan Tingkat Kemewahan (Grade). "
        "Suatu produk perangkat lunak dapat memiliki fitur yang sederhana (Low Grade) namun memiliki Kualitas Sangat Tinggi (High Quality) "
        "karena bekerja secara sempurna tanpa bug, stabil, dan sesuai dengan seluruh kebutuhan pengguna.")
    add_body_p(doc,
        "Contoh Riil pada Proyek PRE-02: Membangun sistem prediksi keterlambatan proyek tidak memerlukan pengadaan superkomputer kluster "
        "berbiaya ratusan juta atau arsitektur Deep Learning 100-layer yang membengkakkan tagihan cloud GPU. Model yang lebih sederhana seperti "
        "Logistic Regression atau LightGBM yang dioptimasi dengan kalibrasi probabilitas (Platt Scaling) terbukti menghasilkan kualitas prediksi "
        "yang sangat unggul (Expected Calibration Error rendah dan Brier Score < 0.18), namun dengan biaya komputasi yang hampir gratis dan "
        "latensi inferensi milidetik. Menghilangkan data duplikat dan menerapkan feature engineering yang bersih justru menurunkan beban biaya komputasi server.")

    # =========================================================================
    # REFLEKSI 5
    # =========================================================================
    add_heading_2(doc, "5. Bagaimana Anda akan memastikan bahwa kualitas tetap terjaga di tengah tekanan untuk menyelesaikan proyek tepat waktu dan sesuai anggaran?")
    add_body_p(doc,
        "Ketika proyek berada di bawah tekanan kompresi jadwal (schedule compression) dan batasan anggaran ketat, "
        "kualitas sering kali menjadi korban pertama yang dikorbankan melalui pemotongan sesi pengujian (skimping on testing). "
        "Sebagai Manajer Proyek, saya akan menerapkan empat strategi pertahanan kualitas yang kokoh:")
    add_num_p(doc, "Menegakkan 'Definition of Done' (DoD) yang Tidak Dapat Dinegosiasikan: Menetapkan aturan bahwa tidak ada deliverables yang dapat disetujui (sign-off) jika belum lolos serangkaian kriteria penerimaan kualitas minimal, seperti passing automated unit tests dan peer review approval.", 1)
    add_num_p(doc, "Otomatisasi Penuh Saluran Pengujian (Automated Quality Gates): Memanfaatkan pipeline CI/CD yang menolak proses build secara otomatis jika terjadi regresi performa atau penurunan coverage pengujian. Otomatisasi menjamin pengujian berjalan cepat tanpa menambah jam kerja lembur tim.", 2)
    add_num_p(doc, "Mengendalikan Ruang Lingkup, Bukan Memotong Kualitas (Scope Triaging): Jika keterlambatan tak terhindarkan, solusi yang tepat adalah menegosiasikan pemangkasan fitur opsional (descope lower-priority features) bersama sponsor melalui Change Control Board, bukan menurunkan standar pengujian dari fitur-fitur kritis yang tetap dirilis.", 3)
    add_num_p(doc, "Kewaspadaan Terhadap Utang Teknis (Technical Debt Management): Menyadarkan seluruh pemangku kepentingan bahwa mengorbankan kualitas demi mengejar deadline semu akan menciptakan 'utang teknis' beracun yang pada akhirnya melumpuhkan sistem dan memicu pembengkakan biaya pemeliharaan di masa depan.", 4)
