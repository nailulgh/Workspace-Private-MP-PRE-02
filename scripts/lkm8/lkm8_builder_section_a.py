# -*- coding: utf-8 -*-
"""
Section A Builder: Tugas Tertulis Individu
Pertemuan 8: Manajemen Kualitas Proyek
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_a(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "A. TUGAS TERTULIS INDIVIDU")

    # =========================================================================
    # SOAL 1
    # =========================================================================
    add_heading_2(doc, "1. Apa yang dimaksud dengan manajemen kualitas proyek dan mengapa area pengetahuan ini penting dalam keberhasilan proyek?")
    add_body_p(doc,
        "Manajemen Kualitas Proyek (Project Quality Management) adalah area pengetahuan manajemen proyek yang mencakup proses-proses "
        "untuk menggabungkan kebijakan kualitas organisasi mengenai perencanaan, pengelolaan, dan pengendalian persyaratan kualitas proyek "
        "dan produk guna memenuhi ekspektasi serta kebutuhan para pemangku kepentingan (PMBOK® Guide 6th Edition, PMI, 2017). "
        "Secara konseptual, kualitas dalam manajemen proyek tidak didefinisikan sebagai kemewahan (luxury) atau kecanggihan fitur tanpa batas, "
        "melainkan didefinisikan secara presisi sebagai dua pilar utama:")
    add_bullet_p(doc, "Kesesuaian terhadap Persyaratan (Conformance to Requirements): Derajat di mana proses eksekusi dan artefak yang diproduksi secara ketat mematuhi spesifikasi desain, standar teknis, dan kontrak yang telah disepakati.")
    add_bullet_p(doc, "Kesesuaian untuk Digunakan (Fitness for Use): Kemampuan produk atau hasil akhir proyek untuk benar-benar memuaskan kebutuhan operasional pengguna nyata tanpa menimbulkan friksi atau kegagalan fungsional (Juran & Godfrey, 1999).")
    
    add_body_p(doc,
        "Area pengetahuan ini memegang peranan krusial dalam menentukan keberhasilan proyek karena tiga alasan strategis:")
    add_num_p(doc, "Menjaga Keseimbangan Segitiga Manajemen Proyek (Iron Triangle): Menghasilkan proyek tepat waktu dan sesuai anggaran menjadi sia-sia jika produk yang diserahkan cacat, tidak stabil, atau ditolak pengguna. Kualitas menjadi penentu nilai substantif dari deliverables proyek.", 1)
    add_num_p(doc, "Mencegah Eskalasi Biaya Pengerjaan Ulang (Rework): Dalam rekayasa perangkat lunak dan sistem kecerdasan buatan, bug atau cacat arsitektur yang terdeteksi di fase akhir membutuhkan biaya perbaikan hingga puluhan kali lipat lebih besar dibandingkan jika dicegah sejak awal.", 2)
    add_num_p(doc, "Menjamin Keandalan Sistem Kritis (Konteks Proyek PRE-02): Pada proyek pengembangan Sistem Peringatan Dini Keterlambatan Proyek Berbasis Machine Learning (PRE-02), kualitas kalibrasi probabilitas dan validasi model menentukan keakuratan sinyal risiko. Kualitas yang buruk akan menghasilkan alarm palsu (false alarms) yang merusak kredibilitas pengambilan keputusan manajerial.", 3)

    # =========================================================================
    # SOAL 2
    # =========================================================================
    add_heading_2(doc, "2. Sebutkan dan jelaskan tiga proses dalam manajemen kualitas proyek beserta hubungan antar ketiganya!")
    add_body_p(doc,
        "Berdasarkan standar PMBOK® Guide 6th Edition (PMI, 2017, hal. 271–305), manajemen kualitas proyek terdiri dari tiga proses "
        "komprehensif yang berjalan beriringan sepanjang siklus hidup proyek:")
    
    add_num_p(doc, "Plan Quality Management (Merencanakan Manajemen Kualitas): Proses mengidentifikasi standar kualitas, regulasi, dan spesifikasi yang relevan bagi proyek dan deliverables-nya, serta mendokumentasikan bagaimana proyek akan mendemonstrasikan kepatuhan terhadap standar tersebut. Proses ini menghasilkan Quality Management Plan dan Quality Metrics yang terukur.", 1)
    add_num_p(doc, "Manage Quality (Mengelola Kualitas): Proses menerjemahkan rencana kualitas ke dalam aktivitas kualitas yang dapat dieksekusi, memastikan bahwa kebijakan kualitas organisasi diterapkan secara konsisten, serta mengevaluasi efektivitas proses produksi melalui audit kualitas dan analisis proses (Quality Assurance). Fokusnya adalah pada perbaikan proses (process improvement).", 2)
    add_num_p(doc, "Control Quality (Mengendalikan Kualitas): Proses memantau dan mencatat hasil pelaksanaan aktivitas kualitas untuk menilai kinerja teknis, memverifikasi kesesuaian deliverables fisik terhadap kriteria penerimaan (acceptance criteria) melalui inspeksi dan pengujian, serta merekomendasikan penyesuaian jika ditemukan cacat (Quality Control).", 3)

    add_body_p(doc,
        "Hubungan Antar Ketiga Proses: Ketiga proses ini membentuk siklus umpan balik tertutup (closed-loop feedback cycle) "
        "yang mengadopsi siklus Deming (Plan-Do-Check-Act):")
    add_bullet_p(doc, "Plan Quality Management menetapkan tolok ukur (standar acuan dan metrik keberhasilan) pada fase perencanaan (Plan).")
    add_bullet_p(doc, "Manage Quality menerapkan prosedur penjaminan kualitas selama proses eksekusi berlangsung (Do), memastikan tim bekerja dengan metode yang benar guna mencegah timbulnya cacat.")
    add_bullet_p(doc, "Control Quality menginspeksi dan menguji luaran fisik/artefak (Check). Jika ditemukan penyimpangan atau bug, data hasil pengujian (quality control measurements) diumpankan kembali ke proses Manage Quality untuk dianalisis akar penyebabnya (Root Cause Analysis) dan dilakukan tindakan korektif/preventif (Act).")

    # =========================================================================
    # SOAL 3
    # =========================================================================
    add_heading_2(doc, "3. Apa perbedaan antara kualitas produk (product quality) dan kualitas proses (process quality) dalam konteks manajemen proyek?")
    add_body_p(doc,
        "Dalam manajemen proyek modern, kualitas dibedakan secara tegas antara apa yang dihasilkan (produk/artefak) "
        "dan bagaimana artefak tersebut diproduksi (metodologi/alur kerja):")

    headers_q3 = ["Dimensi Komparasi", "Kualitas Produk (Product Quality)", "Kualitas Proses (Process Quality)"]
    data_q3 = [
        ["Fokus Utama", "Luaran fisik, deliverables, atau artefak akhir yang diserahkan ke pengguna.", "Metode kerja, prosedur, alur eksekusi, dan teknik rekayasa yang digunakan tim."],
        ["Orientasi Pengujian", "Verifikasi dan validasi spesifikasi fungsional, performa, stabilitas, dan keamanan produk.", "Audit kepatuhan SOP, tinjauan arsitektur (code review), integrasi berkelanjutan (CI/CD), dan efisiensi alur."],
        ["Tujuan Akhir", "Memastikan produk bebas dari bug dan bekerja sesuai kebutuhan bisnis pemangku kepentingan.", "Mencegah kesalahan sistemik sejak awal dan memungkinkan peningkatan berkesinambungan (continuous improvement)."],
        ["Domain Pengawasan", "Ranah Control Quality (Quality Control / QC) melalui inspeksi, testing, dan QA testing.", "Ranah Manage Quality (Quality Assurance / QA) melalui quality audit dan process analysis."],
        ["Contoh Nyata pada Proyek PRE-02", "Model ML memiliki nilai AUC-ROC ≥ 0.80, Brier Score < 0.18, ECE rendah, dan respons inferensi API < 200 ms.", "Penerapan automated testing, standardisasi konvensi penamaan kode, cross-validation terisolasi, dan versioning dataset."],
    ]
    create_table(doc, headers_q3, data_q3)

    add_body_p(doc,
        "Hubungan Kausal: Kualitas produk yang unggul merupakan konsekuensi logis dari kualitas proses yang matang. "
        "Proses yang tidak disiplin dan semrawut hampir pasti menghasilkan produk yang sarat cacat tersembunyi, "
        "meskipun diuji berulang kali di akhir proyek.")

    # =========================================================================
    # SOAL 4
    # =========================================================================
    add_heading_2(doc, "4. Sebutkan dan jelaskan empat komponen Cost of Quality (COQ)! Mana yang paling efektif untuk diinvestasikan dan mengapa?")
    add_body_p(doc,
        "Cost of Quality (COQ) mencakup seluruh biaya moneter yang dikeluarkan untuk memastikan kualitas yang sesuai serta biaya "
        "yang timbul akibat terjadinya kualitas yang buruk (PMI, 2017). COQ diklasifikasikan ke dalam dua pilar utama dengan empat komponen:")
    
    add_num_p(doc, "Biaya Kesesuaian (Cost of Conformance): Uang yang dihabiskan untuk mencegah dan mendeteksi kesalahan.", 1)
    add_subbullet_p(doc, "Prevention Costs (Biaya Pencegahan): Investasi yang dikeluarkan untuk mendesain proses kerja agar bebas dari kesalahan sejak awal. Mencakup pelatihan tim, standardisasi coding guidelines, perencanaan arsitektur modular, dan pengadaan tools otomatis.")
    add_subbullet_p(doc, "Appraisal Costs (Biaya Penilaian/Pemeriksaan): Biaya yang dikeluarkan untuk mengevaluasi, mengukur, dan mengaudit produk guna memastikan kepatuhan standar sebelum produk dirilis. Mencakup kegiatan unit testing, peer code review, continuous integration testing, dan kalibrasi alat ukur.")

    add_num_p(doc, "Biaya Ketidaksesuaian (Cost of Nonconformance): Uang yang terbuang sia-sia akibat terjadinya kegagalan kualitas.", 2)
    add_subbullet_p(doc, "Internal Failure Costs (Biaya Kegagalan Internal): Biaya yang timbul ketika cacat atau bug terdeteksi oleh tim sebelum produk diserahkan ke klien/pengguna akhir. Mencakup biaya perbaikan bug (bug fixing), refactoring darurat, pelatihan ulang model ML yang overfit, dan testing ulang (regression testing).")
    add_subbullet_p(doc, "External Failure Costs (Biaya Kegagalan Eksternal): Biaya yang muncul ketika cacat lolos ke lingkungan produksi dan dialami langsung oleh pengguna/klien. Mencakup klaim garansi, kompensasi kerugian, downtime server, hilangnya reputasi organisasi, hingga tuntutan hukum.")

    add_body_p(doc,
        "Komponen yang Paling Efektif untuk Diinvestasikan: Prevention Costs (Biaya Pencegahan) adalah komponen investasi yang paling efektif dan berdaya guna tinggi. "
        "Alasannya, setiap rupiah yang dialokasikan untuk pencegahan akan mengeliminasi potensi munculnya cacat di sumbernya (source elimination), "
        "sehingga secara drastis memangkas biaya appraisal dan mencegah munculnya biaya kegagalan internal maupun eksternal yang eksponensial.")

    # =========================================================================
    # SOAL 5
    # =========================================================================
    add_heading_2(doc, "5. Mengapa biaya pencegahan (prevention cost) dianggap lebih efektif dibandingkan biaya inspeksi dan biaya kegagalan? Jelaskan dengan logika sederhana!")
    add_body_p(doc,
        "Keunggulan biaya pencegahan dapat dipahami melalui The 1-10-100 Rule (Aturan Kelipatan Sepuluh Kualitas) "
        "yang dirumuskan oleh George Labovitz dan Yu Sang Chang:")
    add_bullet_p(doc, "Fase Pencegahan ($1): Menghabiskan Rp 1 untuk melatih tim tentang clean architecture dan penanganan data outliers mencegah masuknya logika salah ke dalam kode sumber.")
    add_bullet_p(doc, "Fase Penilaian/Inspeksi ($10): Jika kesalahan lolos dan baru ditemukan saat fase code review atau QA testing, dibutuhkan Rp 10 untuk mendokumentasikan tiket bug, mengisolasi fungsi bermasalah, dan memverifikasi perbaikan.")
    add_bullet_p(doc, "Fase Kegagalan Eksternal ($100+): Jika bug tersebut lolos ke server production—misalnya model ML PRE-02 salah memprediksi keterlambatan proyek bernilai ratusan miliar—maka dibutuhkan lebih dari Rp 100 untuk menangani krisis, merilis patch darurat, mengganti kerugian proyek, dan memulihkan reputasi tim pengembang.")

    add_body_p(doc,
        "Logika Sederhana: Mencegah kebakaran dengan memasang kabel standar dan alarm asap selalu jauh lebih murah, aman, "
        "dan rasional daripada membayar gaji regu pemadam kebakaran setiap hari (inspeksi), apalagi menanggung biaya membangun kembali "
        "gedung yang telah hangus terbakar (biaya kegagalan). Kualitas yang dibangun dari awal (built-in quality) menghasilkan produktivitas "
        "yang stabil dan bebas stres manajerial.")

    # =========================================================================
    # SOAL 6
    # =========================================================================
    add_heading_2(doc, "6. Sebutkan dan jelaskan tujuan dari tujuh alat kualitas dasar (Seven Basic Quality Tools)!")
    add_body_p(doc,
        "Tujuh Alat Kualitas Dasar (Seven Basic Quality Tools) adalah kumpulan teknik grafis dan statistik yang dipopulerkan "
        "oleh Dr. Kaoru Ishikawa untuk mengidentifikasi, memetakan, dan memecahkan permasalahan kualitas operasional:")

    headers_q6 = ["No", "Nama Alat Kualitas", "Tujuan dan Fungsi Operasional dalam Proyek"]
    data_q6 = [
        ["1", "Cause-and-Effect Diagram\n(Fishbone / Ishikawa)", "Mengidentifikasi dan memetakan secara sistematis seluruh akar penyebab potensial (root causes) dari suatu permasalahan kualitas tertentu dengan mengelompokkannya ke dalam kategori terstruktur (seperti 5M: Man, Machine, Method, Material, Measurement)."],
        ["2", "Flowchart (Diagram Alir)", "Memvisualisasikan urutan langkah-langkah dalam suatu proses dari awal hingga akhir guna mengidentifikasi titik-titik rawan kegagalan (bottlenecks), redundansi tugas, atau loop perulangan yang tidak efisien."],
        ["3", "Check Sheet (Lembar Periksa)", "Menyediakan formulir terstruktur untuk mencatat, mengumpulkan, dan mengompilasi data frekuensi kejadian atau cacat secara real-time di lapangan secara konsisten dan mudah dianalisis."],
        ["4", "Pareto Diagram (Diagram Pareto)", "Grafik batang gabungan kurva persentase kumulatif yang mengurutkan penyebab masalah dari frekuensi tertinggi ke terendah, berlandaskan Aturan 80/20 untuk memisahkan masalah krusial (vital few) dari masalah sepele (trivial many)."],
        ["5", "Histogram", "Grafik batang yang menampilkan distribusi frekuensi sebaran data numerik kontinu guna melihat kecenderungan pemusatan data (mean/median), variasi rentang, dan pola sebaran (normal atau condong)."],
        ["6", "Control Chart (Bagan Kendali)", "Grafik visual batas toleransi statistik (Upper & Lower Control Limits) dari waktu ke waktu untuk memantau apakah stabilitas proses berada dalam kendali statistik (in statistical control) atau mengalami anomali sistemik."],
        ["7", "Scatter Diagram (Diagram Tebar)", "Grafik pemetaan titik data pada koordinat kartesius (sumbu X dan Y) untuk menganalisis hubungan korelasi, arah tren, dan kekuatan dependensi antara dua variabel berbeda."],
    ]
    create_table(doc, headers_q6, data_q6)

    # =========================================================================
    # SOAL 7
    # =========================================================================
    add_heading_2(doc, "7. Jelaskan perbedaan antara Manage Quality (Quality Assurance) dan Control Quality (Quality Control)!")
    add_body_p(doc,
        "Meskipun sering disamakan dalam percakapan informal, PMBOK Guide 6th Edition membedakan Manage Quality (QA) "
        "dan Control Quality (QC) secara fundamental:")

    headers_q7 = ["Parameter Pembeda", "Manage Quality (Quality Assurance / QA)", "Control Quality (Quality Control / QC)"]
    data_q7 = [
        ["Definisi Inti", "Aktivitas proaktif yang menjamin proses pelaksanaan proyek mematuhi standar kualitas organisasi.", "Aktivitas reaktif yang memverifikasi apakah luaran produk fisik telah memenuhi kriteria teknis spesifik."],
        ["Fokus Utama", "Fokus pada PROSES (pencegahan timbulnya cacat melalui desain alur kerja yang benar).", "Fokus pada PRODUK (pendeteksian dan isolasi cacat fisik pada deliverables yang sudah dibuat)."],
        ["Waktu Pelaksanaan", "Berjalan terus-menerus sepanjang siklus pelaksanaan kerja (proaktif selama eksekusi).", "Dilakukan pada interval akhir fase atau setelah deliverable selesai diproduksi (reaktif inspeksi)."],
        ["Teknik Utama", "Quality Audit, Process Analysis, Root Cause Analysis, Design for X (DfX).", "Testing, Inspection, Sampling Statistik, Pengujian Toleransi Metrik."],
        ["Hasil Luaran (Output)", "Pembaruan proses (Quality Reports), Change Requests perbaikan alur, Lessons Learned.", "Deliverables yang Terverifikasi (Verified Deliverables), Data Pengukuran Kualitas, Cacat yang Ditolak."],
        ["Analogi Sederhana", "Memastikan resep masakan, kebersihan dapur, dan kualifikasi koki higienis.", "Mencicipi sup sebelum disajikan ke meja tamu untuk memastikan rasanya tidak keasinan."],
    ]
    create_table(doc, headers_q7, data_q7)

    # =========================================================================
    # SOAL 8
    # =========================================================================
    add_heading_2(doc, "8. Apa yang dimaksud dengan continuous improvement dan bagaimana konsep ini diterapkan dalam manajemen kualitas proyek?")
    add_body_p(doc,
        "Continuous Improvement (Peningkatan Berkesinambungan / Kaizen) adalah filosofi manajemen kualitas modern "
        "yang meyakini bahwa setiap proses kerja selalu dapat disempurnakan secara bertahap dan konsisten sepanjang waktu. "
        "Tujuannya adalah mengeliminasi pemborosan (waste), mereduksi variasi kesalahan proses, serta meningkatkan efisiensi dan kepuasan pemangku kepentingan.")
    
    add_body_p(doc, "Implementasi Konsep Continuous Improvement dalam Manajemen Proyek:")
    add_num_p(doc, "Penerapan Siklus Deming (PDCA Cycle): Tim proyek merencanakan standar kualitas (Plan), mengeksekusi dan mengukur kinerja (Do), meninjau hasil evaluasi terhadap baseline (Check), lalu merevisi SOP kerja atau parameter teknis jika ditemukan inefisiensi (Act).", 1)
    add_num_p(doc, "Pelaksanaan Retrospektif Berkala (Agile/Lessons Learned): Di setiap akhir sprint atau milestone proyek, tim mengidentifikasi hambatan teknis yang terjadi dan menyepakati minimal satu tindakan perbaikan konkret untuk diterapkan pada iterasi berikutnya.", 2)
    add_num_p(doc, "Integrasi Pembaruan Aset Proses Organisasi (OPAs): Pelajaran yang dipetik (lessons learned) dari penanganan kendala kualitas tidak hanya disimpan secara pasif, tetapi dibukukan ke dalam repositori pengetahuan organisasi agar dapat diadopsi oleh proyek-proyek masa depan.", 3)
    add_num_p(doc, "Penerapan Kontekstual pada Proyek PRE-02: Dalam riset model prediksi keterlambatan, continuous improvement diwujudkan melalui eksperimen iteratif kalibrasi probabilitas—mulai dari Baseline Logistic Regression, evaluasi Expected Calibration Error (ECE), lalu mengoptimalkan pembobotan ensemble gradient boosting dan regularisasi dropout pada MLP guna menekan angka Brier Score secara progresif.", 4)
