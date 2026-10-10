# -*- coding: utf-8 -*-
"""
Section C Builder: Refleksi Mandiri
Pertemuan 12: Manajemen Risiko Proyek (Bagian II)
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_c(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']

    add_heading_1(doc, "C. REFLEKSI MANDIRI")

    add_body_p(doc,
        "Bagian ini menyajikan perenungan kritis dan evaluasi personal mahasiswa terhadap konsep analisis risiko kuantitatif "
        "dan perumusan strategi respons risiko, serta refleksinya dalam mengawal proyek riset PRE-02 (Sistem Prediksi "
        "Probabilitas Keterlambatan Proyek Berbasis Machine Learning sebagai Early Warning System):")

    # =========================================================================
    # REFLEKSI 1
    # =========================================================================
    add_heading_2(doc, "1. Dari berbagai strategi respons risiko yang dipelajari (Avoid, Transfer, Mitigate, Accept), strategi mana yang menurut Anda paling sesuai untuk proyek kelompok Anda saat ini, dan mengapa?")
    add_body_p(doc,
        "Berdasarkan karakteristik riset rekayasa perangkat lunak dan machine learning PRE-02, strategi respons risiko "
        "yang paling dominan dan sesuai untuk diterapkan saat ini adalah kombinasi STRATEGI MITIGATE (MEMITIGASI) "
        "untuk ancaman teknis dan STRATEGI EXPLOIT (MENGEKSPLOITASI) untuk peluang komputasi.")
    
    add_body_p(doc,
        "Rasionalitas Pemilihan Strategi Mitigate:")
    add_bullet_p(doc, "Sifat Inheren Kompleksitas Model ML: Risiko teknis seperti overfitting pada data sintetis, penurunan skor kalibrasi Brier Score, atau tingginya latensi komputasi inferensi tidak dapat dihindari sepenuhnya (Avoid) karena pemodelan probabilistik dan algoritma ensemble adalah esensi utama penelitian PRE-02. Risiko tersebut juga tidak dapat dialihkan (Transfer) kepada pihak ketiga karena proyek dikerjakan secara mandiri dalam lingkungan akademik.")
    add_bullet_p(doc, "Tindakan Reduksi Terarah: Strategi mitigasi memungkinkan tim mengambil langkah rekayasa proaktif yang terukur—seperti menerapkan regularisasi bobot jaringan, validasi silang bertingkat (Stratified K-Fold), kuantisasi model ONNX runtime, dan teknik kalibrasi Isotonic Regression—sehingga skor risiko dapat ditekan ke dalam batas toleransi yang aman.")

    # =========================================================================
    # REFLEKSI 2
    # =========================================================================
    add_heading_2(doc, "2. Menurut Anda, apa perbedaan paling mendasar antara analisis risiko kualitatif dan kuantitatif, dan dalam situasi apa Anda akan memilih untuk melakukan analisis kuantitatif?")
    add_body_p(doc,
        "Perbedaan paling mendasar antara kedua pendekatan tersebut terletak pada sifat representasi data dan tingkat objektivitasnya. "
        "Analisis kualitatif berfokus pada kategorisasi deskriptif berbasis persepsi subjektif ahli (menghasilkan peringkat ordinal: "
        "Tinggi, Sedang, Rendah pada matriks P&I), sedangkan analisis kuantitatif berfokus pada kuantifikasi numerik berbasis matematika statistik "
        "(menghasilkan nilai moneter eksak dalam Rupiah, distribusi waktu probabilitas dalam hari, dan nilai EMV terukur).")

    add_body_p(doc,
        "Situasi di Mana Saya Memilih Melakukan Analisis Kuantitatif:")
    add_num_p(doc, "Ketika Diperlukan Penetapan Anggaran Cadangan Kontingensi yang Akurat: Jika sponsor atau manajemen menuntut justifikasi angka pasti untuk alokasi dana cadangan kontingensi biaya dan buffer jadwal, analisis kuantitatif (seperti EMV dan simulasi Monte Carlo) mutlak diperlukan guna memberikan dasar matematika yang dapat dipertanggungjawabkan.", 1)
    add_num_p(doc, "Pada Titik Keputusan Strategis yang Kompleks (Trade-Off Analysis): Ketika tim dihadapkan pada pilihan arsitektur biner atau keputusan pemilihan vendor pihak ketiga (misalnya analisis Make-or-Buy), teknik kuantitatif Decision Tree Analysis memberikan perbandingan moneter yang sangat transparan.", 2)
    add_num_p(doc, "Ketersediaan Data Historis yang Memadai: Analisis kuantitatif hanya efektif jika proyek memiliki data empiris atau estimasi probabilitas yang cukup andal; jika data sangat minim, analisis kualitatif lebih bijak digunakan untuk menghindari kesalahan ilusi presisi semu.", 3)

    # =========================================================================
    # REFLEKSI 3
    # =========================================================================
    add_heading_2(doc, "3. Berdasarkan proyek kelompok Anda, apakah ada risiko yang menurut Anda perlu dianalisis lebih lanjut secara kuantitatif? Jelaskan mengapa!")
    add_body_p(doc,
        "Ya, terdapat satu risiko kritis pada proyek PRE-02 yang sangat mendesak untuk dianalisis lebih lanjut secara kuantitatif, yaitu: "
        "\"Risiko Kegagalan SLA Latensi Inferensi API Akibat Kompleksitas Arsitektur Ensemble Model ML (Deep Learning MLP vs Gradient Boosting)\".")

    add_body_p(doc,
        "Alasan Mengapa Perlu Dianalisis Kuantitatif:")
    add_bullet_p(doc, "Dampak Terhadap Kebutuhan Biaya Komputasi Cloud: Jika model ensemble yang dilatih memiliki latensi inferensi di atas batas toleransi 200 ms, sistem membutuhkan infrastruktur server komputasi GPU/vCPU berkapasitas tinggi di cloud. Menggunakan analisis Decision Tree dan EMV, tim dapat memodelkan biaya: apakah lebih hemat mengalokasikan waktu sprint 1 minggu untuk melakukan optimasi kompresi model (kuantisasi int8 dan pruning) dengan biaya tenaga kerja tertentu, ataukah langsung menyewa instance server cloud berdaya tinggi dengan biaya langganan moneter bulanan.")
    add_bullet_p(doc, "Keputusan Trade-Off Akurasi vs Efisiensi: Analisis kuantitatif memungkinkan tim menghitung Expected Value dari penurunan marjinal skor AUC-ROC (misal turun 0.015) dibandingkan penghematan waktu inferensi sebesar 60%, sehingga keputusan arsitektur final didasarkan pada data kuantitatif yang optimal bagi sistem peringatan dini.")

    # =========================================================================
    # REFLEKSI 4
    # =========================================================================
    add_heading_2(doc, "4. Menurut Anda, apa tantangan terbesar dalam mengimplementasikan strategi respons risiko dalam proyek nyata?")
    add_body_p(doc,
        "Berdasarkan pengamatan dan pengalaman selama pengerjaan proyek, tantangan terbesar dalam mengimplementasikan strategi respons risiko "
        "bukanlah pada tahap merencanakannya di atas kertas dokumen, melainkan pada tahap eksekusinya di dunia nyata (Implement Risk Responses). "
        "Tantangan-tantangan utama tersebut meliputi:")

    add_num_p(doc, "Keengganan Mengalokasikan Sumber Daya Sebelum Krisis Terjadi (The Prevention Paradox): Sering kali tim merasa enggan menyisihkan waktu dan kapasitas tenaga kerja saat ini untuk aktivitas mitigasi preventif (seperti menulis unit testing otomatis atau menyiapkan script backup) karena merasa risiko tersebut \"belum tentu terjadi\". Akibatnya, rencana mitigasi tertunda hingga risiko benar-benar meledak menjadi isu kritis.", 1)
    add_num_p(doc, "Ambiguitas Kepemilikan Risiko (Diffusion of Responsibility): Ketika risiko tidak ditetapkan memiliki satu orang penanggung jawab tunggal (Risk Owner), tidak ada anggota tim yang merasa bertanggung jawab memantau indikator pemicu (risk triggers). Rencana respons akhirnya hanya menjadi dokumen formalitas tanpa eksekutor aktif.", 2)
    add_num_p(doc, "Munculnya Risiko Sekunder (Secondary Risks): Tindakan mitigasi yang dieksekusi sering kali secara tidak sengaja memicu risiko baru. Misalnya, ketika tim melakukan kuantisasi model untuk memitigasi latensi, timbul risiko sekunder berupa penurunan akurasi probabilitas kalibrasi model yang menuntut pengujian regresi tambahan.", 3)

    # =========================================================================
    # REFLEKSI 5
    # =========================================================================
    add_heading_2(doc, "5. Bagaimana Anda akan memantau efektivitas strategi respons risiko yang telah direncanakan untuk proyek kelompok Anda?")
    add_body_p(doc,
        "Untuk memastikan bahwa rencana respons risiko yang telah disepakati benar-benar berjalan efektif dan memberikan perlindungan "
        "nyata terhadap sasaran riset PRE-02, saya akan menerapkan empat mekanisme pemantauan berkelanjutan (Monitor Risks):")

    add_num_p(doc, "Audit Metrik Kinerja Teknis Berkala (Technical Performance Analysis): Memantau secara berkala metrik teknis kunci model—seperti Brier Score, ECE (Expected Calibration Error), kurva AUC-ROC, dan latensi inferensi API—setiap kali kode baru digabungkan ke repositori. Jika metrik menyimpang dari ambang batas toleransi, hal ini menjadi sinyal objektif bahwa tindakan mitigasi perlu disesuaikan.", 1)
    add_num_p(doc, "Peninjauan Status Pemicu Risiko pada Setiap Sprint Standup: Menjadikan pengecekan indikator pemicu (risk triggers) sebagai rutinitas pemantauan tim. Setiap pemilik risiko melaporkan apakah gejala awal ancaman mulai teramati sehingga rencana kontingensi dapat diaktifkan tepat waktu.", 2)
    add_num_p(doc, "Analisis Cadangan Kontingensi (Reserve Analysis): Membandingkan sisa cadangan waktu (buffer jadwal) dan kapasitas kerja tim terhadap profil risiko yang tersisa. Jika risiko-risiko fase awal telah berhasil diatasi, cadangan penyangga dapat dilepaskan untuk memperkuat fase integrasi antarmuka.", 3)
    add_num_p(doc, "Pembaruan Berkala Dokumen Risk Register (Continuous Retrospective): Menjadikan Risk Register sebagai dokumen hidup yang diperbarui setiap akhir sprint. Risiko yang telah tertangani ditutup, risiko baru dicatat, dan efektivitas strategi respons sebelumnya dievaluasi secara terbuka dalam sesi retrospektif tim.", 4)
