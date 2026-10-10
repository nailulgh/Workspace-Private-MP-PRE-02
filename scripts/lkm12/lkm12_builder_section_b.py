# -*- coding: utf-8 -*-
"""
Section B Builder: Tugas Latihan Individu
Pertemuan 12: Perhitungan Expected Monetary Value (EMV) dan Decision Tree Analysis
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_b(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "B. TUGAS LATIHAN INDIVIDU: ANALISIS EXPECTED MONETARY VALUE (EMV) DAN DECISION TREE")

    # =========================================================================
    # BAGIAN 1: PERHITUNGAN EMV TIGA RISIKO PROYEK
    # =========================================================================
    add_heading_2(doc, "1. Latihan Perhitungan Expected Monetary Value (EMV) Tiga Risiko Proyek")
    add_body_p(doc,
        "Diketahui sebuah proyek pengembangan teknologi informasi memiliki tiga peristiwa risiko yang telah diidentifikasi "
        "dan dinilai secara kuantitatif dengan rincian data sebagai berikut (sesuai data silabus acuan):")

    headers_emv_raw = ["Identifikasi Risiko", "Probabilitas Keterjadian (P)", "Dampak Moneter Potensial (I)"]
    data_emv_raw = [
        ["Risiko A", "30% (0,30)", "Rp 100.000.000"],
        ["Risiko B", "50% (0,50)", "Rp 50.000.000"],
        ["Risiko C", "20% (0,20)", "Rp 200.000.000"]
    ]
    create_table(doc, headers_emv_raw, data_emv_raw)

    # 1.1 Perhitungan EMV Masing-Masing Risiko (Asumsi Awal: Ketiganya adalah Ancaman)
    add_heading_3(doc, "a. Perhitungan Nilai EMV untuk Masing-Masing Risiko (Skenario Seluruhnya Ancaman)")
    add_body_p(doc,
        "Dalam skenario dasar di mana ketiga peristiwa risiko di atas diklasifikasikan sebagai risiko negatif (ancaman/threats), "
        "dampak moneter bernilai negatif (-) karena memicu pengeluaran atau kerugian biaya proyek. Formulasi EMV = P × (-I):")
    
    add_bullet_p(doc, "EMV Risiko A: P(A) × Impact(A) = 0,30 × (-Rp 100.000.000) = -Rp 30.000.000")
    add_bullet_p(doc, "EMV Risiko B: P(B) × Impact(B) = 0,50 × (-Rp 50.000.000) = -Rp 25.000.000")
    add_bullet_p(doc, "EMV Risiko C: P(C) × Impact(C) = 0,20 × (-Rp 200.000.000) = -Rp 40.000.000")

    # 1.2 Total EMV
    add_heading_3(doc, "b. Perhitungan Total Nilai EMV Proyek")
    add_body_p(doc,
        "Total EMV dihitung dengan menjumlahkan aljabar nilai EMV dari seluruh risiko individu:")
    add_bullet_p(doc, "Total EMV = EMV(A) + EMV(B) + EMV(C)")
    add_bullet_p(doc, "Total EMV = (-Rp 30.000.000) + (-Rp 25.000.000) + (-Rp 40.000.000)")
    add_bullet_p(doc, "Total EMV = -Rp 95.000.000 (Kerugian Moneter yang Diharapkan)")

    # 1.3 Interpretasi Total EMV
    add_heading_3(doc, "c. Interpretasi Manajerial dari Total EMV (-Rp 95.000.000)")
    add_body_p(doc,
        "Secara statistik, nilai Total EMV sebesar -Rp 95.000.000 merepresentasikan nilai rata-rata kerugian moneter terbobot probabilitas "
        "yang dihadapi oleh portofolio proyek akibat paparan ketiga ancaman tersebut. Nilai ini tidak berarti bahwa proyek pasti merugi "
        "Rp 95.000.000 pada realitasnya, melainkan berfungsi sebagai estimasi matematis objektif untuk menetapkan besaran Cadangan Kontingensi Biaya "
        "(Cost Contingency Reserve). Manajer Proyek wajib mencadangkan anggaran penyangga sebesar Rp 95.000.000 di dalam baseline anggaran proyek "
        "guna menyerap dampak finansial yang mungkin terjadi sepanjang siklus hidup proyek.")

    # 1.4 Skenario Modifikasi: Risiko A adalah Peluang, B dan C adalah Ancaman
    add_heading_3(doc, "d. Analisis Skenario Diferensiasi: Risiko A adalah Peluang (+), Risiko B & C adalah Ancaman (-)")
    add_body_p(doc,
        "Jika Risiko A diklasifikasikan sebagai risiko positif atau peluang (opportunity)—misalnya penghematan biaya lisensi perangkat lunak "
        "karena diskon institusi—sedangkan Risiko B dan Risiko C tetap sebagai ancaman, maka perhitungan aljabar EMV berubah menjadi:")
    
    add_bullet_p(doc, "EMV Risiko A (Peluang / Positif): P(A) × (+Impact A) = 0,30 × (+Rp 100.000.000) = +Rp 30.000.000")
    add_bullet_p(doc, "EMV Risiko B (Ancaman / Negatif): P(B) × (-Impact B) = 0,50 × (-Rp 50.000.000) = -Rp 25.000.000")
    add_bullet_p(doc, "EMV Risiko C (Ancaman / Negatif): P(C) × (-Impact C) = 0,20 × (-Rp 200.000.000) = -Rp 40.000.000")
    add_bullet_p(doc, "Total Net EMV Proyek Baru: (+Rp 30.000.000) + (-Rp 25.000.000) + (-Rp 40.000.000) = -Rp 35.000.000")

    add_body_p(doc,
        "Interpretasi Perbedaan Skenario:\n"
        "Keberadaan peluang A memberikan dampak kompensasi positif (+Rp 30.000.000) yang meng-offset sebagian dari eksposur kerugian ancaman B dan C (-Rp 65.000.000). "
        "Akibatnya, eksposur kerugian bersih (net risk exposure) proyek menyusut secara drastis dari -Rp 95.000.000 menjadi -Rp 35.000.000. "
        "Implikasi manajerialnya adalah Manajer Proyek hanya perlu menyiapkan cadangan kontingensi biaya bersih sebesar Rp 35.000.000, "
        "dengan asumsi tim mengeksekusi strategi respons aktif (misal: Exploit atau Enhance) untuk menjamin peluang A terwujud.")

    # 1.5 Prioritas Tertinggi Berdasarkan EMV
    add_heading_3(doc, "e. Penentuan Prioritas Risiko Tertinggi Berdasarkan Analisis Nilai EMV")
    add_body_p(doc,
        "Berdasarkan besaran nilai mutlak EMV (|EMV|), urutan peringkat risiko adalah sebagai berikut:\n"
        "1. Risiko C: |EMV| = Rp 40.000.000 (Peringkat 1 / Prioritas Tertinggi)\n"
        "2. Risiko A: |EMV| = Rp 30.000.000 (Peringkat 2 / Prioritas Sedang)\n"
        "3. Risiko B: |EMV| = Rp 25.000.000 (Peringkat 3 / Prioritas Terendah)")
    
    add_body_p(doc,
        "Justifikasi dan Rasionalitas:\n"
        "Risiko C merupakan prioritas tertinggi yang wajib mendapatkan alokasi sumber daya mitigasi agresif terlebih dahulu. "
        "Meskipun probabilitas terjadinya paling kecil di antara ketiganya (hanya 20%), dampak moneter yang ditimbulkan sangat masif "
        "(Rp 200.000.000), sehingga secara statistik menghasilkan bobot kerugian ekspektasi terbesar (-Rp 40.000.000). "
        "Sebaliknya, Risiko B yang memiliki probabilitas tertinggi (50%) justru menghasilkan eksposur EMV paling rendah (-Rp 25.000.000) "
        "karena dampaknya relatif kecil. Analisis EMV membuktikan bahwa mengelola risiko tidak boleh hanya terpaku pada probabilitas semata, "
        "melainkan wajib mempertimbangkan besaran kuantitatif paparan moneternya.")

    headers_emv_summary = ["Peristiwa Risiko", "Tipe", "Probabilitas (P)", "Dampak (I)", "Nilai EMV (Rp)", "Peringkat Prioritas"]
    data_emv_summary = [
        ["Risiko C", "Ancaman", "20%", "-Rp 200.000.000", "-Rp 40.000.000", "Prioritas 1 (Kritis)"],
        ["Risiko A", "Ancaman / Peluang", "30%", "±Rp 100.000.000", "±Rp 30.000.000", "Prioritas 2 (Moderat)"],
        ["Risiko B", "Ancaman", "50%", "-Rp 50.000.000", "-Rp 25.000.000", "Prioritas 3 (Rendah)"]
    ]
    create_table(doc, headers_emv_summary, data_emv_summary)

    # =========================================================================
    # BAGIAN 2: DECISION TREE ANALYSIS (VENDOR X VS VENDOR Y)
    # =========================================================================
    add_heading_2(doc, "2. Tugas Latihan Tambahan: Decision Tree Analysis (Pemilihan Alternatif Vendor)")
    add_body_p(doc,
        "Kasus Studi: Seorang Manajer Proyek dihadapkan pada keputusan strategis untuk memilih salah satu dari dua vendor "
        "guna mengerjakan aktivitas kritis proyek. Data perbandingan penawaran biaya dan risiko denda keterlambatan disajikan berikut:")
    
    add_bullet_p(doc, "Vendor X: Biaya penawaran dasar = Rp 150.000.000. Probabilitas terlambat = 30% (0,30) dengan denda keterlambatan = Rp 50.000.000. Probabilitas tepat waktu = 70% (0,70) tanpa denda.")
    add_bullet_p(doc, "Vendor Y: Biaya penawaran dasar = Rp 180.000.000. Probabilitas terlambat = 10% (0,10) dengan denda keterlambatan = Rp 80.000.000. Probabilitas tepat waktu = 90% (0,90) tanpa denda.")

    # 2.1 Struktur Decision Tree
    add_heading_3(doc, "a. Struktur Diagram dan Pemodelan Decision Tree")
    add_body_p(doc,
        "Model pohon keputusan (Decision Tree) terdiri dari Simpul Keputusan (Decision Node □), "
        "Simpul Ketidakpastian/Peluang (Chance Nodes ○), dan Cabang Hasil Akhir (Payoff Branches):")

    headers_dt = ["Simpul Keputusan (□)", "Alternatif Cabang", "Simpul Peluang (○)", "Kondisi Kinerja Vendor", "Probabilitas", "Biaya Dasar + Denda", "Total Pengeluaran (Payoff)"]
    data_dt = [
        ["Keputusan Memilih Vendor", "Cabang 1: Vendor X", "Simpul Peluang X", "Tepat Waktu\nTerlambat (Denda Rp 50 Jt)", "0,70\n0,30", "Rp 150 Jt + Rp 0\nRp 150 Jt + Rp 50 Jt", "Rp 150.000.000\nRp 200.000.000"],
        ["Keputusan Memilih Vendor", "Cabang 2: Vendor Y", "Simpul Peluang Y", "Tepat Waktu\nTerlambat (Denda Rp 80 Jt)", "0,90\n0,10", "Rp 180 Jt + Rp 0\nRp 180 Jt + Rp 80 Jt", "Rp 180.000.000\nRp 260.000.000"]
    ]
    create_table(doc, headers_dt, data_dt)

    # 2.2 Perhitungan EMV Masing-Masing Vendor
    add_heading_3(doc, "b. Perhitungan Nilai EMV untuk Masing-Masing Alternatif Vendor")
    add_body_p(doc,
        "Perhitungan Expected Monetary Value dilakukan dengan mengalikan probabilitas masing-masing cabang hasil dengan total biaya pengeluaran:")
    
    add_body_p(doc, "1. Perhitungan EMV Pengeluaran Total untuk Vendor X:")
    add_bullet_p(doc, "Kondisi Tepat Waktu (Probabilitas 0,70): Biaya = Rp 150.000.000")
    add_bullet_p(doc, "Kondisi Terlambat (Probabilitas 0,30): Biaya = Rp 150.000.000 + Rp 50.000.000 = Rp 200.000.000")
    add_bullet_p(doc, "EMV(Vendor X) = (0,70 × Rp 150.000.000) + (0,30 × Rp 200.000.000)")
    add_bullet_p(doc, "EMV(Vendor X) = Rp 105.000.000 + Rp 60.000.000 = Rp 165.000.000")
    add_body_p(doc, "(Catatan: Dapat juga dihitung dari Biaya Pokok Rp 150 Jt + Ekspektasi Denda (0,30 × Rp 50 Jt = Rp 15 Jt) = Rp 165.000.000).")

    add_body_p(doc, "2. Perhitungan EMV Pengeluaran Total untuk Vendor Y:")
    add_bullet_p(doc, "Kondisi Tepat Waktu (Probabilitas 0,90): Biaya = Rp 180.000.000")
    add_bullet_p(doc, "Kondisi Terlambat (Probabilitas 0,10): Biaya = Rp 180.000.000 + Rp 80.000.000 = Rp 260.000.000")
    add_bullet_p(doc, "EMV(Vendor Y) = (0,90 × Rp 180.000.000) + (0,10 × Rp 260.000.000)")
    add_bullet_p(doc, "EMV(Vendor Y) = Rp 162.000.000 + Rp 26.000.000 = Rp 188.000.000")
    add_body_p(doc, "(Catatan: Dapat juga dihitung dari Biaya Pokok Rp 180 Jt + Ekspektasi Denda (0,10 × Rp 80 Jt = Rp 8 Jt) = Rp 188.000.000).")

    # 2.3 Keputusan Berdasarkan Analisis EMV
    add_heading_3(doc, "c. Keputusan Rekomendasi Berdasarkan Analisis Kuantitatif EMV")
    add_body_p(doc,
        "Rekomendasi Pemilihan: VENDOR X SEBAIKNYA DIPILIH.\n"
        "Alasan Analitis:\n"
        "Dalam konteks manajemen biaya proyek, kriteria seleksi yang optimal adalah meminimalkan total ekspektasi pengeluaran dana moneter (expected cost minimization). "
        "Nilai EMV pengeluaran Vendor X adalah sebesar Rp 165.000.000, sedangkan EMV pengeluaran Vendor Y adalah sebesar Rp 188.000.000. "
        "Dengan memilih Vendor X, proyek secara statistik diperkirakan menghemat anggaran sebesar Rp 23.000.000 "
        "(Rp 188.000.000 - Rp 165.000.000) dibandingkan jika mengontrak Vendor Y.")

    # 2.4 Analisis dari Perspektif Risk-Averse
    add_heading_3(doc, "d. Analisis Keputusan Berdasarkan Sikap Preferensi Risiko (Risk-Averse Perspective)")
    add_body_p(doc,
        "Pertanyaan: Apakah pilihan akan berbeda jika Manajer Proyek memiliki preferensi menghindari risiko (risk-averse)?\n"
        "Jawaban: YA, PILIHAN DAPAT BERUBAH MEMILIH VENDOR Y.\n"
        "Penjelasan Analitis Mendalam:")
    
    add_num_p(doc, "Karakteristik Pengambil Keputusan Risk-Averse: Seseorang yang risk-averse sangat memprioritaskan kepastian dan memiliki toleransi yang sangat rendah terhadap probabilitas kegagalan, terutama jika aktivitas tersebut bersifat sangat kritis (Critical Activity) terhadap reputasi organisasi atau jadwal rilis akhir proyek.", 1)
    add_num_p(doc, "Tingkat Keandalan Vendor Y (90% Kepastian Sukses): Vendor Y menawarkan tingkat kepastian penyelesaian tepat waktu sebesar 90% (hanya 10% peluang terlambat), sementara Vendor X memiliki risiko kegagalan keterlambatan yang cukup tinggi, yaitu 30% (hampir sepertiga kemungkinan gagal).", 2)
    add_num_p(doc, "Konsep Premi Asuransi Kualitas (Risk Premium): Selisih biaya dasar penawaran sebesar Rp 30.000.000 (Vendor Y Rp 180 Jt vs Vendor X Rp 150 Jt) dipandang oleh manajer risk-averse sebagai \"premi asuransi\" yang wajar untuk dibayarkan demi membeli kepastian ketepatan waktu sebesar 90%. Manajer risk-averse lebih rela membayar pasti lebih mahal di depan daripada menanggung stres operasional dan ancaman reputasi akibat 30% risiko kegagalan Vendor X.", 3)
    add_num_p(doc, "Konteks Kritis Proyek PRE-02: Jika modul yang dialihdayakan adalah modul inti pengolahan pipeline fitur Machine Learning yang menjadi prasyarat mutlak seluruh pengujian berikutnya, memilih Vendor Y menjadi keputusan yang sangat rasional bagi manajer yang risk-averse guna menjamin kesinambungan proyek tanpa gangguan jadwal kritis.", 4)
