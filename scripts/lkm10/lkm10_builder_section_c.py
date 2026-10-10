# -*- coding: utf-8 -*-
"""
Section C Builder: Refleksi Mandiri
Pertemuan 10: Manajemen Komunikasi Proyek
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_c(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']

    add_heading_1(doc, "C. REFLEKSI MANDIRI")

    add_body_p(doc,
        "Bagian ini memuat introspeksi kritis dan evaluasi personal mahasiswa terhadap pemahaman konseptual manajemen komunikasi "
        "proyek serta penerapannya secara nyata dalam dinamika tim proyek riset kelompok PRE-02 (Sistem Prediksi Probabilitas Keterlambatan "
        "Proyek Berbasis Machine Learning sebagai Early Warning System):")

    # =========================================================================
    # REFLEKSI 1
    # =========================================================================
    add_heading_2(doc, "1. Berdasarkan pengalaman Anda dalam proyek kelompok, kesalahan komunikasi apa yang paling sering terjadi dan mengapa kesalahan tersebut terjadi?")
    add_body_p(doc,
        "Berdasarkan pengalaman empiris selama pengerjaan proyek riset PRE-02, kesalahan komunikasi yang paling sering muncul adalah "
        "terjadinya asumsi implisit tanpa verifikasi tertulis (unverified implicit assumptions) mengenai spesifikasi antarmuka teknis "
        "dan format data model Machine Learning.")
    
    add_body_p(doc,
        "Akar Penyebab Kesalahan Tersebut:")
    add_bullet_p(doc, "Ambiguitas Format Data Output: Anggota tim pemodelan data science mengasumsikan bahwa hasil prediksi model cukup disediakan berupa nilai probabilitas float kontinu [0.0 - 1.0] yang belum terkalibrasi melalui fungsi aktivasi sigmoid/predict_proba. Di sisi lain, anggota tim rekayasa antarmuka (UI/Backend) mengasumsikan bahwa model akan langsung mengembalikan label biner ('Terlambat' vs 'Tepat Waktu') beserta kategori warna status bahaya. Akibat tiadanya kontrak data formal (data contract) di awal, modul pipeline sempat mengalami kegagalan integrasi (data type mismatch).")
    add_bullet_p(doc, "Ketergantungan Berlebih pada Komunikasi Lisan Spontan: Diskusi teknis penting sering kali dilakukan melalui obrolan pesan instan yang cepat tenggelam atau percakapan lisan tanpa adanya notula tertulis (minutes of meeting) atau pembaruan dokumentasi pada repositori sentral, sehingga memicu perbedaan memori ingatan antar-anggota.")
    add_bullet_p(doc, "Kutukan Pengetahuan (Curse of Knowledge): Masing-masing anggota mengira bahwa rekan tim lainnya telah memahami alur pikir teknisnya secara otomatis, padahal pemahaman tersebut belum pernah dieksplisitkan secara sistematis.")

    # =========================================================================
    # REFLEKSI 2
    # =========================================================================
    add_heading_2(doc, "2. Menurut Anda, bagaimana cara paling efektif untuk meningkatkan komunikasi dalam tim proyek, terutama ketika anggota tim memiliki gaya komunikasi dan latar belakang yang berbeda?")
    add_body_p(doc,
        "Dalam tim proyek teknologi informasi multidisiplin seperti tim PRE-02—yang memadukan kepribadian analitis-probabilistik "
        "(Data Scientist), kepribadian terstruktur-deterministik (Software Architect), dan kepribadian berorientasi hasil-tenggat waktu (Project Manager)—"
        "strategi paling efektif untuk meningkatkan komunikasi berpusat pada tiga pilar fundamental:")

    add_num_p(doc, "Membangun Shared Mental Model Melalui Kamus Data Bersama (Glosarium Proyek): Menghilangkan disparitas persepsi dengan menyusun kamus definisi istilah resmi di awal proyek. Istilah seperti \"kalibrasi probabilitas\", \"Early Warning Threshold\", \"Dependency Score\", dan \"Brier Score\" didefinisikan secara baku dalam satu dokumen acuan yang disepakati bersama, sehingga tidak ada lagi salah tafsir antar-disiplin ilmu.", 1)
    add_num_p(doc, "Membudayakan Keselamatan Psikologis (Psychological Safety) yang Inklusif: Menjamin bahwa anggota tim yang berkepribadian introvert atau lebih berhati-hati merasa aman untuk menyampaikan keraguan, menanyakan konsep yang belum dipahami, atau mengoreksi asumsi anggota lain tanpa rasa takut direndahkan. Pemimpin tim harus secara aktif memfasilitasi giliran berbicara (equal conversational turn-taking) dalam setiap sesi rapat.", 2)
    add_num_p(doc, "Mengadopsi Komunikasi Visual dan Berbasis Bukti Nyata: Mengurangi perdebatan retoris abstrak dengan menggunakan artefak visual nyata, seperti diagram alur pipeline data, skema arsitektur API Swagger, atau grafik distribusi probabilitas. Menampilkan artefak visual memungkinkan seluruh anggota tim—terlepas dari perbedaan gaya komunikasi verbalnya—untuk melihat realitas teknis yang sama persis.", 3)

    # =========================================================================
    # REFLEKSI 3
    # =========================================================================
    add_heading_2(doc, "3. Dari tiga metode komunikasi (interaktif, push, pull), metode mana yang menurut Anda paling sering digunakan dalam proyek kelompok Anda dan metode mana yang masih perlu ditingkatkan?")
    add_body_p(doc,
        "Evaluasi kritis terhadap penerapan tiga metode komunikasi dalam dinamika internal proyek PRE-02 menghasilkan temuan berikut:")

    add_bullet_p(doc, "Metode yang Paling Dominan Digunakan: Kombinasi Push Communication (melalui grup percakapan WhatsApp untuk pengumuman jadwal dan tautan dokumen) dan Interactive Communication (diskusi tatap muka di kampus dan koordinasi audio-video daring). Metode ini sangat efektif untuk interaksi cepat dan pengambilan keputusan mendesak.")
    add_bullet_p(doc, "Metode yang Sangat Mendesak untuk Ditingkatkan: Pull Communication (Komunikasi Tarik/Mandiri).")
    
    add_body_p(doc,
        "Analisis Kebutuhan Peningkatan Pull Communication:\n"
        "Saat ini, tim masih menghadapi tantangan redundansi koordinasi—di mana anggota tim sering kali mengajukan pertanyaan berulang di grup chat "
        "mengenai struktur dataset terbaru, bobot fitur, atau progres eksperimen model ML. Hal ini terjadi karena arsip dokumentasi teknis, "
        "log pelatihan model (MLflow/Weights & Biases), dan repositori GitHub belum dioptimalkan sebagai sistem Single Source of Truth yang mudah diakses. "
        "Dengan meningkatkan metode Pull Communication (menata rapi wiki GitHub, README, dan dashboard metrik proyek), setiap anggota tim dapat menarik "
        "informasi secara mandiri kapan pun dibutuhkan, mereduksi distraksi pesan instan, dan meningkatkan efisiensi waktu kerja secara drastis.")

    # =========================================================================
    # REFLEKSI 4
    # =========================================================================
    add_heading_2(doc, "4. Bagaimana Anda akan menerapkan lima langkah komunikasi efektif (Listen Intently, Think Clearly, Discuss Openly, Develop Sensitivity, Respond Quickly to Needs) dalam peran Anda di proyek kelompok?")
    add_body_p(doc,
        "Dalam kapasitas saya sebagai anggota tim pengembang inti riset PRE-02, saya berkomitmen menerapkan lima langkah komunikasi efektif PMI "
        "secara operasional dan disiplin melalui tindakan nyata berikut:")

    add_num_p(doc, "Penerapan Listen Intently: Ketika rekan tim Backend atau Project Manager memaparkan kendala latensi server atau tenggat waktu rilis, saya akan mendengarkan secara utuh tanpa membuka editor kode atau memainkan ponsel. Saya akan mencatat poin-poin utama dan melakukan teknik rephrasing (\"Apakah maksud Anda arsitektur model saat ini terlalu lambat saat melayani request inferensi?\") untuk memastikan keselarasan persepsi.", 1)
    add_num_p(doc, "Penerapan Think Clearly: Sebelum memberikan respons atau membela model Machine Learning yang telah saya latih, saya akan mengambil jeda sejenak untuk menganalisis data empiris. Saya memisahkan antara fakta metrik (misal: waktu inferensi 250 ms vs target SLA 100 ms) dari ego profesional, sehingga argumen yang saya sampaikan berlandaskan fakta logis dan solusi teknis terstruktur.", 2)
    add_num_p(doc, "Penerapan Discuss Openly: Saya akan secara proaktif dan transparan membuka kendala eksperimen model di forum tim—misalnya jika model mengalami overfitting pada data sintetis atau nilai ECE (Expected Calibration Error) masih melebihi batas toleransi. Saya tidak akan menutupi kekurangan data demi pujian semu, melainkan mengajak rekan tim membedahnya bersama.", 3)
    add_num_p(doc, "Penerapan Develop Sensitivity: Saya menyadari bahwa rekan tim memiliki beban akademis paralel dari mata kuliah lain. Oleh karena itu, saya akan menyampaikan masukan kode program dengan nada konstruktif dan empatik (misal: menggunakan format \"Saya menyarankan optimasi di fungsi ini agar lebih efisien\" daripada \"Kode Anda lambat\"), serta memilih waktu komunikasi yang tepat dan tidak mengganggu jam istirahat.", 4)
    add_num_p(doc, "Penerapan Respond Quickly to Needs: Berkomitmen untuk tidak membiarkan pertanyaan rekan tim di repositori GitHub atau saluran komunikasi terkatung-katung. Saya menetapkan standar pribadi untuk merespons pertanyaan pemblokir kerja (blocking issues) maksimal dalam waktu 2 jam pada jam kerja operasional proyek guna menjamin kelancaran sprint tim.", 5)

    # =========================================================================
    # REFLEKSI 5
    # =========================================================================
    add_heading_2(doc, "5. Menurut Anda, seberapa penting kontrak tim (team contract) dalam mengatur komunikasi dan ekspektasi tim, dan apa yang harus dimuat dalam kontrak tim untuk proyek kelompok Anda?")
    add_body_p(doc,
        "Kontrak Tim (Team Charter / Team Contract) memiliki tingkat urgensi yang fundamental dan mutlak diperlukan dalam proyek rekayasa perangkat lunak. "
        "Kontrak tim bertindak sebagai landasan psikososial formal yang menyepakati aturan main (ground rules) sebelum friksi atau konflik timbul. "
        "Tanpa kontrak tim yang jelas, tim akan beroperasi di bawah asumsi individual yang bertabrakan, rentan mengalami sindrom penunggang bebas (free-rider syndrome), "
        "dan terjebak dalam ambiguitas peran.")

    add_body_p(doc,
        "Berdasarkan karakteristik riset PRE-02, klausul-klausul esensial yang wajib dimuat dalam Kontrak Tim kelompok kami meliputi:")
    
    add_bullet_p(doc, "Saluran Komunikasi Resmi dan Protokol Aksesibilitas: Menetapkan WhatsApp sebagai saluran komunikasi darurat/koordinasi cepat, Discord untuk sesi kerja bareng (co-working) virtual, dan GitHub Issues/Discussions untuk seluruh pelaporan bug dan keputusan teknis resmi.")
    add_bullet_p(doc, "Service Level Agreement (SLA) Respons Komunikasi: Menyepakati batas waktu respons pesan: maksimal 2 jam untuk isu kritis yang menghambat pekerjaan rekan lain (blocking issue), dan maksimal 12 jam untuk diskusi umum non-kritis.")
    add_bullet_p(doc, "Protokol Rapat dan Kehadiran (Meeting Etiquette): Menetapkan kewajiban hadir tepat waktu pada sesi Sprint Sync mingguan, mematikan mikrofon saat orang lain berbicara, menyampaikan bahan evaluasi sebelum rapat dimulai, dan kewajiban notula tertulis bergiliran.")
    add_bullet_p(doc, "Mekanisme Pengambilan Keputusan dan Resolusi Perbedaan Teknis: Mengadopsi prinsip voting mayoritas jika konsensus diskusi teknis (misalnya pemilihan pustaka XGBoost vs LightGBM) tidak tercapai dalam waktu 30 menit, dengan hak veto strategis di tangan Project Manager.")
    add_bullet_p(doc, "Komitmen Kualitas Kode dan Akuntabilitas Kontribusi: Kesepakatan bahwa seluruh kode wajib melewati uji coba lokal dan lolos review minimal satu rekan tim (peer pull request review) sebelum digabungkan (merged) ke branch utama.")
    add_bullet_p(doc, "Prosedur Eskalasi Konflik Bertingkat: Jika terjadi ketidaksepakatan atau kelalaian anggota tim yang berlarut-larut, masalah wajib diselesaikan terlebih dahulu secara internal (face-to-face problem solving) sebelum dieskalasikan secara formal kepada Dosen Pengampu.")
