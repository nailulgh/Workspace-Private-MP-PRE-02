# -*- coding: utf-8 -*-
"""
Section C Builder: Refleksi Mandiri
Pertemuan 7: Manajemen Biaya Proyek
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_c(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "C. REFLEKSI MANDIRI")

    add_heading_2(doc, "1. Menurut Anda, apa manfaat terbesar Earned Value Management bagi seorang manajer proyek dalam mengendalikan biaya dan jadwal proyek secara bersamaan?")
    add_body_p(doc, 
        "Menurut saya, manfaat terbesar EVM adalah kemampuannya menyediakan indikator kuantitatif yang objektif tanpa bias subjektif. "
        "Sebagai contoh, tanpa EVM, seorang manajer mungkin merasa aman karena baru menghabiskan 30% anggaran, namun dengan EVM "
        "ia bisa menyadari bahwa progres fisiknya (Earned Value) ternyata baru 15%. Hal ini memungkinkan deteksi dini terhadap "
        "potensi keterlambatan (schedule slip) dan pembengkakan dana (cost overrun) sehingga tindakan korektif bisa dilakukan "
        "sebelum semuanya terlambat.")

    add_heading_2(doc, "2. Menurut Anda, apa kesulitan terbesar dalam mengestimasi biaya proyek, terutama ketika menggunakan teknik bottom-up estimating?")
    add_body_p(doc, 
        "Kesulitan terbesarnya adalah kebutuhan akan detail yang sangat tinggi (granularity). Bottom-up mengharuskan setiap "
        "Work Package sudah terdefinisi secara sempurna di awal (WBS sudah matang). Pada proyek IT atau rekayasa perangkat "
        "lunak seperti pembuatan model ML, mendefinisikan seluruh variabel secara absolut di awal sering kali mustahil "
        "karena tingginya faktor eksplorasi atau \"Unknown-Unknowns\". Akibatnya, estimasi ini sangat memakan waktu, tenaga, "
        "dan rawan lumpuh di fase perencanaan jika informasinya belum memadai.")

    add_heading_2(doc, "3. Dari empat teknik estimasi biaya yang dipelajari, teknik mana yang menurut Anda paling sesuai untuk proyek kelompok Anda, dan mengapa?")
    add_body_p(doc, 
        "Untuk proyek PRE-02, teknik Parametric Estimating yang dikombinasikan dengan Bottom-Up (Hybrid) adalah yang paling "
        "ideal. Karena komponen seperti komputasi GPU (Cloud) dapat dihitung dengan model matematis parametrik (misal: "
        "Rp 30.000 per jam × 50 jam training). Sementara untuk alokasi jam kerja SDM (Data Scientist & ML Engineer) pada setiap "
        "task spesifik (seperti Hyperparameter Tuning), Bottom-Up dapat diaplikasikan agar biaya tenaga ahli bisa dipertanggungjawabkan "
        "secara rinci per Work Package.")

    add_heading_2(doc, "4. Bagaimana Anda akan menggunakan metrik EVM (CPI dan SPI) untuk mengambil keputusan jika proyek menunjukkan CPI < 1 dan SPI > 1 (di atas anggaran tetapi di depan jadwal)?")
    add_body_p(doc, 
        "Kondisi CPI < 1 (Cost Overrun) dan SPI > 1 (Ahead of Schedule) menunjukkan bahwa proyek berjalan sangat cepat, namun "
        "mengorbankan efisiensi dana, misalnya karena sering melakukan kerja lembur (overtime), membayar biaya kirim ekspres (premium), "
        "atau menyewa resource tambahan yang mahal. "
        "Keputusan saya sebagai Manajer Proyek adalah melakukan \"Resource Optimization\". Saya akan mengurangi jam lembur atau "
        "menghentikan penyewaan resource premium tersebut, agar kecepatan progres kembali ke batas normal yang dijadwalkan, "
        "sekaligus menekan biaya aktual (AC) agar CPI bisa berangsur pulih mendekati angka 1.0.")

    add_heading_2(doc, "5. Menurut Anda, mengapa estimasi biaya cenderung menjadi lebih akurat seiring berjalannya proyek, dan bagaimana Anda mengelola ketidakpastian biaya di awal proyek?")
    add_body_p(doc, 
        "Fenomena ini dikenal dengan \"Cone of Uncertainty\". Seiring berjalannya proyek, asumsi-asumsi liar mulai tervalidasi, "
        "risiko-risiko berhasil dimitigasi, dan definisi ruang lingkup menjadi semakin solid. Hal inilah yang menyempitkan variansi "
        "estimasi biaya. Untuk mengelola ketidakpastian di fase inisiasi proyek (saat tingkat ketidakpastian masih sangat lebar), "
        "saya menggunakan strategi alokasi Management Reserve dan Contingency Reserve yang proporsional, serta mengimplementasikan "
        "revisi anggaran berkala ketika milestone tertentu (seperti selesainya desain arsitektur) telah terlewati.")
