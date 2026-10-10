# -*- coding: utf-8 -*-
"""
Section B Builder: Tugas Latihan Individu (Perhitungan EVM)
Pertemuan 7: Manajemen Biaya Proyek
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_b(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "B. TUGAS LATIHAN INDIVIDU: PERHITUNGAN EVM")

    # Deskripsi Kasus
    add_body_p(doc,
        "Studi Kasus: Sebuah proyek memiliki total anggaran (BAC) sebesar Rp 100.000.000. "
        "Pada tanggal pelaporan, rencana proyek menargetkan penyelesaian 50% pekerjaan. "
        "Namun, tim proyek melaporkan bahwa pekerjaan yang faktual telah selesai adalah 40%. "
        "Sedangkan Biaya aktual (kas) yang telah dikeluarkan adalah Rp 45.000.000.")

    # 1. Hitung PV, EV, AC
    add_heading_2(doc, "1. Perhitungan Planned Value (PV), Earned Value (EV), dan Actual Cost (AC)")
    add_bullet_p(doc, "Planned Value (PV) = Target Rencana (%) × BAC = 50% × Rp 100.000.000 = Rp 50.000.000")
    add_bullet_p(doc, "Earned Value (EV) = Progres Aktual (%) × BAC = 40% × Rp 100.000.000 = Rp 40.000.000")
    add_bullet_p(doc, "Actual Cost (AC) = Biaya Kas yang Telah Keluar = Rp 45.000.000")

    # 2. Hitung CV, SV, CPI, SPI
    add_heading_2(doc, "2. Perhitungan Variance dan Index Kinerja (CV, SV, CPI, SPI)")
    add_bullet_p(doc, "Cost Variance (CV) = EV - AC = Rp 40.000.000 - Rp 45.000.000 = -Rp 5.000.000")
    add_bullet_p(doc, "Schedule Variance (SV) = EV - PV = Rp 40.000.000 - Rp 50.000.000 = -Rp 10.000.000")
    add_bullet_p(doc, "Cost Performance Index (CPI) = EV / AC = Rp 40.000.000 / Rp 45.000.000 = 0.8889")
    add_bullet_p(doc, "Schedule Performance Index (SPI) = EV / PV = Rp 40.000.000 / Rp 50.000.000 = 0.80")

    # 3. Interpretasi
    add_heading_2(doc, "3. Interpretasi Hasil Analisis EVM")
    add_body_p(doc, "Berdasarkan nilai perhitungan di atas, status kesehatan proyek menunjukkan indikasi peringatan dini:")
    add_bullet_p(doc, "Cost Variance (CV) negatif (-Rp 5 Juta) dan CPI (0.88) < 1.0 mengindikasikan bahwa proyek mengalami pembengkakan biaya (Over Budget). Setiap Rp 1 uang yang dihabiskan hanya menghasilkan Rp 0.88 nilai pekerjaan. Terdapat pemborosan sebesar 11%.")
    add_bullet_p(doc, "Schedule Variance (SV) negatif (-Rp 10 Juta) dan SPI (0.80) < 1.0 mengindikasikan bahwa proyek mengalami keterlambatan progres (Behind Schedule). Kinerja penyelesaian pekerjaan proyek hanya berjalan 80% dari kecepatan yang dijadwalkan (tertinggal 20%).")

    # 4. Proyeksi EAC
    add_heading_2(doc, "4. Proyeksi Total Biaya Proyek (EAC)")
    add_body_p(doc,
        "Estimate at Completion (EAC) adalah peramalan total anggaran yang dibutuhkan hingga proyek usai "
        "dengan asumsi tren inefisiensi (CPI=0.8889) akan terus berlanjut. "
        "Rumus: EAC = BAC / CPI = Rp 100.000.000 / (40/45) = Rp 112.500.000. "
        "Kesimpulannya, jika kinerja tidak diperbaiki, proyek yang semula dianggarkan Rp 100 Juta akan membengkak menjadi Rp 112,5 Juta.")

    # 5. Proyeksi ETC
    add_heading_2(doc, "5. Proyeksi Sisa Kebutuhan Dana (ETC)")
    add_body_p(doc,
        "Estimate to Complete (ETC) adalah sisa modal tunai yang harus disiapkan untuk mendanai "
        "sisa pekerjaan proyek yang belum terselesaikan. "
        "Rumus: ETC = EAC - AC = Rp 112.500.000 - Rp 45.000.000 = Rp 67.500.000. "
        "Tim proyek masih membutuhkan suntikan dana sebesar Rp 67,5 Juta agar proyek dapat tuntas.")
