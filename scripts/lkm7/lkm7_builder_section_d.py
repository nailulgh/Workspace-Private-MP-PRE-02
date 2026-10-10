# -*- coding: utf-8 -*-
"""
Section D Builder: Daftar Pustaka
Pertemuan 7: Manajemen Biaya Proyek
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_d(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_body_p = helpers['add_body_p']

    add_heading_1(doc, "D. DAFTAR PUSTAKA")

    def add_hanging_indent_p(document, text):
        p = document.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        return p

    add_hanging_indent_p(doc, 
        "Kerzner, H. (2017). Project Management: A Systems Approach to Planning, Scheduling, and Controlling (12th ed.). John Wiley & Sons.")
    
    add_hanging_indent_p(doc, 
        "Larson, E. W., & Gray, C. F. (2021). Project Management: The Managerial Process (8th ed.). McGraw-Hill Education.")
    
    add_hanging_indent_p(doc, 
        "Project Management Institute. (2017). A Guide to the Project Management Body of Knowledge (PMBOK® Guide) (6th ed.). Project Management Institute, Inc.")
    
    add_hanging_indent_p(doc, 
        "Schwalbe, K. (2019). Information Technology Project Management (9th ed.). Cengage Learning.")
