# -*- coding: utf-8 -*-
"""
Section C Builder: Refleksi Mandiri
Pertemuan 9: Manajemen Sumber Daya Proyek
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_c(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "C. REFLEKSI MANDIRI")

    add_body_p(doc,
        "Bagian ini memuat refleksi kritis dan introspeksi personal saya terhadap penguasaan materi Manajemen Sumber Daya Proyek, "
        "dinamika pembentukan tim (Tuckman Model), manajemen konflik, dan akuntabilitas kerja (RACI Chart) "
        "yang dipelajari pada Pertemuan 9, serta korelasinya dengan eksekusi riset proyek kelompok PRE-02 "
        "(Sistem Prediksi Probabilitas Keterlambatan Proyek Berbasis Machine Learning sebagai Early Warning System):")

    # =========================================================================
    # REFLEKSI 1
    # =========================================================================
    add_heading_2(doc, "1. Menurut Anda, gaya kepemimpinan apa yang paling efektif dalam mengelola tim proyek, dan mengapa Anda memilih gaya kepemimpinan tersebut?")
    add_body_p(doc,
        "Menurut refleksi saya, tidak ada satu gaya kepemimpinan tunggal yang absolut untuk segala situasi. "
        "Gaya kepemimpinan yang paling efektif adalah Kepemimpinan Situasional (Situational Leadership) yang dipadukan dengan Servant Leadership (Kepemimpinan Pelayan).")
    add_body_p(doc,
        "Alasan Pemilihan: Dalam proyek rekayasa perangkat lunak dan kecerdasan buatan seperti PRE-02, tingkat kematangan tim dan karakteristik tugas berubah secara dinamis:")
    add_bullet_p(doc, "Pada fase inisiasi dan perencanaan awal (tahap Forming), pemimpin harus bersikap Direktif untuk memberikan kejelasan visi, membagi arsitektur modul, dan menetapkan standar pengkodean.")
    add_bullet_p(doc, "Ketika tim telah memasuki fase eksekusi teknis dan mencapai kematangan (Norming menuju Performing), pemimpin harus bertransformasi menjadi Servant Leader yang bersikap fasilitatif dan suportif—berfokus pada menyingkirkan hambatan logistik (seperti akses server GPU, perizinan repositori), memberdayakan otonomi anggota, serta melindungi tim dari interupsi eksternal.")
    add_body_p(doc,
        "Pendekatan ini membangun rasa kepemilikan psikologis (psychological ownership) yang tinggi di kalangan anggota tim teknis yang berpendidikan tinggi.")

    # =========================================================================
    # REFLEKSI 2
    # =========================================================================
    add_heading_2(doc, "2. Bagaimana cara Anda biasanya mengelola konflik dalam tim, dan strategi mana yang menurut Anda paling efektif untuk situasi yang berbeda?")
    add_body_p(doc,
        "Secara personal, pendekatan utama saya dalam mengelola konflik adalah dengan memisahkan orang dari masalah (Separate the People from the Problem) "
        "sebagaimana prinsip negosiasi Harvard (Fisher & Ury). Saya berusaha mendengarkan secara aktif alasan di balik tuntutan lawan bicara, "
        "mencari data objektif penguji, dan tidak membawa perbedaan pandangan ke ranah ego pribadi.")
    add_body_p(doc, "Penerapan Strategi Efektif Berdasarkan Konteks Situasi:")
    add_bullet_p(doc, "Strategi Kolaboratif (Collaborating): Paling efektif untuk konflik teknis yang memiliki dampak arsitektural jangka panjang, seperti penentuan algoritma inti ML atau perancangan skema database. Kedua pihak diajak merumuskan solusi berbasis eksperimen empiris.")
    add_bullet_p(doc, "Strategi Kompromi (Compromising): Paling efektif ketika proyek menghadapi tenggat waktu (deadline) yang sangat ketat dan kedua opsi memiliki kelebihan yang seimbang, misalnya pembagian jadwal giliran presentasi atau alokasi jam kerja.")
    add_bullet_p(doc, "Strategi Forcing (Memaksakan): Hanya saya gunakan dalam situasi darurat kritis yang menyangkut kepatuhan etika data, keamanan sistem, atau tenggat waktu pengumpulan tugas yang tidak dapat diundur sama sekali.")
    add_bullet_p(doc, "Strategi Smoothing (Menghaluskan): Digunakan sementara untuk mendinginkan tensi emosional anggota tim sebelum sesi mediasi formal dimulai.")

    # =========================================================================
    # REFLEKSI 3
    # =========================================================================
    add_heading_2(doc, "3. Berdasarkan Tuckman Model, pada tahap apa tim Anda saat ini, dan apa yang dapat Anda lakukan untuk membantu tim mencapai tahap Performing?")
    add_body_p(doc,
        "Berdasarkan evaluasi terhadap dinamika pengerjaan proyek kelompok PRE-02 hingga pertemuan ke-9 ini, "
        "saya mengidentifikasi bahwa tim kami saat ini berada pada Masa Transisi dari Tahap Norming menuju Tahap Performing.")
    add_body_p(doc, "Bukti Perilaku yang Mendukung:")
    add_bullet_p(doc, "Norma kerja telah terbentuk: Pembagian peran antara penyiapan data 7 fitur monitoring, eksperimen pemodelan machine learning, dan integrasi modul telah disepakati tanpa perdebatan personal.")
    add_bullet_p(doc, "Komunikasi berjalan terbuka melalui kanal koordinasi rutin dan saling meninjau draf deliverables sebelum finalisasi.")
    add_body_p(doc, "Kontribusi Konkret Saya untuk Mendorong Tim Mencapai Tahap Performing Penuh:")
    add_num_p(doc, "Meningkatkan Akuntabilitas Proaktif: Menyelesaikan penugasan modul saya lebih cepat dari tenggat waktu dan menawarkan bantuan teknis kepada rekan yang menghadapi kendala komputasi.", 1)
    add_num_p(doc, "Membangun Dokumentasi Pengetahuan Terbuka: Mendokumentasikan hasil pengujian dan parameter kalibrasi probabilitas secara transparan di repositori bersama agar rekan tim tidak perlu mengulang pengujian yang sama.", 2)
    add_num_p(doc, "Mendorong Budaya Umpan Balik Positif: Memberikan apresiasi atas kontribusi rekan serta menyampaikan kritik perbaikan secara konstruktif dengan berbasis bukti data.", 3)

    # =========================================================================
    # REFLEKSI 4
    # =========================================================================
    add_heading_2(doc, "4. Menurut Anda, apa tantangan terbesar dalam mengelola tim proyek yang terdiri dari anggota dengan latar belakang dan kepribadian yang beragam?")
    add_body_p(doc,
        "Tantangan terbesar terletak pada Menyelaraskan Ekspektasi dan Mengatasi Hambatan Komunikasi Kognitif (Cognitive & Communication Barriers). "
        "Dalam tim proyek teknologi informasi:")
    add_bullet_p(doc, "Perbedaan Pola Pikir Antara Anggota Analitik vs Kreatif/Sosial: Anggota dengan latar belakang Data Science cenderung berpikir sangat matematis, berorientasi pada ketelitian data, dan membutuhkan waktu riset eksploratif. Sebaliknya, anggota yang berorientasi pada manajemen atau UI/UX menginginkan eksekusi visual yang cepat, praktis, dan langsung berinteraksi dengan pengguna. Menjembatani dua kecepatan dan gaya berpikir ini membutuhkan kesabaran luar biasa.")
    add_bullet_p(doc, "Perbedaan Gaya Komunikasi (Introvert vs Ekstrovert): Anggota yang introvert sering kali memiliki analisis teknis yang sangat tajam namun enggan menyuarakannya di rapat pleno, sementara anggota yang ekstrovert dapat mendominasi percakapan tanpa sengaja. Manajer proyek harus mampu menciptakan ruang aman psikologis (psychological safety) agar setiap anggota merasa didengar secara setara.")

    # =========================================================================
    # REFLEKSI 5
    # =========================================================================
    add_heading_2(doc, "5. Bagaimana RACI Chart dapat membantu mengurangi konflik dalam tim proyek, terutama terkait ketidakjelasan peran dan tanggung jawab?")
    add_body_p(doc,
        "RACI Chart berfungsi sebagai 'Peta Hukum dan Demarkasi Teritori Kerja' yang secara radikal memangkas potensi konflik peran (role ambiguity & role conflict) "
        "melalui tiga mekanisme kerja utama:")
    add_bullet_p(doc, "Menghilangkan Asumsi Liar (Unspoken Assumptions): Banyak konflik tim dipicu oleh kalimat 'Saya kira kamu yang mengerjakannya'. Dengan RACI Chart, kolom Responsible (R) dan Accountable (A) tertulis secara hitam di atas putih, sehingga tidak ada ruang untuk saling melempar tanggung jawab.")
    add_bullet_p(doc, "Mencegah Intervensi Otoritas Berlebih (Micromanagement & Turf Wars): Kode 'Accountable (A)' yang dipegang tunggal memberikan batas kekuasaan yang jelas. Anggota lain memahami batas keterlibatannya hanya sebatas 'Consulted (C)' atau 'Informed (I)', sehingga tidak terjadi perebutan hak veto dalam persetujuan deliverables.")
    add_bullet_p(doc, "Menjaga Keseimbangan Beban Kerja (Workload Transparency): Meninjau matriks RACI secara vertikal memungkinkan tim melihat apakah ada satu individu yang kelebihan huruf 'R' dan 'A' (yang berpotensi memicu burnout dan kemarahan), sehingga realokasi beban kerja dapat dilakukan secara adil dan transparan sebelum memicu konflik antar-anggota.")
