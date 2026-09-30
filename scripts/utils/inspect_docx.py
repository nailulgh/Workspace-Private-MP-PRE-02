import docx
from docx.oxml.ns import qn

doc = docx.Document('MP_Tugas_Individu/LKM2_Manajemen_Proyek.docx')

print('=== SECTIONS ===')
for s in doc.sections:
    print('Top margin:', s.top_margin.cm, 'Bottom:', s.bottom_margin.cm, 'Left:', s.left_margin.cm, 'Right:', s.right_margin.cm)
    print('Page width:', s.page_width.cm, 'height:', s.page_height.cm)

print('\n=== PARAGRAPH SAMPLES ===')
for i in [0, 1, 3, 4, 10, 15, 17, 19, 20, 21, 22, 23, 24, 25, 26, 81, 82, 90]:
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        font_names = set(r.font.name for r in p.runs if r.font.name)
        font_sizes = set(r.font.size.pt for r in p.runs if r.font.size)
        bolds = set(r.font.bold for r in p.runs)
        print(f'P{i:02d}: text="{p.text[:40]}..." | align={p.alignment} | fonts={font_names} | sizes={font_sizes} | bold={bolds} | space_before={p.paragraph_format.space_before} | space_after={p.paragraph_format.space_after} | line_spacing={p.paragraph_format.line_spacing}')

print('\n=== TABLE SAMPLES ===')
for t_idx, t in enumerate(doc.tables):
    print(f'Table {t_idx}: style={t.style.name}, rows={len(t.rows)}, cols={len(t.columns)}')
    cell0 = t.rows[0].cells[0]
    shd = cell0._tc.get_or_add_tcPr().find(qn('w:shd'))
    shd_val = shd.get(qn('w:fill')) if shd is not None else None
    print(f'  Row 0 Cell 0 fill: {shd_val}')
