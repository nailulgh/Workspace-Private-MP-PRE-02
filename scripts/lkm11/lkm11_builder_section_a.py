# -*- coding: utf-8 -*-
"""
Section A Builder: Tugas Tertulis Individu
Pertemuan 11: Manajemen Risiko Proyek (Bagian I)
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_a(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "A. TUGAS TERTULIS INDIVIDU")

    # =========================================================================
    # SOAL 1
    # =========================================================================
    add_heading_2(doc, "1. Apa yang dimaksud dengan risiko dalam konteks proyek dan mengapa manajemen risiko penting bagi keberhasilan proyek?")
    add_body_p(doc,
        "Dalam standar manajemen proyek internasional (PMBOK® Guide 6th Edition, PMI, 2017), risiko proyek (project risk) didefinisikan "
        "sebagai suatu peristiwa atau kondisi yang tidak pasti (uncertain event or condition), yang jika terjadi, akan memberikan dampak "
        "(impact/effect) terhadap satu atau lebih tujuan proyek—seperti ruang lingkup (scope), jadwal (schedule), biaya (cost), "
        "maupun kualitas (quality). Risiko memiliki sifat probabilistik, di mana probabilitas keterjadiannya berada pada rentang terbuka 0% < P < 100%. "
        "Jika suatu peristiwa telah pasti terjadi (P = 100%), maka peristiwa tersebut bukan lagi sebuah risiko melainkan sebuah fakta, isu (issue), "
        "atau kendala (constraint) yang membutuhkan penanganan korektif langsung.")
    
    add_body_p(doc,
        "Manajemen risiko proyek (Project Risk Management) adalah seni dan ilmu mengidentifikasi, menganalisis, merencanakan respons, "
        "serta mengendalikan risiko sepanjang siklus hidup proyek. Urgensi manajemen risiko terhadap keberhasilan proyek didasari oleh faktor-faktor fundamental berikut:")
    
    add_num_p(doc, "Mereduksi Ketidakpastian Menjadi Parameter yang Terkelola: Proyek secara inheren merupakan usaha unik yang sarat kebaruan dan ketidakpastian. Tanpa manajemen risiko, proyek beroperasi dalam mode reaktif (\"memadamkan api\"), di mana tim baru bertindak setelah krisis terjadi. Manajemen risiko mengubah pendekatan reaktif menjadi pendekatan proaktif yang terencana.", 1)
    add_num_p(doc, "Perlindungan Terhadap Tiga Kendala Utama (Triple Constraints): Kejadian tak terduga seperti kegagalan dependensi server, pergantian personil kunci, atau eskalasi cacat software dapat langsung memicu pembengkakan anggaran (cost overrun) dan keterlambatan jadwal. Manajemen risiko melindungi baseline proyek melalui pencadangan kontingensi (contingency reserves) yang terhitung.", 2)
    add_num_p(doc, "Meningkatkan Akurasi Pengambilan Keputusan: Dengan memetakan profil risiko, Manajer Proyek dan pemangku kepentingan dapat membuat keputusan berbasis data (risk-informed decision making), seperti menentukan arsitektur sistem yang lebih andal atau menegosiasikan target jadwal yang lebih realistis.", 3)
    add_num_p(doc, "Konteks Kritis Proyek Riset PRE-02: Dalam pengembangan sistem Machine Learning untuk memprediksi probabilitas keterlambatan proyek PRE-02, risiko inheren seperti ketidakseimbangan data (class imbalance), overfitting pada data sintetis, dan pergeseran distribusi fitur (data drift) dapat melumpuhkan daya prediksi model. Manajemen risiko memastikan strategi mitigasi—seperti kalibrasi Platt Scaling dan validasi silang bertingkat—disiapkan sejak awal siklus eksperimen.", 4)

    # =========================================================================
    # SOAL 2
    # =========================================================================
    add_heading_2(doc, "2. Sebutkan dan jelaskan tujuh proses dalam manajemen risiko proyek menurut PMBOK Guide!")
    add_body_p(doc,
        "Dalam PMBOK® Guide 6th Edition (PMI, 2017, Bab 11, hal. 395–458), area pengetahuan Manajemen Risiko Proyek "
        "disempurnakan menjadi tujuh proses terstruktur yang berjalan secara iteratif:")

    headers_q2 = ["Nama Proses", "Kelompok Proses", "Tujuan & Mekanisme Kerja Utama", "Output Kunci"]
    data_q2 = [
        [
            "1. Plan Risk Management\n(Merencanakan Manajemen Risiko)",
            "Planning\n(Perencanaan)",
            "Mendefinisikan bagaimana aktivitas manajemen risiko akan distrukturkan dan dilaksanakan sepanjang proyek. Menetapkan metodologi, peran, anggaran risiko, waktu, kategori risiko (RBS), serta ambang batas toleransi risiko.",
            "• Risk Management Plan"
        ],
        [
            "2. Identify Risks\n(Mengidentifikasi Risiko)",
            "Planning\n(Perencanaan)",
            "Menentukan risiko-risiko individu maupun risiko proyek secara keseluruhan yang dapat memengaruhi proyek dan mendokumentasikan karakteristiknya secara rinci melalui teknik kolaboratif dan analitis.",
            "• Risk Register awal\n• Risk Report awal\n• Pembaruan Project Documents"
        ],
        [
            "3. Perform Qualitative Risk Analysis\n(Melakukan Analisis Risiko Kualitatif)",
            "Planning\n(Perencanaan)",
            "Memprioritaskan risiko individu proyek untuk analisis atau tindakan lebih lanjut dengan menilai dan menggabungkan probabilitas kemunculan (probability) dan dampak (impact)-nya ke dalam matriks penilaian.",
            "• Pembaruan Risk Register (Risk Score, Prioritas, Watch List)\n• Pembaruan Risk Report"
        ],
        [
            "4. Perform Quantitative Risk Analysis\n(Melakukan Analisis Risiko Kuantitatif)",
            "Planning\n(Perencanaan)",
            "Menganalisis secara numerik efek gabungan dari risiko individu yang diprioritaskan terhadap sasaran proyek secara keseluruhan menggunakan simulasi Monte Carlo, Decision Tree, atau Expected Monetary Value (EMV).",
            "• Pembaruan Risk Report (analisis probabilitas kuantitatif biaya/jadwal)\n• Cadangan kontingensi terkuantisasi"
        ],
        [
            "5. Plan Risk Responses\n(Merencanakan Respons Risiko)",
            "Planning\n(Perencanaan)",
            "Mengembangkan opsi, memilih strategi, dan menyepakati tindakan terencana untuk menangani paparan risiko keseluruhan serta menangani risiko individu (strategi mitigasi, transfer, penghindaran, atau pemanfaatan peluang).",
            "• Change Requests\n• Pembaruan Risk Register (Rencana Aksi Respons, Pemilik Risiko, Trigger)\n• Pembaruan Project Management Plan"
        ],
        [
            "6. Implement Risk Responses\n(Melaksanakan Respons Risiko)",
            "Executing\n(Pelaksanaan)",
            "Proses melaksanakan rencana respons risiko yang telah disepakati sebelumnya ke dalam aktivitas operasional proyek saat pemicu risiko (risk trigger) terjadi guna memastikan respons benar-benar dieksekusi.",
            "• Change Requests\n• Pembaruan Project Documents (Issue Log, Lessons Learned Register)"
        ],
        [
            "7. Monitor Risks\n(Memantau Risiko)",
            "Monitoring & Controlling\n(Pemantauan & Pengendalian)",
            "Melacak risiko yang telah teridentifikasi, memantau risiko residu, mengidentifikasi risiko baru, mengevaluasi efektivitas respons risiko sepanjang proyek, serta meninjau ulang cadangan kontingensi.",
            "• Work Performance Information\n• Change Requests\n• Pembaruan Risk Register & Risk Report\n• Pembaruan Organizational Process Assets"
        ]
    ]
    create_table(doc, headers_q2, data_q2)

    # =========================================================================
    # SOAL 3
    # =========================================================================
    add_heading_2(doc, "3. Apa perbedaan antara risiko negatif (ancaman) dan risiko positif (peluang)? Berikan contoh masing-masing dalam konteks proyek!")
    add_body_p(doc,
        "Dalam paradigma modern manajemen proyek (PMI, 2017), risiko tidak selalu bermakna buruk. Risiko dipandang secara dwidimensi "
        "yang mencakup dua sisi mata uang ketidakpastian:")
    
    add_bullet_p(doc, "Risiko Negatif (Threats / Ancaman): Peristiwa atau kondisi tidak pasti yang, jika terjadi, akan memberikan dampak merugikan (detrimental effect) terhadap pencapaian sasaran proyek (misalnya memicu keterlambatan jadwal, pembengkakan biaya, penurunan kualitas, atau pengurangan ruang lingkup). Tujuan pengelolaannya adalah meminimalkan probabilitas dan/atau dampaknya.")
    add_bullet_p(doc, "Risiko Positif (Opportunities / Peluang): Peristiwa atau kondisi tidak pasti yang, jika terjadi, akan memberikan dampak menguntungkan (beneficial effect) bagi proyek (misalnya penyelesaian tugas lebih cepat, penghematan anggaran, peningkatan kinerja sistem, atau adopsi fitur unggulan). Tujuan pengelolaannya adalah memaksimalkan probabilitas dan/atau dampaknya.")

    add_body_p(doc,
        "Tabel berikut mengilustrasikan perbandingan komprehensif beserta contoh penerapannya dalam proyek riset rekayasa machine learning PRE-02:")

    headers_q3 = ["Indikator Pembanding", "Risiko Negatif (Ancaman / Threat)", "Risiko Positif (Peluang / Opportunity)"]
    data_q3 = [
        [
            "Dampak terhadap Tujuan Proyek",
            "Merugikan: Menghambat jadwal, meningkatkan pengeluaran biaya, menurunkan performa model, atau memicu kegagalan integrasi.",
            "Menguntungkan: Mempercepat waktu rilis, mereduksi konsumsi sumber daya komputasi, atau meningkatkan akurasi sistem di atas target."
        ],
        [
            "Strategi Respons Standar PMI",
            "1. Escalate (Eskalasi)\n2. Avoid (Menghindari)\n3. Mitigate (Memitigasi)\n4. Transfer (Mengalihkan)\n5. Accept (Menerima)",
            "1. Escalate (Eskalasi)\n2. Exploit (Mengeksploitasi)\n3. Enhance (Meningkatkan)\n4. Share (Berbagi)\n5. Accept (Menerima)"
        ],
        [
            "Contoh Riil 1 dalam Proyek PRE-02",
            "Ancaman Kebocoran Data (Data Leakage): Terjadinya kontaminasi informasi antara data latih (training set) dan data uji (test set) saat ekstraksi fitur, yang menyebabkan akurasi model tampak tinggi semu saat pelatihan namun gagal total saat evaluasi data baru.",
            "Peluang Rilis Library Open-Source Teroptimasi (Exploit): Tersedianya rilis pustaka LightGBM versi terbaru yang mendukung kuantisasi terakselerasi CPU, sehingga waktu inferensi model dapat dipangkas sebesar 40% tanpa kehilangan skor AUC-ROC."
        ],
        [
            "Contoh Riil 2 dalam Proyek PRE-02",
            "Ancaman Ketidakseimbangan Kelas (Class Imbalance): Jumlah proyek terlambat dalam dataset historis jauh lebih sedikit daripada proyek tepat waktu, memicu bias model Logistic Regression yang cenderung memprediksi seluruh proyek sebagai aman.",
            "Peluang Kolaborasi Dataset Terbuka (Enhance): Peluang kemitraan dengan repository proyek open-source yang menyediakan data telemetri ribuan tugas proyek nyata, meningkatkan variasi fitur dan generalisasi model PRE-02 secara signifikan."
        ]
    ]
    create_table(doc, headers_q3, data_q3)

    # =========================================================================
    # SOAL 4
    # =========================================================================
    add_heading_2(doc, "4. Mengapa identifikasi risiko perlu dilakukan sedini mungkin dalam siklus proyek? Jelaskan dengan mengacu pada pepatah \"a stitch in time saves nine\"!")
    add_body_p(doc,
        "Pepatah klasik \"a stitch in time saves nine\" bermakna bahwa tindakan preventif kecil yang dilakukan tepat waktu di awal "
        "dapat mencegah kebutuhan melakukan pekerjaan perbaikan yang sembilan kali lebih besar dan melelahkan di kemudian hari. "
        "Dalam konteks manajemen proyek rekayasa perangkat lunak dan teknologi informasi, prinsip ini selaras sempurna dengan konsep "
        "Cost of Changes Curve yang dirumuskan oleh Barry Boehm (1981) dan standar PMI (2017).")

    add_body_p(doc,
        "Alasan-alasan mendasar mengapa identifikasi risiko wajib dilakukan sedini mungkin meliputi:")
    
    add_num_p(doc, "Eksponensialitas Biaya Perbaikan (Cost of Changes): Pada fase inisiasi dan perencanaan awal proyek, mengidentifikasi risiko kesalahan arsitektur (misalnya pemilihan framework model yang tidak kompatibel dengan pipeline backend) hanya membutuhkan revisi dokumen spesifikasi, diskusi teknis singkat, atau perubahan konfigurasi repositori yang berbiaya mendekati nol. Namun, jika risiko tersebut baru disadari pada fase pengujian akhir atau penerapan produksi, memperbaikinya menuntut perombakan arsitektur menyeluruh, penulisan ulang ribuan baris kode, pengujian regresi ulang, dan keterlambatan rilis yang berbiaya puluhan hingga ratusan kali lipat lebih mahal.", 1)
    add_num_p(doc, "Tingkat Fleksibilitas dan Kebebasan Bermanuver (Degrees of Freedom): Di awal proyek, Manajer Proyek memiliki derajat kebebasan tertinggi dalam memilih strategi alternatif—seperti mengganti algoritma, menyesuaikan ruang lingkup fitur, atau mengatur ulang struktur tim. Seiring berjalannya proyek, komitmen sumber daya dan keterikatan desain semakin terkunci, sehingga opsi mitigasi yang tersedia semakin sempit dan mahal.", 2)
    add_num_p(doc, "Menghilangkan Efek Domino Kerusakan (Cascading Failure Prevention): Satu risiko teknis yang tidak teridentifikasi di awal dapat memicu rentetan kegagalan di modul lain. Misalnya dalam proyek PRE-02, ketidakjelasan definisi fitur \"Dependency Score\" yang dibiarkan sejak awal akan merusak tahapan pengumpulan data, mengacaukan proses pelatihan model, dan akhirnya membatalkan seluruh hasil evaluasi probabilitas.", 3)
    add_num_p(doc, "Pembangunan Cadangan Kontingensi yang Realistis: Identifikasi risiko dini memungkinkan tim mengalokasikan cadangan kontingensi biaya dan waktu (contingency buffer) secara terukur di dalam baseline jadwal, bukan sekadar menebak-nebak.", 4)

    # =========================================================================
    # SOAL 5
    # =========================================================================
    add_heading_2(doc, "5. Sebutkan dan jelaskan minimal empat teknik identifikasi risiko (brainstorming, Delphi, SWOT, checklist) beserta kelebihan masing-masing!")
    add_body_p(doc,
        "PMBOK® Guide (PMI, 2017, hal. 414–417) merekomendasikan berbagai teknik terstruktur untuk mengidentifikasi risiko proyek secara komprehensif. "
        "Berikut adalah penjelasan mendalam tentang empat teknik utama beserta kelebihan dan karakteristiknya:")

    headers_q5 = ["Teknik Identifikasi", "Mekanisme Pelaksanaan & Karakteristik", "Kelebihan Utama", "Keterbatasan & Mitigasi"]
    data_q5 = [
        [
            "1. Brainstorming\n(Curah Pendapat)",
            "Sesi diskusi interaktif yang melibatkan tim proyek, pakar lintas fungsi, dan pemangku kepentingan untuk menghasilkan daftar risiko secara spontan dan kreatif tanpa penilaian awal.",
            "• Menghasilkan daftar risiko dalam jumlah besar secara cepat\n• Mendorong keterbukaan dan sinergi multidisiplin\n• Membangun rasa kepemilikan tim (team buy-in) terhadap risiko",
            "• Rentan dominasi oleh individu yang vokal atau berstatus tinggi\n• Mitigasi: Dipandu fasilitator netral menggunakan teknik silent brainstorming/sticky notes."
        ],
        [
            "2. Delphi Technique\n(Teknik Delphi)",
            "Teknik komunikasi terstruktur di mana panel pakar independen menjawab kuesioner risiko secara anonim dalam beberapa putaran (rounds), dengan fasilitator merangkum konsensus di setiap putaran.",
            "• Menghilangkan bias dominasi sosial, hierarki jabatan, dan groupthink\n• Menghasilkan opini pakar yang objektif dan independen\n• Menjangkau pakar geografis luas",
            "• Memerlukan waktu pelaksanaan yang cukup lama (multi-putaran)\n• Bergantung pada komitmen dan responsivitas para pakar terpilih."
        ],
        [
            "3. SWOT Analysis\n(Strengths, Weaknesses, Opportunities, Threats)",
            "Menganalisis proyek dari empat perspektif: kekuatan internal (Strengths), kelemahan internal (Weaknesses), peluang eksternal (Opportunities), dan ancaman eksternal (Threats).",
            "• Menghubungkan faktor risiko internal organisasi dengan lingkungan eksternal\n• Secara alami mengidentifikasi risiko positif (peluang) dan risiko negatif (ancaman) secara seimbang\n• Berorientasi strategis",
            "• Cenderung menghasilkan analisis tingkat makro yang membutuhkan dekomposisi lanjutan agar dapat ditindaklanjuti pada level tugas teknis."
        ],
        [
            "4. Checklist Analysis\n(Analisis Daftar Periksa)",
            "Mengembangkan daftar risiko potensial berdasarkan data historis, arsip pelajaran yang dipetik (lessons learned), dan taksonomi risiko dari proyek-proyek serupa sebelumnya.",
            "• Sangat cepat, terstruktur, dan mudah digunakan\n• Memastikan risiko umum yang lazim terjadi di industri tidak terlewatkan\n• Memanfaatkan aset proses organisasi (OPA)",
            "• Rentan menimbulkan rasa aman palsu (complacency) di mana tim hanya terpaku pada daftar checklist dan gagal mendeteksi risiko baru yang belum pernah terjadi sebelumnya."
        ]
    ]
    create_table(doc, headers_q5, data_q5)

    # =========================================================================
    # SOAL 6
    # =========================================================================
    add_heading_2(doc, "6. Apa yang dimaksud dengan Probability and Impact Matrix dan bagaimana cara menggunakannya dalam analisis risiko kualitatif?")
    add_body_p(doc,
        "Probability and Impact Matrix (Matriks Probabilitas dan Dampak / P&I Matrix) adalah alat visual dan analitis standar "
        "(PMI, 2017) yang digunakan dalam proses Perform Qualitative Risk Analysis untuk memetakan dan menentukan tingkat prioritas risiko "
        "berdasarkan kombinasi dua variabel penilaian: probabilitas terjadinya risiko (Likelihood/Probability) dan besarnya dampak "
        "(Consequence/Impact) terhadap tujuan proyek jika risiko tersebut benar-benar terjadi.")

    add_body_p(doc,
        "Mekanisme Operasional Penggunaan Matriks Probabilitas dan Dampak:")
    add_num_p(doc, "Penetapan Skala Penilaian: Tim proyek menetapkan skala ordinal (misalnya: 1 = Sangat Rendah hingga 5 = Sangat Tinggi) atau skala kardinal (misalnya: probabilitas 0.1, 0.3, 0.5, 0.7, 0.9 dan dampak berbobot 0.05, 0.10, 0.20, 0.40, 0.80) lengkap dengan definisi kriteria terukur untuk biaya, waktu, dan mutu.", 1)
    add_num_p(doc, "Penilaian Parameter Setiap Risiko: Untuk setiap risiko yang tercatat dalam Risk Register, tim penilai/pakar memberikan skor probabilitas (P) dan skor dampak (I) berdasarkan data historis, estimasi ahli, atau penilaian situasional.", 2)
    add_num_p(doc, "Pemetaan Koordinat ke dalam Matriks Grid: Mengintegrasikan nilai P dan I ke dalam sel kisi-kisi matriks untuk mendapatkan Risk Score dan mengkategorikan posisi risiko ke dalam tiga zona tingkat keparahan risiko:", 3)
    add_subbullet_p(doc, "Zona Merah (High Risk / Prioritas Tinggi): Risiko dengan skor tinggi (misal kuadran kanan atas). Membutuhkan tindakan mitigasi agresif, rencana respons formal, dan pengawasan ketat oleh pimpinan proyek.")
    add_subbullet_p(doc, "Zona Kuning (Medium Risk / Prioritas Sedang): Risiko dengan skor moderat. Memerlukan rencana kontingensi terencana dan pemantauan berkala.")
    add_subbullet_p(doc, "Zona Hijau (Low Risk / Prioritas Rendah): Risiko dengan skor rendah. Cukup dimasukkan ke dalam daftar pantau (watch list) tanpa alokasi anggaran mitigasi aktif.")
    add_num_p(doc, "Penetapan Ambang Batas Aksi (Risk Threshold): Matriks membantu organisasi menentukan ambang batas di mana suatu risiko memerlukan eskalasi ke komite manajemen atau cukup dikelola di tingkat operasional tim pengembang.", 4)

    # =========================================================================
    # SOAL 7
    # =========================================================================
    add_heading_2(doc, "7. Jelaskan bagaimana risk score dihitung dan bagaimana skor tersebut digunakan untuk menentukan prioritas risiko!")
    add_body_p(doc,
        "Risk Score (Skor Risiko atau Tingkat Keparahan Risiko / Severity Score) adalah representasi kuantitatif dari paparan risiko "
        "yang dihasilkan dari perkalian matematis antara nilai probabilitas terjadinya risiko dengan nilai dampaknya terhadap sasaran proyek. "
        "Formulasi fundamental perhitungannya dirumuskan sebagai berikut:")
    
    # Formula box display
    p_box = doc.add_paragraph()
    p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box.paragraph_format.space_before = Pt(4)
    p_box.paragraph_format.space_after = Pt(8)
    p_box.paragraph_format.line_spacing = 1.15
    r_box = p_box.add_run("Risk Score = Probability (P) × Impact (I)")
    r_box.font.name = "Times New Roman"
    r_box.font.size = Pt(13)
    r_box.font.bold = True

    add_body_p(doc,
        "Pemanfaatan Risk Score dalam Penentuan Prioritas Risiko:")
    add_num_p(doc, "Normalisasi dan Pemeringkatan Relatif (Risk Ranking): Dalam matriks berdimensi 5 × 5 dengan skala ordinal 1 hingga 5, Risk Score berkisar antara 1 (skor minimum: 1 × 1) hingga 25 (skor maksimum: 5 × 5). Nilai numerik ini memungkinkan Manajer Proyek melakukan sorting (pemeringkatan dari tertinggi ke terendah) secara objektif terhadap seluruh daftar risiko.", 1)
    add_num_p(doc, "Segmentasi Kebijakan Tindakan (Risk Categorization): Berdasarkan konvensi manajemen proyek standar, rentang skor dikelompokkan menjadi tiga tingkatan aksi manajerial:", 2)
    add_subbullet_p(doc, "Prioritas Tinggi (Skor 15–25): Status Kritis / Bahaya. Menjadi fokus utama sprint manajemen risiko. Harus memiliki pemilik risiko resmi (Risk Owner), strategi mitigasi terperinci, alokasi anggaran kontingensi cadangan, dan dipantau setiap hari/minggu.")
    add_subbullet_p(doc, "Prioritas Sedang (Skor 6–14): Status Waspada. Membutuhkan rencana respons kontingensi yang siap dieksekusi jika indikator pemicu (trigger) muncul, serta ditinjau dalam rapat status berkala.")
    add_subbullet_p(doc, "Prioritas Rendah (Skor 1–5): Status Dapat Ditoleransi. Tidak memerlukan tindakan pencegahan aktif; risiko dicatat dalam Watch List untuk dipantau secara pasif agar tidak berkembang menjadi ancaman yang lebih besar.")
    add_num_p(doc, "Alokasi Efisien Sumber Daya yang Terbatas: Organisasi proyek selalu memiliki keterbatasan waktu, tenaga kerja, dan anggaran. Risk score memastikan bahwa sumber daya pencegahan difokuskan secara optimal pada 20% risiko paling berbahaya yang berpotensi menimbulkan 80% kehancuran proyek (Prinsip Pareto Risiko).", 3)

    # =========================================================================
    # SOAL 8
    # =========================================================================
    add_heading_2(doc, "8. Apa yang dimaksud dengan Risk Register dan informasi apa saja yang harus dicatat di dalamnya pada tahap awal setelah identifikasi risiko?")
    add_body_p(doc,
        "Risk Register (Daftar Risiko) adalah dokumen repositori sentral dinamis (living document) yang berfungsi mencatat seluruh rincian "
        "risiko proyek yang telah diidentifikasi, hasil analisis risiko, rencana respons, serta status pemantauan risiko sepanjang siklus hidup proyek "
        "(PMBOK® Guide 6th Edition, PMI, 2017). Risk Register diinisiasi pada proses Identify Risks dan secara bertahap diperbarui (progressive elaboration) "
        "seiring berjalannya proses analisis kualitatif, kuantitatif, perencanaan respons, hingga pemantauan risiko.")

    add_body_p(doc,
        "Pada tahap awal—tepat setelah proses identifikasi risiko (Identify Risks) selesai dilaksanakan dan sebelum analisis kualitatif/kuantitatif mendalam—"
        "informasi penting yang wajib dicatat di dalam Risk Register mencakup elemen-elemen berikut:")

    add_num_p(doc, "ID Risiko (Risk Identifier): Kode pengenal unik untuk setiap risiko (misalnya: R-001, R-002, TECH-01) guna memfasilitasi pelacakan dan referensi silang dokumen.", 1)
    add_num_p(doc, "Deskripsi Risiko Terstruktur (Risk Description / Event Statement): Pernyataan risiko yang dirumuskan secara spesifik dan tidak ambigu, idealnya menggunakan format standar kausal: \"Karena [penyebab/faktor pemicu], [peristiwa risiko] mungkin terjadi, yang dapat mengakibatkan [dampak terhadap sasaran proyek]\".", 2)
    add_num_p(doc, "Kategori Risiko (Risk Category / RBS Level): Klasifikasi sumber risiko berdasarkan Risk Breakdown Structure (RBS), seperti Risiko Teknis (algoritma, arsitektur, data), Risiko Manajemen (jadwal, komunikasi, alokasi SDM), Risiko Komersial/Eksternal, atau Risiko Kualitas.", 3)
    add_num_p(doc, "Akar Penyebab Risiko (Root Causes): Sumber fundamental atau kondisi laten yang memicu terjadinya peristiwa risiko (diidentifikasi melalui teknik 5-Whys atau Ishikawa Diagram).", 4)
    add_num_p(doc, "Gejala dan Pemicu Risiko (Risk Triggers / Warning Signs): Indikator atau sinyal peringatan dini teramati yang menunjukkan bahwa peristiwa risiko sedang atau akan segera terjadi (misalnya: latensi API mulai melambat di atas 150 ms).", 5)
    add_num_p(doc, "Pemilik Risiko Potensial (Potential Risk Owner): Anggota tim proyek atau pemangku kepentingan yang dinominasikan untuk memantau risiko dan bertanggung jawab memimpin respons jika risiko terjadi.", 6)
    add_num_p(doc, "Daftar Respons Potensial Awal (Potential Risk Responses): Gagasan awal atau opsi strategi penanganan spontan yang diusulkan oleh tim selama sesi identifikasi (misalnya: menyiapkan modul fallback, kompresi bobot model, atau negosiasi ulang tenggat waktu).", 7)
