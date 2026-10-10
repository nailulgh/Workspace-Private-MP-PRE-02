# -*- coding: utf-8 -*-
"""
Section B Builder: Tugas Latihan & Studi Kasus Individu
Pertemuan 11: Perancangan Probability & Impact Matrix serta Simulasi Penilaian Risiko Proyek PRE-02
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_b(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "B. TUGAS LATIHAN DAN STUDI KASUS INDIVIDU: PERANCANGAN MATRIKS PROBABILITAS DAN DAMPAK SERTA SIMULASI PENILAIAN RISIKO")

    add_body_p(doc,
        "Untuk mengoperasionalkan konsep analisis risiko kualitatif sesuai standar PMBOK® Guide (PMI, 2017), "
        "tugas latihan individu ini menyajikan perancangan terstruktur skala penilaian risiko, pemetaan kisi-kisi "
        "Probability and Impact Matrix (P&I Matrix) 5 × 5, serta simulasi pengisian Risk Register awal pada proyek "
        "riset PRE-02 (Sistem Prediksi Probabilitas Keterlambatan Proyek Berbasis Machine Learning sebagai Early Warning System).")

    # =========================================================================
    # 1. PENETAPAN KRITERIA SKALA PROBABILITAS DAN DAMPAK
    # =========================================================================
    add_heading_2(doc, "1. Perancangan Kriteria Skala Penilaian Kualitatif (Skala Ordinal 1–5)")
    add_body_p(doc,
        "Agar penilaian risiko tidak bersifat subjektif atau bias, ditetapkan skala kardinal/ordinal 1 hingga 5 dengan definisi kriteria terukur:")

    # Tabel Skala Probabilitas
    add_heading_3(doc, "a. Definisi Kriteria Skala Probabilitas (Likelihood Scale)")
    headers_prob = ["Nilai", "Tingkat Probabilitas", "Rentang Frekuensi / Probabilitas Terjadinya", "Deskripsi Kualitatif"]
    data_prob = [
        ["1", "Sangat Rendah (Very Low)", "P < 10%", "Peristiwa hampir mustahil terjadi; hanya mungkin timbul dalam kondisi yang sangat luar biasa."],
        ["2", "Rendah (Low)", "10% ≤ P < 30%", "Kecil kemungkinan terjadi, namun pernah terdokumentasi dalam proyek-proyek serupa sebelumnya."],
        ["3", "Sedang (Moderate)", "30% ≤ P < 50%", "Memiliki kemungkinan seimbang untuk terjadi selama siklus hidup pengembangan sistem."],
        ["4", "Tinggi (High)", "50% ≤ P < 70%", "Sangat mungkin terjadi; kondisi pemicu risiko telah teramati dalam operasional proyek."],
        ["5", "Sangat Tinggi (Very High)", "P ≥ 70%", "Hampir pasti terjadi kecuali dilakukan tindakan intervensi pencegahan aktif seketika."]
    ]
    create_table(doc, headers_prob, data_prob)

    # Tabel Skala Dampak
    add_heading_3(doc, "b. Definisi Kriteria Skala Dampak terhadap Sasaran Proyek (Impact Scale)")
    headers_impact = ["Nilai", "Tingkat Dampak", "Dampak Jadwal (Schedule)", "Dampak Biaya (Cost)", "Dampak Kualitas Model ML & Teknis"]
    data_impact = [
        ["1", "Sangat Rendah", "Keterlambatan < 2 hari kalender; tidak berdampak pada jalur kritis.", "Peningkatan biaya tidak signifikan (< 2% anggaran baseline).", "Penurunan metrik akurasi/AUC-ROC < 0.01; tidak terasa pada performa."],
        ["2", "Rendah", "Keterlambatan 3–5 hari; dapat diserap oleh float jalur non-kritis.", "Peningkatan biaya 2% – 5% dari total anggaran proyek.", "Penurunan AUC-ROC 0.01 – 0.03; model tetap memenuhi batas toleransi."],
        ["3", "Sedang", "Keterlambatan 6–10 hari; memengaruhi milestone utama sprint.", "Peningkatan biaya 5% – 10%; membutuhkan realokasi anggaran.", "Penurunan AUC-ROC 0.03 – 0.06; kurva kalibrasi probabilitas terdistorsi."],
        ["4", "Tinggi", "Keterlambatan 11–20 hari; menggeser baseline rilis akhir.", "Peningkatan biaya 10% – 20%; memerlukan persetujuan sponsor.", "Akurasi model anjlok signifikan; Brier Score > 0.25 (sistem tidak andal)."],
        ["5", "Sangat Tinggi", "Keterlambatan > 20 hari; kegagalan total penyelesaian proyek.", "Peningkatan biaya > 20%; pembengkakan kritis tak terkendali.", "Kegagalan sistem total; model tidak dapat inferensi atau crash produksi."]
    ]
    create_table(doc, headers_impact, data_impact)

    # =========================================================================
    # 2. PROBABILITY & IMPACT MATRIX GRID 5x5
    # =========================================================================
    add_heading_2(doc, "2. Perancangan Probability & Impact Matrix (Kisi-Kisi 5 × 5)")
    add_body_p(doc,
        "Berdasarkan formula Risk Score = P × I, disusun kisi-kisi matriks 5 × 5 yang memetakan skor risiko "
        "ke dalam tiga zona ambang batas tindakan (Risk Severity Thresholds):")
    
    add_bullet_p(doc, "Zona Merah (Prioritas Tinggi / High Risk, Skor 15–25): Tindakan mitigasi wajib dirancang, dialokasikan anggaran cadangan, dan dipantau intensif setiap minggu.")
    add_bullet_p(doc, "Zona Kuning (Prioritas Sedang / Medium Risk, Skor 6–14): Rencana kontingensi disiapkan dan dipantau secara berkala pada setiap sprint review.")
    add_bullet_p(doc, "Zona Hijau (Prioritas Rendah / Low Risk, Skor 1–5): Dicatat dalam Watch List untuk pemantauan pasif tanpa alokasi sumber daya aktif.")

    headers_matrix = ["Probabilitas (P) ↓ / Dampak (I) →", "1 (Sangat Rendah)", "2 (Rendah)", "3 (Sedang)", "4 (Tinggi)", "5 (Sangat Tinggi)"]
    data_matrix = [
        ["5 (Sangat Tinggi)", "5 (Hijau / Rendah)", "10 (Kuning / Sedang)", "15 (Merah / Tinggi)", "20 (Merah / Tinggi)", "25 (Merah / Kritis)"],
        ["4 (Tinggi)", "4 (Hijau / Rendah)", "8 (Kuning / Sedang)", "12 (Kuning / Sedang)", "16 (Merah / Tinggi)", "20 (Merah / Tinggi)"],
        ["3 (Sedang)", "3 (Hijau / Rendah)", "6 (Kuning / Sedang)", "9 (Kuning / Sedang)", "12 (Kuning / Sedang)", "15 (Merah / Tinggi)"],
        ["2 (Rendah)", "2 (Hijau / Rendah)", "4 (Hijau / Rendah)", "6 (Kuning / Sedang)", "8 (Kuning / Sedang)", "10 (Kuning / Sedang)"],
        ["1 (Sangat Rendah)", "1 (Hijau / Rendah)", "2 (Hijau / Rendah)", "3 (Hijau / Rendah)", "4 (Hijau / Rendah)", "5 (Hijau / Rendah)"]
    ]
    create_table(doc, headers_matrix, data_matrix)

    # =========================================================================
    # 3. SIMULASI PENGISIAN RISK REGISTER AWAL PROYEK PRE-02
    # =========================================================================
    add_heading_2(doc, "3. Simulasi Penilaian dan Pengisian Risk Register Awal Proyek PRE-02")
    add_body_p(doc,
        "Berikut adalah simulasi pencatatan dan penilaian kualitatif untuk lima risiko representatif (empat ancaman dan satu peluang) "
        "yang diidentifikasi pada sistem riset machine learning PRE-02:")

    headers_reg = ["ID", "Tipe", "Pernyataan Kejadian Risiko", "Akar Penyebab", "P", "I", "Skor", "Prioritas", "Rencana Tindak Lanjut"]
    data_reg = [
        [
            "R-01",
            "Ancaman\n(Threat)",
            "Terjadi overfitting model pada data sintetis sehingga akurasi prediksi probabilitas anjlok saat diuji data riil.",
            "Keterbatasan variasi sampel data historis proyek dan bobot parameter sintetis yang kurang representatif.",
            "4",
            "4",
            "16",
            "Tinggi\n(Merah)",
            "Terapkan validasi silang bertingkat (Stratified K-Fold), regularisasi L2/Dropout, dan teknik kalibrasi probabilitas Isotonic Regression."
        ],
        [
            "R-02",
            "Ancaman\n(Threat)",
            "Latensi waktu inferensi API prediksi melebihi batas toleransi SLA (> 250 ms per request).",
            "Ukuran ensemble model deep learning MLP/Boosting terlalu besar dan belum dioptimasi untuk komputasi runtime.",
            "3",
            "3",
            "9",
            "Sedang\n(Kuning)",
            "Lakukan kuantisasi model ke format ONNX Runtime int8 dan caching fitur input yang sering diakses."
        ],
        [
            "R-03",
            "Ancaman\n(Threat)",
            "Keterlambatan integrasi modul backend pipeline dengan dashboard antarmuka visualisasi risiko.",
            "Miskomunikasi spesifikasi skema JSON dan perbedaan tipe data output probabilitas antara Data Scientist dan Backend Dev.",
            "3",
            "4",
            "12",
            "Sedang\n(Kuning)",
            "Definisikan kontrak data formal (OpenAPI/Swagger spec) dan adopsi pengujian otomatis antarmuka (API contract testing)."
        ],
        [
            "R-04",
            "Ancaman\n(Threat)",
            "Keterbatasan kuota komputasi cloud/GPU gratisan selama fase hyperparameter tuning.",
            "Ketergantungan pada infrastruktur pihak ketiga gratis yang memiliki pembatasan durasi eksekusi session.",
            "2",
            "2",
            "4",
            "Rendah\n(Hijau)",
            "Masukkan ke dalam Watch List; siapkan lingkungan fallback eksekusi lokal menggunakan CPU multiprocessing bila kuota habis."
        ],
        [
            "R-05",
            "Peluang\n(Opp.)",
            "Tersedianya library LightGBM versi terbaru yang mendukung paralelisasi CPU native dan reduksi memori hingga 50%.",
            "Perkembangan rilis open-source framework komputasi machine learning global yang aktif diperbarui komunitas.",
            "4",
            "4",
            "16",
            "Tinggi\n(Eksploitasi)",
            "Strategi Exploit: Segera lakukan uji kompatibilitas library baru pada branch eksperimen dan migrasikan pipeline ke versi tersebut."
        ]
    ]
    create_table(doc, headers_reg, data_reg)
