# -*- coding: utf-8 -*-
"""
Section C Builder: Refleksi Mandiri
Pertemuan 11: Manajemen Risiko Proyek (Bagian I)
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_c(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']

    add_heading_1(doc, "C. REFLEKSI MANDIRI")

    add_body_p(doc,
        "Bagian ini menyajikan perenungan kritis dan evaluasi personal mahasiswa terhadap konsep teoritis "
        "manajemen risiko proyek serta dinamika penerapannya pada proyek riset kelompok PRE-02 (Sistem Prediksi "
        "Probabilitas Keterlambatan Proyek Berbasis Machine Learning sebagai Early Warning System):")

    # =========================================================================
    # REFLEKSI 1
    # =========================================================================
    add_heading_2(doc, "1. Menurut Anda, mengapa identifikasi risiko perlu dilakukan sedini mungkin dalam proyek, dan apa yang terjadi jika identifikasi risiko dilakukan terlambat?")
    add_body_p(doc,
        "Identifikasi risiko sedini mungkin merupakan fondasi paling esensial dalam menjaga stabilitas proyek. "
        "Pada tahap awal inisiasi dan perencanaan, tingkat ketidakpastian (uncertainty) berada pada titik tertingginya, "
        "namun pada saat yang sama, biaya untuk melakukan penyesuaian atau perubahan arah (cost of changes) berada pada tingkat terendah. "
        "Dengan mengenali potensi ancaman di awal, tim memiliki keleluasaan merancang arsitektur sistem yang resilien (tahan banting), "
        "menentukan metodologi yang tepat, serta mengalokasikan cadangan waktu dan anggaran kontingensi secara terukur.")
    
    add_body_p(doc,
        "Jika identifikasi risiko dilakukan terlambat atau diabaikan hingga fase eksekusi akhir, dampak yang ditimbulkan sangat destruktif:")
    add_bullet_p(doc, "Pergeseran Paradigma Menjadi Reaktif: Tim proyek terperangkap dalam kepanikan \"memadamkan api\" (firefighting mode), di mana seluruh fokus dan energi habis terkuras untuk merespons krisis darurat alih-alih menyelesaikan pekerjaan bernilai tambah.")
    add_bullet_p(doc, "Pembengkakan Biaya dan Keterlambatan Eksponensial: Masalah teknis yang baru terungkap di akhir (misalnya format probabilitas output model tidak kompatibel dengan visualisasi antarmuka) menuntut perombakan arsitektur besar-besaran, penulisan ulang modul, dan pengujian ulang menyeluruh yang menguras anggaran dan merusak jadwal baseline rilis.")
    add_bullet_p(doc, "Degradasi Kualitas Produk Akhir: Di bawah tekanan tenggat waktu yang kian mendekat, tim sering kali terpaksa mengambil jalan pintas (technical debt) yang mengorbankan kualitas dan keandalan sistem.")

    # =========================================================================
    # REFLEKSI 2
    # =========================================================================
    add_heading_2(doc, "2. Berdasarkan proyek kelompok Anda, risiko apa yang menurut Anda paling mungkin terjadi dan mengapa?")
    add_body_p(doc,
        "Dalam konteks riset proyek PRE-02, risiko yang memiliki probabilitas tertinggi untuk terjadi adalah "
        "Ketidakseimbangan Distribusi Data (Class Imbalance) dan Kebocoran Data (Data Leakage) pada saat rekayasa fitur data sintetis.")
    
    add_body_p(doc,
        "Analisis Mengapa Risiko Ini Paling Mungkin Terjadi:")
    add_bullet_p(doc, "Karakteristik Data Proyek Riil: Secara empiris di industri perangkat lunak, proporsi proyek yang mengalami keterlambatan parah biasanya lebih sedikit dibandingkan proyek yang berhasil diselesaikan, atau sebaliknya tergantung kompleksitas modul. Ketika dataset sintetis dibangkitkan dari 7 parameter pemantauan proyek, rentang fitur dapat condong ke salah satu kelas, memicu bias pada algoritma Logistic Regression dan Random Forest.")
    add_bullet_p(doc, "Kerentanan Pipeline Pra-Pemrosesan: Mahasiswa sering kali melakukan penskalaan data (feature scaling seperti StandardScaler) atau seleksi fitur sebelum membagi data menjadi train set dan test set. Kesalahan metodologis ini menyebabkan informasi data uji \"bocor\" ke model saat pelatihan, menghasilkan skor AUC-ROC yang tampak sempurna saat eksperimen awal (overfitting semu), namun akurasi prediksinya anjlok drastis ketika diuji pada skenario data proyek baru.")
    add_bullet_p(doc, "Dampak Langsung pada Sasaran Riset: Jika risiko ini terwujud, fungsi aktivasi sigmoid/predict_proba tidak akan terkalibrasi dengan baik, sehingga probabilitas numerik kontinu [0.0 - 1.0] yang dihasilkan gagal berfungsi sebagai sistem peringatan dini (Early Warning System) yang tepercaya.")

    # =========================================================================
    # REFLEKSI 3
    # =========================================================================
    add_heading_2(doc, "3. Apa tantangan terbesar yang Anda hadapi saat melakukan brainstorming identifikasi risiko dengan tim?")
    add_body_p(doc,
        "Saat melaksanakan sesi curah pendapat (brainstorming) identifikasi risiko bersama rekan-rekan tim proyek, "
        "tantangan terbesar yang muncul berakar pada tiga dimensi psikologis dan komunikasi kelompok:")

    add_num_p(doc, "Bias Optimisme Berlebihan (Optimism Bias / Planning Fallacy): Sebagian anggota tim memiliki kecenderungan psikologis untuk berasumsi bahwa rencana teknis yang disusun akan berjalan mulus tanpa kendala (\"kode program pasti bebas bug\", \"server komputasi tidak akan pernah error\"). Mengubah pola pikir optimis naif ini menjadi pola pikir kritis berbasis manajemen risiko menuntut fasilitasi yang intensif.", 1)
    add_num_p(doc, "Kerancuan Antara Gejala (Symptoms), Akar Masalah (Root Causes), dan Risiko (Risk Events): Anggota tim sering kali mencatat isu operasional harian atau fakta saat ini sebagai risiko. Misalnya, mereka menyebut \"laptop lemot\" sebagai risiko, padahal itu adalah kendala saat ini. Fasilitator harus secara konsisten membimbing tim untuk merumuskan pernyataan risiko menggunakan struktur standar kausal: \"Karena [akar penyebab], [kejadian tidak pasti] dapat terjadi, yang berdampak pada [tujuan proyek]\".", 2)
    add_num_p(doc, "Keengganan Mengungkapkan Kelemahan Teknis (Psychological Safety Barrier): Timbul rasa sungkan atau takut dianggap tidak kompeten jika seorang anggota mengungkapkan potensi kegagalan pada modul yang menjadi tanggung jawabnya sendiri. Membangun ruang diskusi yang aman secara psikologis—di mana pengungkapan risiko dipandang sebagai tindakan profesional yang menyelamatkan tim—menjadi tantangan kepemimpinan tersendiri.", 3)

    # =========================================================================
    # REFLEKSI 4
    # =========================================================================
    add_heading_2(doc, "4. Menurut Anda, bagaimana cara membedakan antara risiko yang perlu mendapat perhatian segera (prioritas tinggi) dan risiko yang hanya perlu dipantau (prioritas rendah)?")
    add_body_p(doc,
        "Untuk mendemarkasi secara objektif antara risiko berkategori prioritas tinggi dan prioritas rendah, "
        "saya menerapkan kombinasi analisis kuantitatif Risk Score dan kriteria kualitatif multidimensi:")

    add_num_p(doc, "Penerapan Ambang Batas Risk Score (Threshold Cutoff): Melalui matriks P&I 5 × 5, risiko dengan Risk Score ≥ 15 (Zona Merah) secara otomatis diklasifikasikan sebagai Prioritas Tinggi. Risiko dengan skor 1–5 (Zona Hijau) diklasifikasikan sebagai Prioritas Rendah.", 1)
    add_num_p(doc, "Urgensi Waktu dan Kedekatan Kejadian (Risk Proximity / Urgency): Suatu risiko—meskipun skor matematisnya moderat—wajib dinaikkan statusnya menjadi prioritas tinggi jika jendela waktu terjadinya sangat dekat (misalnya kegagalan integrasi API yang dijadwalkan besok lusa), karena tim memiliki waktu reaksi yang sangat sempit untuk merespons.", 2)
    add_num_p(doc, "Keterkaitan dengan Jalur Kritis (Critical Path Impact): Jika peristiwa risiko berpotensi langsung menunda aktivitas yang berada di jalur kritis jadwal proyek atau menghancurkan metrik evaluasi utama (AUC-ROC < 0.80), maka risiko tersebut wajib mendapatkan perhatian prioritas tinggi seketika.", 3)
    add_num_p(doc, "Strategi Pengelolaan yang Berbeda: Risiko prioritas tinggi memerlukan tindakan mitigasi aktif (proactive mitigation plan), alokasi cadangan kontingensi, dan penunjukan pemilik risiko spesifik. Sebaliknya, risiko prioritas rendah cukup didaftarkan ke dalam Watch List dan ditinjau secara pasif pada evaluasi berkala tanpa membuang sumber daya operasional.", 4)

    # =========================================================================
    # REFLEKSI 5
    # =========================================================================
    add_heading_2(doc, "5. Bagaimana Anda akan menggunakan Risk Register yang telah disusun untuk membantu pengelolaan proyek ke depan?")
    add_body_p(doc,
        "Risk Register bukanlah dokumen statis administratif yang disusun sekali untuk formalitas laporan lalu disimpan di laci. "
        "Saya akan mentransformasikannya menjadi dokumen dinamis yang hidup (living document) dan instrumen navigasi operasional proyek "
        "melalui empat langkah operasional konkret:")

    add_num_p(doc, "Agenda Tetap dalam Setiap Rapat Mingguan (Standing Agenda Item): Menjadikan peninjauan Risk Register sebagai agenda wajib berdurasi 10 menit dalam setiap rapat koordinasi atau Sprint Review tim PRE-02. Tim akan meninjau apakah status risiko berubah, apakah pemicu risiko (risk trigger) mulai muncul, dan apakah ada risiko baru yang teridentifikasi seiring progres riset.", 1)
    add_num_p(doc, "Mekanisme Pemicu Tindakan Respons Cepat (Early Warning Triggers): Memanfaatkan kolom \"Risk Triggers\" dalam register sebagai sistem alarm dini. Ketika indikator peringatan terdeteksi (misalnya metrik akurasi validasi model mulai menunjukkan gejala stagnasi), pemilik risiko yang telah ditunjuk dapat segera mengeksekusi rencana kontingensi tanpa perlu menunggu persetujuan birokratis yang panjang.", 2)
    add_num_p(doc, "Pengelolaan dan Audit Cadangan Kontingensi (Contingency Reserves Management): Menggunakan Risk Register untuk memantau penggunaan cadangan waktu float dan anggaran komputasi. Jika suatu risiko berhasil dimitigasi atau masa ancamannya telah terlewati, cadangan kontingensi yang dialokasikan untuk risiko tersebut dapat dilepaskan kembali ke cadangan umum proyek.", 3)
    add_num_p(doc, "Repositori Pelajaran yang Dipetik (Lessons Learned Repository): Pada fase penutupan proyek nantinya, seluruh catatan historis mengenai efektivitas respons risiko di dalam Risk Register akan didokumentasikan sebagai aset proses organisasi (OPA) berharga untuk penelitian-penelitian rekayasa perangkat lunak dan machine learning di masa depan.", 4)
