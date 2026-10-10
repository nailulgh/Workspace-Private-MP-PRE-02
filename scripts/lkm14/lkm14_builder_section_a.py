# -*- coding: utf-8 -*-
"""
Section A Builder: Tugas Tertulis Individu
Pertemuan 14: Pendekatan Agile/Hybrid & Review Perkuliahan
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_a(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "A. TUGAS TERTULIS INDIVIDU")

    # =========================================================================
    # SOAL 1
    # =========================================================================
    add_heading_2(doc, "1. Apa yang dimaksud dengan pendekatan prediktif (Waterfall) dalam manajemen proyek dan jenis proyek apa yang paling cocok menggunakan pendekatan ini?")
    add_body_p(doc,
        "Pendekatan Prediktif (sering disebut sebagai Traditional Project Management atau siklus hidup Waterfall) "
        "adalah paradigma manajemen proyek berbasis rencana (plan-driven methodology) di mana ruang lingkup (scope), "
        "waktu (schedule), dan biaya (budget) ditentukan serta dirinci secara komprehensif pada fase-fase awal siklus hidup proyek "
        "(PMBOK® Guide 6th Edition, PMI, 2017, hal. 19; Larson & Gray, 2021). Karakteristik fundamental dari pendekatan ini adalah "
        "aliran eksekusi yang bersifat linear dan sekuensial (sequential linear flow), di mana suatu fase proyek (misalnya: Feasibility -> "
        "Requirements Analysis -> System Architecture & Design -> Implementation / Coding -> Verification & Testing -> Maintenance) "
        "harus diselesaikan dan disetujui secara formal melalui milestone review sebelum fase berikutnya dapat dimulai.")
    
    add_body_p(doc,
        "Dalam pendekatan prediktif, perubahan terhadap baseline ruang lingkup dipandang sebagai ancaman yang harus dikendalikan secara "
        "ketat melalui prosedur Change Control Board (CCB) formal guna mencegah terjadinya pembengkakan biaya dan jadwal (scope creep). "
        "Pengiriman produk atau hasil akhir proyek (deliverable delivery) umumnya dilakukan dalam satu pelepasan tunggal berskala besar "
        "(single big-bang release) pada akhir masa proyek.")

    add_body_p(doc,
        "Jenis proyek yang paling cocok (ideal) menggunakan pendekatan prediktif memiliki karakteristik sebagai berikut:")
    add_num_p(doc, "Persyaratan dan Kebutuhan yang Stabil (Clear & Stable Requirements): Proyek di mana spesifikasi kebutuhan teknis, operasional, dan fungsional sudah dipahami secara menyeluruh dan disepakati oleh semua pemangku kepentingan sejak hari pertama, serta probabilitas terjadinya perubahan di tengah jalan sangat rendah.", 1)
    add_num_p(doc, "Teknologi dan Solusi yang Telah Terbukti (Proven Technology): Proyek yang mengandalkan teknologi, metode arsitektur, dan perkakas yang telah mapan, sehingga tidak menuntut fase eksplorasi ilmiah atau eksperimentasi algoritma yang tidak pasti.", 2)
    add_num_p(doc, "Biaya Perubahan yang Sangat Mahal (High Cost of Change): Proyek di mana kesalahan desain yang baru terdeteksi pada fase konstruksi akhir akan menimbulkan kerugian finansial bencana atau membahayakan keselamatan jiwa. Contoh klasik mencakup proyek konstruksi gedung bertingkat, pembangunan jembatan bentang panjang, perakitan infrastruktur hardware data center, perakitan satelit kedirgantaraan, dan manufaktur sirkuit perangkat keras terintegrasi.", 3)
    add_num_p(doc, "Lingkungan Kepatuhan dan Regulasi Ketat (Heavy Compliance & Regulation): Proyek yang secara hukum mewajibkan dokumentasi audit menyeluruh, verifikasi bertahap, dan pengesahan otoritas resmi pada setiap gerbang fase (phase gates), seperti proyek pengembangan perangkat medis kelas III atau sistem kendali pembangkit listrik.", 4)

    # =========================================================================
    # SOAL 2
    # =========================================================================
    add_heading_2(doc, "2. Apa yang dimaksud dengan pendekatan Agile dalam manajemen proyek dan jenis proyek apa yang paling cocok menggunakan pendekatan ini?")
    add_body_p(doc,
        "Pendekatan Agile adalah paradigma manajemen proyek yang adaptif dan berorientasi pada nilai (value-driven and change-driven methodology) "
        "yang menekankan fleksibilitas, kolaborasi intensif, penyampaian hasil secara berulang (iterative) dan bertahap (incremental), "
        "serta kemampuan merespons perubahan kebutuhan pasar secara cepat dan elegan (Agile Practice Guide, PMI & Agile Alliance, 2017, hal. 13–18; "
        "Highsmith, 2009). Berakar dari Agile Manifesto (2001), pendekatan ini mendudukkan individu dan interaksi di atas proses dan perkakas, "
        "perangkat lunak yang berfungsi di atas dokumentasi komprehensif, kolaborasi pelanggan di atas negosiasi kontrak, serta respons terhadap perubahan di atas ketaatan pada rencana kaku.")

    add_body_p(doc,
        "Dalam metodologi Agile, siklus proyek dipecah menjadi rentang waktu pendek yang berulang (disebut Iterasi atau Sprint, biasanya 1 hingga 4 minggu). "
        "Pada setiap akhir iterasi, tim berkomitmen menghasilkan sebuah inkremen produk yang berpotensi dapat dirilis (potentially shippable increment) "
        "dan langsung diinspeksi oleh pemangku kepentingan. Umpan balik (feedback loops) yang diterima kemudian digunakan untuk menyempurnakan daftar "
        "prioritas pekerjaan (backlog) pada iterasi berikutnya.")

    add_body_p(doc,
        "Jenis proyek yang paling cocok (ideal) menggunakan pendekatan Agile meliputi:")
    add_num_p(doc, "Tingkat Ketidakpastian dan Kompleksitas Tinggi (High Uncertainty & Dynamic Domain): Proyek yang beroperasi di lingkungan VUCA (Volatility, Uncertainty, Complexity, Ambiguity), di mana solusi akhir belum dapat didefinisikan secara presisi di awal dan membutuhkan validasi hipotesis berkelanjutan.", 1)
    add_num_p(doc, "Persyaratan yang Berkembang Berkelanjutan (Evolving Requirements): Proyek inovasi digital di mana kebutuhan pengguna akhir, dinamika kompetitor bisnis, atau tren teknologi berkembang dengan sangat dinamis, sehingga rencana jangka panjang yang kaku akan cepat menjadi usang (obsolete).", 2)
    add_num_p(doc, "Pengembangan Produk Perangkat Lunak dan Digital (Software & SaaS Products): Proyek pembuatan aplikasi web, aplikasi mobile, platform e-commerce, dan sistem SaaS di mana biaya perubahan kode program relatif rendah dan kecepatan rilis ke pasar (time-to-market) merupakan keunggulan kompetitif utama.", 3)
    add_num_p(doc, "Riset Inovasi, Eksperimentasi Data Science, dan Machine Learning: Proyek penelitian analitik seperti proyek riset PRE-02, di mana performa model prediksi (akurasi, Brier score, kalibrasi) bergantung pada siklus coba-dan-evaluasi fitur, arsitektur model, dan tuning hyperparameter yang tidak dapat diprediksi secara linear.", 4)

    # =========================================================================
    # SOAL 3
    # =========================================================================
    add_heading_2(doc, "3. Sebutkan dan jelaskan tiga peran (roles) dalam Scrum beserta tanggung jawab masing-masing!")
    add_body_p(doc,
        "Scrum adalah kerangka kerja (framework) Agile yang paling luas diadopsi di industri rekayasa perangkat lunak modern. "
        "Berdasarkan The Scrum Guide (Schwaber & Sutherland, 2020), Scrum Team terdiri dari satu kesatuan unit profesional tanpa "
        "hierarki sub-tim, yang memegang tiga peran (akuntabilitas) spesifik yang saling melengkapi:")

    headers_q3 = ["Peran Scrum (Role)", "Fokus & Akuntabilitas Utama", "Rincian Tanggung Jawab Operasional"]
    data_q3 = [
        [
            "Product Owner (PO)",
            "Memaksimalkan Nilai Produk (Maximizing Product Value) dan Pengelolaan Product Backlog",
            "• Mendefinisikan visi produk dan merumuskan Product Goal yang terukur.\n"
            "• Mengembangkan, mengartikulasikan, dan memperjelas item-item Product Backlog (User Stories).\n"
            "• Memprioritaskan urutan backlog berdasarkan nilai bisnis, risiko, ketergantungan, dan ROI.\n"
            "• Menjembatani aspirasi bisnis pemangku kepentingan (stakeholders) dengan tim pengembang.\n"
            "• Memastikan Product Backlog bersifat transparan, dapat dipahami, dan selalu mutakhir."
        ],
        [
            "Scrum Master (SM)",
            "Efektivitas Scrum Team dan Penerapan Prinsip Scrum (Servant Leader & Coach)",
            "• Memastikan Scrum Team dan organisasi memahami teori serta praktik empiris Scrum.\n"
            "• Memfasilitasi seluruh even Scrum agar berlangsung produktif dan tetap dalam batas waktu (time-boxed).\n"
            "• Bertindak sebagai penyingkir hambatan (impediment remover) yang menghalangi kemajuan tim pengembang.\n"
            "• Melindungi tim dari gangguan dan intervensi eksternal yang tidak selaras dengan Sprint Goal.\n"
            "• Melatih tim menuju kemandirian (self-management) dan pembentukan budaya perbaikan mutu berkelanjutan."
        ],
        [
            "Developers / Development Team",
            "Menciptakan Usable Increment pada Setiap Akhir Sprint (Engineering Excellence)",
            "• Menyusun rencana teknis harian dan Sprint Backlog untuk mencapai Sprint Goal.\n"
            "• Melakukan estimasi ukuran pekerjaan teknis (story points / jam kerja) secara objektif.\n"
            "• Menanamkan standar kualitas tinggi dengan mematuhi Definition of Done (DoD) pada setiap baris kode/fitur.\n"
            "• Bersifat lintas fungsi (cross-functional) mencakup perancangan, penulisan kode, pengetesan (QA), dan deployment.\n"
            "• Berkolaborasi dan beradaptasi setiap hari melalui pertemuan Daily Scrum."
        ]
    ]
    create_table(doc, headers_q3, data_q3)

    # =========================================================================
    # SOAL 4
    # =========================================================================
    add_heading_2(doc, "4. Sebutkan dan jelaskan lima even (events) dalam Scrum beserta tujuan masing-masing!")
    add_body_p(doc,
        "Scrum menetapkan lima even resmi yang dirancang secara khusus untuk menciptakan keteraturan, menyediakan forum inspeksi "
        "dan adaptasi yang kritis, serta meminimalkan kebutuhan akan rapat-rapat ad-hoc yang tidak terdefinisi (Schwaber & Sutherland, 2020). "
        "Seluruh even ini memiliki batasan waktu ketat (time-boxed):")

    headers_q4 = ["Nama Even Scrum", "Waktu Pelaksanaan & Time-Box", "Tujuan Utama & Mekanisme Kerja"]
    data_q4 = [
        [
            "1. The Sprint",
            "Wadah waktu tetap: 1–4 minggu (konsisten sepanjang proyek)",
            "Merupakan jantung dari irama Scrum; sebuah wadah (container) di mana seluruh even lainnya bernaung. Tujuannya adalah mengubah ide dan backlog menjadi inkremen produk bernilai nyata secara konsisten tanpa ada jeda antar-sprint."
        ],
        [
            "2. Sprint Planning",
            "Awal setiap Sprint\nMaksimal 8 jam (untuk Sprint 1 bulan) / ~2 jam per minggu",
            "Menetapkan arah dan komitmen Sprint dengan menjawab 3 pertanyaan mendasar:\n(1) Mengapa Sprint ini bernilai? (Merumuskan Sprint Goal)\n(2) Apa yang dapat diselesaikan? (Pemilihan item Product Backlog)\n(3) Bagaimana pekerjaan akan diselesaikan? (Penyusunan tugas teknis ke dalam Sprint Backlog)."
        ],
        [
            "3. Daily Scrum",
            "Setiap hari kerja\nMaksimal 15 menit (berdiri / stand-up)",
            "Menginspeksi kemajuan pekerjaan menuju tercapainya Sprint Goal dan mengidentifikasi hambatan (blockers). Tim menyinkronkan aktivitas 24 jam terakhir, rencana 24 jam ke depan, serta menyesuaikan rencana kerja jangka pendek."
        ],
        [
            "4. Sprint Review",
            "Menjelang akhir Sprint\nMaksimal 4 jam (untuk Sprint 1 bulan) / ~1 jam per minggu",
            "Menginspeksi hasil akhir Increment yang telah berstatus 'Done' bersama para pemangku kepentingan (stakeholders). Tim mendemonstrasikan fitur baru yang berfungsi, mengumpulkan umpan balik langsung, dan mendiskusikan penyesuaian Product Backlog berikutnya."
        ],
        [
            "5. Sprint Retrospective",
            "Penutup Sprint (setelah Review)\nMaksimal 3 jam (untuk Sprint 1 bulan) / ~45 mnt per minggu",
            "Melakukan introspeksi internal Scrum Team mengenai efektivitas proses kerja, relasi antar-individu, perkakas (tools), dan Definition of Done. Tujuannya adalah merumuskan rencana tindakan perbaikan nyata (actionable improvements) untuk diterapkan pada Sprint berikutnya."
        ]
    ]
    create_table(doc, headers_q4, data_q4)

    # =========================================================================
    # SOAL 5
    # =========================================================================
    add_heading_2(doc, "5. Sebutkan dan jelaskan tiga artefak (artifacts) dalam Scrum beserta fungsinya!")
    add_body_p(doc,
        "Artefak dalam Scrum merepresentasikan pekerjaan atau nilai yang dirancang untuk memaksimalkan transparansi informasi kunci, "
        "sehingga setiap pihak yang menginspeksinya memiliki basis pemahaman yang sama. Berdasarkan The Scrum Guide (Schwaber & Sutherland, 2020), "
        "setiap artefak memuat komitmen (commitment) formal guna memperkuat fokus dan akuntabilitas tim:")

    add_num_p(doc, "**Product Backlog (Daftar Kebutuhan Produk):** "
        "Merupakan daftar inventaris terurut dan dinamis dari setiap fitur, fungsi, persyaratan teknis, perbaikan bug, dan peningkatan non-fungsional "
        "yang diperlukan untuk menyempurnakan produk. Artefak ini bersifat terus berkembang (living artifact) dan dikelola secara eksklusif oleh Product Owner. "
        "Komitmen yang melekat pada Product Backlog adalah **Product Goal**, yaitu target jangka panjang yang menggambarkan keadaan masa depan produk sebagai sasaran besar yang ingin dicapai oleh Scrum Team.", 1)

    add_num_p(doc, "**Sprint Backlog (Daftar Pekerjaan Sprint):** "
        "Merupakan kumpulan item Product Backlog terpilih yang menjadi target penyelesaian pada Sprint berjalan, dilengkapi dengan rincian rencana "
        "tugas teknis operasional untuk menghasilkan Increment. Artefak ini dimiliki dan diperbarui secara real-time oleh Developers sepanjang Sprint. "
        "Komitmen yang melekat pada Sprint Backlog adalah **Sprint Goal**, yaitu satu tujuan tunggal yang koheren yang memberikan kejelasan fokus bagi Developers mengenai nilai bisnis yang harus dicapai dalam Sprint tersebut.", 2)

    add_num_p(doc, "**Increment (Inkremen Produk Nyata):** "
        "Merupakan langkah konkret kumulatif menuju tercapainya Product Goal. Increment adalah jumlah keseluruhan dari semua item Product Backlog "
        "yang berhasil diselesaikan sepanjang Sprint berjalan dan seluruh Sprint sebelumnya, yang berada dalam kondisi terintegrasi, teruji, dan dapat digunakan secara nyata (usable and releasable). "
        "Komitmen yang melekat pada Increment adalah **Definition of Done (DoD)**, yaitu deskripsi formal tentang kondisi mutu teknis yang harus dipenuhi agar suatu item backlog diakui sah telah selesai tanpa ada utang teknis (technical debt).", 3)

    # =========================================================================
    # SOAL 6
    # =========================================================================
    add_heading_2(doc, "6. Apa yang dimaksud dengan pendekatan Hybrid dalam manajemen proyek dan mengapa pendekatan ini semakin banyak digunakan?")
    add_body_p(doc,
        "Pendekatan Hybrid adalah metodologi manajemen proyek terpadu yang memadukan elemen-elemen dari pendekatan prediktif (plan-driven / Waterfall) "
        "dan pendekatan adaptif (Agile / Scrum / Kanban) dalam satu siklus hidup proyek yang disesuaikan secara proporsional dengan karakteristik "
        "dan konteks spesifik organisasi (PMBOK® Guide 7th Edition, PMI, 2021; Agile Practice Guide, PMI & Agile Alliance, 2017). "
        "Pendekatan ini tidak memaksakan dogma metodologi tunggal (bukan dogmatic purism), melainkan menerapkan prinsip *tailoring* "
        "untuk memilih teknik terbaik yang mampu menghasilkan nilai bisnis maksimal dengan risiko terkendali.")

    add_body_p(doc,
        "Faktor-faktor pendorong mengapa pendekatan Hybrid semakin mendominasi lanskap industri modern mencakup:")
    add_bullet_p(doc, "Kebutuhan Menjembatani Budaya Tata Kelola Korporat dengan Fleksibilitas Eksekusi: Banyak organisasi mapan memiliki struktur tata kelola keuangan, audit, dan kepatuhan hukum yang menuntut estimasi biaya total (BAC) dan batas waktu makro di awal (elemen prediktif), namun tim pengembang memerlukan kelincahan iteratif untuk merespons dinamika teknis (elemen Agile).")
    add_bullet_p(doc, "Mengakomodasi Proyek Multidisiplin (Hardware vs Software): Proyek modern sering menggabungkan komponen fisik (seperti pengadaan infrastruktur server, sensor IoT, atau konstruksi) yang berbiaya perubahan tinggi dan harus direncanakan secara Waterfall, dengan komponen analitik perangkat lunak atau model kecerdasan buatan yang menuntut eksperimentasi bertahap.")
    add_bullet_p(doc, "Mitigasi Risiko yang Lebih Berimbang: Hybrid memungkinkan proyek mengunci arsitektur inti dan mitigasi regulasi di awal proyek untuk menghindari risiko fatalitas hukum, sekaligus menjaga ruang inovasi antarmuka dan logika bisnis agar tetap adaptif terhadap kebutuhan pemangku kepentingan.")
    add_bullet_p(doc, "Dukungan Kerangka Kerja Global: Project Management Institute (PMI) mengadopsi Disciplined Agile (DA) dan Agile-Predictive Continuum untuk membimbing praktisi menyusun formula metodologi hibrida sesuai dengan spektrum Stacey Complexity Model.")

    # =========================================================================
    # SOAL 7
    # =========================================================================
    add_heading_2(doc, "7. Berikan contoh konkret penerapan pendekatan Hybrid dari referensi yang Anda baca!")
    add_body_p(doc,
        "Dalam literatur manajemen proyek kontemporer (PMI, 2017, hal. 26–28; Boehm & Turner, 2004), terdapat beberapa pola arsitektur "
        "hibrida yang telah terbukti secara empiris di industri global:")

    add_num_p(doc, "**Model 'Water-Scrum-Fall' pada Sistem Finansial / Perbankan:** "
        "Pada fase awal (Water), manajemen menggunakan pendekatan Waterfall untuk studi kelayakan bisnis, penetapan anggaran tahunan tetap, "
        "analisis regulasi kepatuhan perbankan (OJK/BI), serta perancangan arsitektur keamanan tingkat tinggi. Setelah cetak biru disahkan, "
        "fase implementasi diserahkan kepada tim pengembang perangkat lunak yang bekerja menggunakan kerangka kerja Scrum (Scrum) "
        "selama beberapa siklus Sprint dua mingguan untuk membangun modul-modul transaksi, antarmuka nasabah, dan API microservices. "
        "Terakhir, fase rilis produksi (Fall) kembali menggunakan pendekatan prediktif formal yang mencakup audit kepatuhan, uji penetrasi keamanan (pentest), "
        "pelatihan staf cabang, dan penggelaran sistem secara bertahap.", 1)

    add_num_p(doc, "**Pengembangan Perangkat Medis Cerdas Berbasis AI (Internet of Medical Things):** "
        "Pabrikasi perangkat keras sensor diagnostik dan verifikasi standar keselamatan hayati (ISO 13485 / regulasi FDA) dikerjakan dengan model linier prediktif "
        "karena pembuatan cetakan perangkat keras tidak dapat diulang sembarangan. Di sisi lain, modul algoritma deep learning penafsir sinyal medis dikembangkan "
        "secara iteratif/Agile dengan melakukan pengujian model, retraining data klinis, dan penyempurnaan akurasi secara berkelanjutan.", 2)

    add_num_p(doc, "**Konteks Nyata pada Proyek Riset PRE-02:** "
        "Pada riset sistem prediksi keterlambatan proyek PRE-02, tim menerapkan pendekatan Hybrid yang sangat terstruktur: "
        "(1) Kerangka makro proyek, penjadwalan checkpoint akademik, perancangan WBS, serta integrasi variabel basis data ERP Farm Nation Enterprise "
        "dikunci secara prediktif untuk menjamin keselarasan dengan silabus perkuliahan; "
        "(2) Proses sains data—mencakup seleksi 7 fitur proyek, imputasi data, hyperparameter tuning Random Forest/XGBoost/MLP, serta kalibrasi probabilitas kontinu—dikerjakan "
        "secara adaptif melalui sprint eksperimen mingguan terukur yang dievaluasi menggunakan metrik Brier Score dan Expected Calibration Error (ECE).", 3)

    # =========================================================================
    # SOAL 8
    # =========================================================================
    add_heading_2(doc, "8. Sebutkan kembali sepuluh area pengetahuan manajemen proyek dan jelaskan secara singkat fokus utama masing-masing!")
    add_body_p(doc,
        "Berdasarkan standar PMBOK® Guide Edisi ke-6 (PMI, 2017), terdapat sepuluh Area Pengetahuan (Knowledge Areas) "
        "yang merepresentasikan kumpulan konsep, istilah, dan aktivitas profesional yang mencakup seluruh aspek tata kelola proyek:")

    headers_q8 = ["No", "Area Pengetahuan (Knowledge Area)", "Fokus Utama & Sasaran Pengelolaan Proyek"]
    data_q8 = [
        ["1", "Project Integration Management\n(Manajemen Integrasi)", "Mengidentifikasi, mendefinisikan, mengombinasikan, menyatukan, dan mengoordinasikan seluruh proses serta aktivitas manajemen proyek dari pembuatan Project Charter hingga penutupan formal, termasuk pengendalian perubahan terpadu (Integrated Change Control)."],
        ["2", "Project Scope Management\n(Manajemen Ruang Lingkup)", "Memastikan proyek mencakup seluruh pekerjaan yang diperlukan, dan HANYA pekerjaan yang diperlukan, untuk menyelesaikan proyek dengan sukses melalui pengumpulan persyaratan, pendefinisian ruang lingkup, penyusunan WBS, validasi, dan pengendalian ruang lingkup."],
        ["3", "Project Schedule Management\n(Manajemen Jadwal)", "Mengelola penyelesaian proyek tepat waktu, mencakup pendefinisian aktivitas, pengurutan ketergantungan tugas (sequencing), estimasi durasi, pengembangan diagram jaringan (PDM) & jalur kritis (CPM), serta pengendalian jadwal baseline."],
        ["4", "Project Cost Management\n(Manajemen Biaya)", "Merencanakan, mengestimasi, menganggarkan, mendanai, dan mengendalikan pengeluaran biaya agar proyek selesai dalam batas anggaran yang disetujui (BAC), serta memantau kinerja finansial menggunakan Earned Value Management (EVM)."],
        ["5", "Project Quality Management\n(Manajemen Kualitas)", "Menggabungkan kebijakan mutu organisasi ke dalam perencanaan, pelaksanaan, dan pengendalian mutu proyek dan produk guna memenuhi ekspektasi pemangku kepentingan melalui penjaminan mutu (QA) dan pengendalian mutu (QC)."],
        ["6", "Project Resource Management\n(Manajemen Sumber Daya)", "Mengidentifikasi, mengakuisisi, dan mengelola sumber daya manusia (tim proyek) dan sumber daya fisik (peralatan, material, fasilitas, server) yang diperlukan untuk keberhasilan penyelesaian proyek secara optimal."],
        ["7", "Project Communications Management\n(Manajemen Komunikasi)", "Memastikan perencanaan, pengumpulan, pembuatan, distribusi, penyimpanan, pengambilan, dan pemantauan informasi proyek dilakukan secara tepat waktu, akurat, dan efektif bagi seluruh pemangku kepentingan."],
        ["8", "Project Risk Management\n(Manajemen Risiko)", "Memaksimalkan peluang positif dan meminimalkan dampak ancaman negatif melalui identifikasi risiko, analisis kualitatif (matriks P&I), analisis kuantitatif (EMV/Decision Tree), perencanaan respons mitigasi, dan pemantauan risiko."],
        ["9", "Project Procurement Management\n(Manajemen Pengadaan)", "Membeli atau mengakuisisi produk, jasa, atau hasil kerja dari luar tim proyek melalui perencanaan kontrak, seleksi penjual, administrasi perjanjian hukum, dan pengendalian kinerja vendor eksternal."],
        ["10", "Project Stakeholder Management\n(Manajemen Pemangku Kepentingan)", "Mengidentifikasi individu, kelompok, atau organisasi yang memengaruhi atau dipengaruhi oleh proyek, menganalisis ekspektasi mereka (Power/Interest Grid), serta merancang strategi pelibatan aktif dalam pengambilan keputusan."]
    ]
    create_table(doc, headers_q8, data_q8, [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.JUSTIFY])

    # =========================================================================
    # SOAL 9
    # =========================================================================
    add_heading_2(doc, "9. Menurut Anda, bagaimana penerapan sepuluh area pengetahuan berbeda dalam proyek dengan pendekatan Agile dibandingkan dengan pendekatan prediktif?")
    add_body_p(doc,
        "Meskipun kesepuluh area pengetahuan manajemen proyek tetap relevan di seluruh tipe proyek, filosofi eksekusi, "
        "keterlibatan tim, serta format artefaknya bertransformasi secara radikal saat berpindah dari paradigma prediktif (plan-driven) "
        "menuju paradigma Agile (value-driven) (PMBOK® Guide 6th Edition, PMI, 2017; Agile Practice Guide, 2017). "
        "Tabel berikut menguraikan komparasi kritis transformasi kesepuluh area pengetahuan tersebut:")

    headers_q9 = ["Area Pengetahuan", "Penerapan pada Pendekatan Prediktif (Waterfall)", "Penerapan pada Pendekatan Adaptif (Agile / Scrum)"]
    data_q9 = [
        [
            "1. Integrasi",
            "Dikoordinasikan secara terpusat oleh Project Manager melalui Project Management Plan formal tebal dan proses CCB yang ketat.",
            "Diintegrasikan secara terdesentralisasi oleh seluruh tim; koordinasi harian via Daily Scrum dan integrasi produk berkelanjutan (CI/CD)."
        ],
        [
            "2. Ruang Lingkup",
            "Didefinisikan secara fixed di awal via Scope Baseline dan WBS berjenjang; perubahan scope dianggap deviasi negatif yang dicegah.",
            "Bersifat fleksibel dan berkembang via Product Backlog dinamis; perubahan diprioritaskan ulang setiap Sprint berdasarkan nilai bisnis tertinggi."
        ],
        [
            "3. Jadwal",
            "Berdasarkan lintasan kritis (CPM), ketergantungan PDM, durasi hari kalender, dan Gantt Chart yang dipantau terhadap Schedule Baseline.",
            "Berdasarkan time-box tetap (Sprint 1–4 minggu), kapasitas kecepatan tim (velocity / story points), dan pelacakan via Burn-down/Burn-up Charts."
        ],
        [
            "4. Biaya",
            "Anggaran fixed berbasis baseline terinci (BAC, PV, EV, AC); variansi biaya dikendalikan melalui analisis matematika EVM kuantitatif.",
            "Biaya per sprint relatif stabil (kapasitas tim tetap); pengeluaran dikendalikan melalui penetapan prioritas fitur bernilai tinggi dalam batas anggaran."
        ],
        [
            "5. Kualitas",
            "Fokus pada inspeksi formal (QC) di akhir fase implementasi terhadap dokumen spesifikasi teknis dan audit kualitas terencana.",
            "Kualitas ditanamkan sejak hari pertama (built-in quality) melalui kepatuhan Definition of Done (DoD), pair programming, dan Automated Testing."
        ],
        [
            "6. Sumber Daya",
            "Struktur tim hierarkis dengan spesialisasi tugas terpisah; penugasan kerja dialokasikan oleh manajer proyek.",
            "Tim lintas fungsi (cross-functional) yang mengelola diri sendiri (self-organizing); kolaborasi kolektif dan kepemilikan bersama (shared ownership)."
        ],
        [
            "7. Komunikasi",
            "Komunikasi formal terdokumentasi (status report berkala, memo, rapat formal); saluran komunikasi dihitung ketat n(n-1)/2.",
            "Komunikasi informal, tatap muka, transparan, dan real-time via papan visual (Kanban board), radiasi informasi, dan pertemuan Scrum harian."
        ],
        [
            "8. Risiko",
            "Identifikasi dan analisis mendalam di awal proyek via Risk Register formal; cadangan kontinjensi dihitung matematis (EMV/Monte Carlo).",
            "Risiko dikelola secara iteratif setiap sprint; tim membuat 'Spike' (riset eksperimen singkat) untuk mereduksi ketidakpastian teknis secara cepat."
        ],
        [
            "9. Pengadaan",
            "Kontrak formal berorientasi kepatuhan ketat (Firm Fixed-Price); negosiasi panjang dan evaluasi vendor berlapis.",
            "Kontrak kolaboratif fleksibel (seperti agile contracting, time-and-materials bergradasi, atau target cost sharing) yang mendukung perubahan iteratif."
        ],
        [
            "10. Stakeholder",
            "Keterlibatan diatur formal sesuai Stakeholder Engagement Plan; keterlibatan puncak terjadi pada fase inisiasi dan fase serah terima akhir.",
            "Pelibatan intensif dan berkelanjutan sepanjang proyek; Product Owner dan pemangku kepentingan hadir langsung dalam setiap Sprint Review."
        ]
    ]
    create_table(doc, headers_q9, data_q9, [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
