# -*- coding: utf-8 -*-
"""
Section B Builder: Tugas Individu - Perbandingan Pendekatan Prediktif dan Agile
Pertemuan 14: Pendekatan Agile/Hybrid & Review Perkuliahan
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_b(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "B. TUGAS INDIVIDU: PERBANDINGAN PENDEKATAN PREDIKTIF DAN AGILE")

    add_body_p(doc,
        "Bagian ini menyajikan analisis komparatif mendalam antara dua paradigma dominan dalam manajemen proyek modern, "
        "yaitu Pendekatan Prediktif (Traditional/Waterfall) dan Pendekatan Adaptif (Agile), serta sintesisnya ke dalam "
        "Pendekatan Hibrida (Hybrid Approach). Analisis ini mencakup matriks perbandingan multidimensi, elaborasi teoritis "
        "tiap aspek komparasi, serta studi kasus penerapan arsitektur metodologi hibrida pada proyek riset PRE-02 "
        "(Sistem Prediksi Probabilitas Keterlambatan Proyek Berbasis Machine Learning sebagai Early Warning System).")

    # =========================================================================
    # 1. MATRIKS PERBANDINGAN KOMPREHENSIF
    # =========================================================================
    add_heading_2(doc, "1. Matriks Komparatif Pendekatan Prediktif vs Agile vs Hybrid")
    add_body_p(doc,
        "Tabel berikut menyajikan matriks perbandingan komprehensif yang memetakan seluruh dimensi krusial manajemen proyek "
        "berdasarkan standar PMBOK® Guide Edisi ke-6 & ke-7 (PMI, 2017, 2021) serta Agile Practice Guide (PMI & Agile Alliance, 2017):")

    headers_matrix = ["Dimensi Perbandingan", "Pendekatan Prediktif (Waterfall)", "Pendekatan Adaptif (Agile)", "Pendekatan Hibrida (Hybrid)"]
    data_matrix = [
        [
            "1. Karakteristik Persyaratan / Ruang Lingkup",
            "Stabil, terdefinisi lengkap, dan dikunci (frozen) di awal proyek; deviasi diminimalkan secara ketat.",
            "Dinamis, fleksibel, berkembang (emergent) seiring waktu melalui umpan balik pengguna dan pasar.",
            "Arsitektur makro dan regulasi dikunci di awal (prediktif); fitur fungsional dan UI berevolusi secara iteratif (agile)."
        ],
        [
            "2. Kolaborasi & Struktur Tim Proyek",
            "Hierarkis dengan peran fungsional terpisah (analis, desainer, koder, tester); penugasan oleh PM.",
            "Lintas fungsi (cross-functional), kolaboratif intensif, mandiri (self-organizing & self-managing).",
            "Manajemen proyek formal berkoordinasi dengan kepemimpinan organisasi; tim teknis bekerja mandiri dalam sprint."
        ],
        [
            "3. Siklus Pengiriman Hasil (Delivery)",
            "Tunggal di akhir proyek (single big-bang delivery) setelah seluruh fase sekuensial tuntas.",
            "Berkala, sering, dan bertahap (incremental delivery) dalam bentuk working increment setiap 1–4 minggu.",
            "Pengiriman rilis bertahap (iterative releases) dengan tetap menghormati tonggak capaian (milestone gates) utama."
        ],
        [
            "4. Manajemen Risiko Proyek",
            "Analisis risiko mendalam di awal via Risk Register; mitigasi berbasis kontinjensi matematis.",
            "Mitigasi risiko berkelanjutan dan empiris melalui inspeksi produk nyata setiap sprint dan teknik Spike.",
            "Risiko kepatuhan dan arsitektur dianalisis di awal; risiko teknis dan performa diuji melalui sprint eksperimen."
        ],
        [
            "5. Pelibatan Pelanggan / Stakeholder",
            "Terkonsentrasi di awal (pengumpulan kebutuhan) dan di akhir proyek (User Acceptance Test / serah terima).",
            "Keterlibatan aktif dan berkelanjutan sepanjang siklus proyek melalui kehadiran rutin di Sprint Review.",
            "Pelibatan terstruktur melalui rapat steering committee pada milestone gerbang, dan review langsung pada akhir siklus sprint."
        ],
        [
            "6. Pengendalian Perubahan (Change Control)",
            "Perubahan dipandang sebagai ancaman; memerlukan prosedur formal persetujuan Change Control Board (CCB).",
            "Perubahan disambut baik (embracing change) sebagai peluang menambah nilai; diprioritaskan di Product Backlog.",
            "Perubahan level strategis/anggaran melalui CCB; perubahan taktis fitur/algoritma dikelola fleksibel oleh Product Owner."
        ],
        [
            "7. Metrik Keberhasilan Proyek",
            "Kepatuhan pada Segitiga Emas (Iron Triangle): On-Time, On-Budget, dan On-Scope sesuai baseline awal.",
            "Penyampaian nilai bisnis nyata (business value delivered), kepuasan pengguna, dan kegunaan perangkat lunak.",
            "Keseimbangan antara kepatuhan anggaran/jadwal korporat dan pencapaian nilai inovasi yang memuaskan pengguna."
        ],
        [
            "8. Dokumentasi & Aliran Komunikasi",
            "Komprehensif, formal, berbasis dokumen spesifikasi tebal (SRS, SDD, Project Management Plan).",
            "Dokumentasi secukupnya (just-enough documentation), komunikasi tatap muka, dan radiasi informasi visual (Kanban).",
            "Dokumentasi arsitektur dan kepatuhan tetap formal; dokumentasi teknis harian dikelola ringkas via repositori kode."
        ],
        [
            "9. Jenis Proyek yang Paling Sesuai",
            "Konstruksi sipil, infrastruktur fisik, pengadaan perangkat keras, proyek kepatuhan regulasi ketat.",
            "Pengembangan aplikasi mobile, startup teknologi, software SaaS, eksplorasi desain, penelitian AI/ML murni.",
            "Integrasi sistem ERP enterprise, transformasi digital perbankan, riset sains data terapan, proyek PRE-02."
        ]
    ]
    create_table(doc, headers_matrix, data_matrix, [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    # =========================================================================
    # 2. ANALISIS MENDALAM DIMENSI PERBANDINGAN
    # =========================================================================
    add_heading_2(doc, "2. Analisis Mendalam Karakteristik dan Filosofi Masing-Masing Pendekatan")
    add_body_p(doc,
        "Perbedaan antara pendekatan prediktif dan adaptif bukan sekadar perbedaan teknis operasional, melainkan perbedaan "
        "paradigma epistemologis dalam memandang ketidakpastian (uncertainty) dan nilai (value) (Boehm & Turner, 2004; Highsmith, 2009):")

    add_heading_3(doc, "a. Paradigma Manajemen Ruang Lingkup dan Ketidakpastian")
    add_body_p(doc,
        "Pendekatan prediktif berpijak pada asumsi bahwa ketidakpastian dapat direduksi secara signifikan melalui perencanaan mendalam "
        "di awal proyek (upfront planning). Dengan menganalisis kebutuhan secara mendalam, manajer proyek percaya bahwa ruang lingkup "
        "dapat dikunci sehingga proyek dapat dieksekusi dengan variansi minimal. Namun, kelemahan mendasar model ini terletak pada fenomena "
        "'the cost of delay' dan keusangan spesifikasi ketika kondisi bisnis berubah sebelum produk diluncurkan. Sebaliknya, pendekatan Agile "
        "berpijak pada prinsip empirisme (transparansi, inspeksi, dan adaptasi). Agile menerima kenyataan bahwa pada sistem perangkat lunak "
        "yang kompleks, pengguna sering kali baru mengetahui apa yang benar-benar mereka butuhkan setelah melihat dan mencoba langsung produk yang berfungsi.")

    add_heading_3(doc, "b. Dinamika Struktur Organisasi dan Psikologi Tim")
    add_body_p(doc,
        "Dalam pendekatan prediktif, struktur tim bersifat komando-dan-kendali (command-and-control). Manajer proyek bertindak sebagai "
        "pengambil keputusan utama yang mendelegasikan tugas terpisah kepada spesialis (misal programmer hanya menulis kode tanpa memikirkan "
        "dampak bisnis langsung). Sebaliknya, Agile memberdayakan tim lintas fungsi (cross-functional teams) yang mandiri (self-managing). "
        "Anggota tim memiliki rasa kepemilikan kolektif (collective code ownership) terhadap seluruh produk. Hal ini meningkatkan motivasi intrinsik, "
        "mempercepat pengambilan keputusan teknis, dan mengikis hambatan silang-departemen (silo mentality).")

    add_heading_3(doc, "c. Manajemen Risiko: Analisis Teoretis vs Validasi Empiris")
    add_body_p(doc,
        "Manajemen risiko prediktif mengandalkan analisis probabilitas dan dampak di atas kertas (Risk Register, simulasi Monte Carlo). "
        "Meskipun sangat berguna untuk risiko eksternal dan finansial, pendekatan ini rentan terhadap risiko integrasi tersembunyi "
        "(hidden technical integration risk) yang baru meledak pada saat pengujian sistem di akhir proyek. Di sisi lain, Agile mengelola "
        "risiko teknis secara langsung melalui rilis bertahap. Jika suatu modul memiliki latensi tinggi atau bug arsitektur, masalah tersebut "
        "langsung terungkap pada akhir Sprint pertama atau kedua, sehingga biaya perbaikannya masih sangat murah (low cost of refactoring).")

    # =========================================================================
    # 3. STUDI KASUS KONTEKSTUAL: PENERAPAN PADA PROYEK RISET PRE-02
    # =========================================================================
    add_heading_2(doc, "3. Studi Kasus Penerapan Metodologi pada Proyek Riset PRE-02")
    add_body_p(doc,
        "Proyek riset kelompok PRE-02 mengusung judul: 'Sistem Prediksi Probabilitas Keterlambatan Proyek Berbasis Machine Learning "
        "sebagai Early Warning System Manajemen Proyek'. Proyek ini bertujuan menghasilkan modul kecerdasan buatan yang mampu "
        "mengeluarkan nilai probabilitas numerik kontinu [0.0 - 1.0] yang terkalibrasi secara presisi, terintegrasi dengan modul "
        "monitoring Super ERP Farm Nation Enterprise.")

    add_body_p(doc,
        "Berdasarkan analisis karakteristik proyek PRE-02, penerapan dogma metodologi tunggal (murni Waterfall atau murni Agile) "
        "akan menimbulkan kegagalan struktural:")
    add_bullet_p(doc, "Jika Menggunakan Murni Waterfall: Riset machine learning tidak dapat diprediksi secara kaku. Eksperimen algoritma (seperti regresi logistik, Random Forest, XGBoost, dan neural network MLP) membutuhkan iterasi berulang pada rekayasa fitur (feature engineering) dan penalaan hyperparameter. Jika seluruh tahapan dikunci secara linear, tim tidak akan memiliki fleksibilitas untuk memperbaiki model saat ditemukan masalah underfitting atau Expected Calibration Error (ECE) yang tinggi.")
    add_bullet_p(doc, "Jika Menggunakan Murni Agile (Scrum Ekstrem): Proyek PRE-02 terikat pada batas waktu semester akademik yang ketat (14 pertemuan, jadwal UAS fixed), kewajiban penyusunan dokumen proposal terstruktur (Project Charter, WBS, PDM, EVM, Risk Register), serta integrasi skema basis data ERP Farm Nation Enterprise yang memiliki dependensi antarmuka tetap. Fleksibilitas tanpa batas akan merusak kepatuhan silabus perkuliahan.")

    add_body_p(doc,
        "Oleh karena itu, pilihan paling optimal, rasional, dan ilmiah bagi proyek PRE-02 adalah **Pendekatan Hibrida (Hybrid Approach)** "
        "dengan arsitektur *Disciplined Agile / Water-Scrum-Fall*, yang dipetakan sebagai berikut:")

    headers_case = ["Fase Proyek PRE-02", "Pendekatan yang Diterapkan", "Mekanisme Kerja & Deliverables yang Dihasilkan"]
    data_case = [
        [
            "Fase 1: Inisiasi & Perencanaan Arsitektur\n(Pertemuan 1 – 4)",
            "Prediktif (Waterfall)\nPlan-Driven Governance",
            "• Penyusunan Project Charter dan analisis kelayakan ilmiah.\n"
            "• Pemetaan 7 variabel fitur kondisi proyek dari ERP Farm Nation Enterprise.\n"
            "• Pembagian WBS hierarkis 4 tingkat dan penetapan baseline ruang lingkup akademik.\n"
            "• Penjadwalan milestone mingguan yang terikat pada kalender semester."
        ],
        [
            "Fase 2: Eksperimentasi & Pemodelan Machine Learning\n(Pertemuan 5 – 12)",
            "Adaptif (Agile / Scrum)\nIterative Sprints",
            "• Eksekusi siklus Sprint 1–2 mingguan untuk eksperimentasi model data science.\n"
            "• Iterasi 1: Baseline Logistic Regression & imputasi missing data.\n"
            "• Iterasi 2: Eksperimen model ensemble Random Forest & XGBoost dengan predict_proba.\n"
            "• Iterasi 3: Eksperimen Deep Learning Neural Network Multi-Layer Perceptron (MLP).\n"
            "• Iterasi 4: Kalibrasi probabilitas menggunakan Platt Scaling & Isotonic Regression untuk mereduksi Brier Score (< 0.18).\n"
            "• Daily stand-up dan evaluasi performa model di setiap akhir sprint."
        ],
        [
            "Fase 3: Integrasi API, Pengujian, & Penutupan\n(Pertemuan 13 – 14)",
            "Hibrida Terpadu\nGovernance & Deployment",
            "• Pembungkusan model terbaik ke dalam REST API endpoint (FastAPI/Flask).\n"
            "• Pengujian integrasi fungsional terhadap modul dashboard ERP.\n"
            "• Evaluasi ambang batas peringatan dini (threshold alert default >= 0.65).\n"
            "• Kompilasi dokumen proposal akhir berjenjang (Checkpoint #1–#6) secara terstruktur untuk evaluasi formal UAS."
        ]
    ]
    create_table(doc, headers_case, data_case, [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    add_body_p(doc,
        "Penerapan pendekatan Hibrida ini terbukti memberikan manfaat ganda bagi tim PRE-02: menjamin kepatuhan akademik 100% "
        "terhadap timeline institusi (aspek prediktif), sekaligus memberikan kebebasan eksplorasi ilmiah bagi tim pengembang untuk menghasilkan "
        "model machine learning dengan akurasi diskriminasi (AUC-ROC >= 0.80) dan kalibrasi probabilitas terbaik (aspek Agile).")
