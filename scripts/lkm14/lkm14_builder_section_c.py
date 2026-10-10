# -*- coding: utf-8 -*-
"""
Section C Builder: Refleksi Mandiri (Akhir Semester)
Pertemuan 14: Pendekatan Agile/Hybrid & Review Perkuliahan
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_c(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']

    add_heading_1(doc, "C. REFLEKSI MANDIRI (AKHIR SEMESTER)")

    add_body_p(doc,
        "Bagian ini menyajikan evaluasi personal, introspeksi kritis, dan kristalisasi pemahaman mahasiswa "
        "setelah menuntaskan seluruh rangkaian perkuliahan Teori Manajemen Proyek selama satu semester penuh (Pertemuan 1 s.d. 14), "
        "khususnya dalam merefleksikan dinamika pengerjaan proyek riset kelompok PRE-02 (Sistem Prediksi Probabilitas Keterlambatan "
        "Proyek Berbasis Machine Learning sebagai Early Warning System):")

    # =========================================================================
    # REFLEKSI 1
    # =========================================================================
    add_heading_2(doc, "1. Materi manajemen proyek apa yang paling bermanfaat bagi Anda selama perkuliahan ini, dan mengapa?")
    add_body_p(doc,
        "Dari sepuluh area pengetahuan yang dipelajari sepanjang semester, materi yang memberikan dampak intelektual dan praktis "
        "paling transformatif bagi saya adalah integrasi antara **Manajemen Jadwal (Jalur Kritis CPM/PDM)**, **Manajemen Biaya (Earned Value Management / EVM)**, "
        "dan **Manajemen Risiko Kuantitatif (Expected Monetary Value & Kalibrasi Probabilitas)**.")

    add_body_p(doc,
        "Alasan mengapa materi-materi tersebut sangat bermanfaat adalah:")
    add_bullet_p(doc, "Mengubah Pola Pikir dari Intuisi Subjektif Menjadi Metrik Objektif: Sebelum mempelajari manajemen proyek, estimasi durasi dan biaya proyek sering kali dilakukan hanya berdasarkan tebakan kasar (gut feeling). Melalui teknik Precedence Diagramming Method (PDM), perhitungan lintasan kritis (Forward/Backward Pass), serta metrik EVM (CV, SV, CPI, SPI, EAC), saya memahami bahwa kesehatan sebuah proyek rekayasa perangkat lunak dapat diukur secara matematis dan presisi.")
    add_bullet_p(doc, "Korelasi Langsung dengan Proyek Riset PRE-02: Pemahaman mendalam tentang Schedule Performance Index (SPI), Cost Performance Index (CPI), dan Earned Value menjadi fondasi filosofis utama dalam merancang variabel fitur pada proyek PRE-02. Sistem early warning machine learning yang kami bangun pada dasarnya adalah otomasi cerdas yang mengevaluasi deviasi progres aktual terhadap baseline jadwal untuk memprediksi probabilitas kegagalan secara dini.")
    add_bullet_p(doc, "Pemahaman Holistik Transisi ke Metodologi Hibrida: Materi penutup tentang Agile dan Hybrid membuka wawasan saya bahwa standar baku PMBOK bukanlah belenggu birokrasi yang kaku, melainkan sebuah 'toolbox' yang sangat fleksibel untuk disesuaikan (tailored) dengan kebutuhan industri teknologi informasi modern.")

    # =========================================================================
    # REFLEKSI 2
    # =========================================================================
    add_heading_2(doc, "2. Apa tantangan terbesar yang Anda hadapi selama mengerjakan Tugas Besar kelompok, dan bagaimana Anda mengatasinya?")
    add_body_p(doc,
        "Tantangan terbesar yang saya hadapi selama pengerjaan Tugas Besar Berjenjang pada proyek riset PRE-02 mencakup dua dimensi utama, "
        "yaitu tantangan teknis sains data dan tantangan manajerial koordinasi tim:")

    add_num_p(doc, "**Tantangan Teknis (Data Imbalance dan Kalibrasi Probabilitas Kontinu):** "
        "Dataset telemetri proyek ERP Farm Nation Enterprise memiliki karakteristik class imbalance yang signifikan (jumlah proyek tepat waktu jauh lebih banyak daripada proyek yang terlambat). Selain itu, model machine learning konvensional cenderung menghasilkan probabilitas overconfident (terlalu yakin mendekati 0 atau 1) yang tidak mencerminkan frekuensi empiris aktual. "
        "Strategi Mengatasinya: Kami menerapkan teknik kalibrasi probabilitas pasca-pelatihan (post-processing calibration) menggunakan Platt Scaling dan Isotonic Regression. Kami juga menggeser metrik evaluasi dari akurasi mentah ke Brier Score dan Expected Calibration Error (ECE) untuk memastikan bahwa probabilitas keterlambatan benar-benar dapat dipercaya sebagai indikator risiko.", 1)

    add_num_p(doc, "**Tantangan Manajerial (Sinkronisasi Beban Kerja dan Manajemen Waktu):** "
        "Setiap anggota tim memiliki ritme kerja, tingkat pemahaman teknis, dan kesibukan mata kuliah lain yang berbeda-beda, sehingga sempat terjadi hambatan ketergantungan tugas (task dependency bottlenecks) pada saat penyusunan WBS dan integrasi dataset. "
        "Strategi Mengatasinya: Kami mengadopsi prinsip Agile Scrum secara mandiri dengan membagi target mingguan ke dalam backlog kecil di Trello/GitHub Project, mengadakan koordinasi singkat berkala (daily stand-up virtual), serta menerapkan komunikasi asinkron yang transparan di grup kerja. Hal ini secara signifikan mereduksi miskomunikasi dan memastikan seluruh Checkpoint #1 s.d. #6 selesai tepat waktu.", 2)

    # =========================================================================
    # REFLEKSI 3
    # =========================================================================
    add_heading_2(doc, "3. Bagaimana Anda akan menerapkan ilmu manajemen proyek yang telah dipelajari di masa depan, baik dalam karir maupun kehidupan sehari-hari?")
    add_body_p(doc,
        "Ilmu manajemen proyek yang diperoleh dari perkuliahan ini memberikan bekal fundamental yang akan saya terapkan "
        "secara berkelanjutan dalam beberapa dimensi kehidupan:")

    add_num_p(doc, "**Dalam Karir Profesional sebagai Software Engineer / Data Scientist:** "
        "Saya tidak lagi memandang pembuatan kode program sebagai aktivitas yang terisolasi dari tujuan bisnis organisasi. Saya akan selalu menerapkan prinsip Work Breakdown Structure (WBS) untuk mendekomposisi modul kompleks, mengestimasi story points secara realistis, menerapkan Definition of Done (DoD) yang ketat untuk mencegah utang teknis (technical debt), serta mengidentifikasi risiko integrasi sistem sejak fase awal arsitektur perangkat lunak.", 1)

    add_num_p(doc, "**Dalam Peran Kepemimpinan Proyek (Project Manager / Scrum Master Masa Depan):** "
        "Saya akan mempraktikkan gaya kepemimpinan servant leadership, memfasilitasi kolaborasi tanpa sekat hierarki, aktif mengelola ekspektasi pemangku kepentingan menggunakan analisis Power/Interest Grid, dan memastikan komunikasi tim berjalan lancar dan terbuka tanpa menyembunyikan masalah kinerja.", 2)

    add_num_p(doc, "**Dalam Kehidupan Akademis dan Manajemen Waktu Pribadi:** "
        "Filosofi Critical Path Method (CPM) sangat relevan untuk mengelola skripsi dan studi tingkat akhir. Saya membedakan antara tugas yang berada pada lintasan kritis (yang jika tertunda akan menunda kelulusan) dan tugas yang memiliki total float (kelonggaran waktu). Hal ini menghindarkan saya dari prokrastinasi dan kelelahan mental (burnout) melalui alokasi energi yang terstruktur dan terukur.", 3)

    # =========================================================================
    # REFLEKSI 4
    # =========================================================================
    add_heading_2(doc, "4. Berikan saran konstruktif untuk perbaikan mata kuliah Manajemen Proyek ke depan!")
    add_body_p(doc,
        "Perkuliahan Teori Manajemen Proyek yang diampu oleh Bapak Dr. Muhammad Ainul Yaqin, S.Si, M.Kom telah dirancang "
        "dengan kurikulum yang sangat komprehensif, terstruktur secara logis dari inisiasi hingga penutupan, serta menuntut kedisiplinan akademis yang tinggi. "
        "Sebagai masukan konstruktif untuk pengembangan mata kuliah di masa mendatang, saya mengusulkan beberapa poin berikut:")

    add_bullet_p(doc, "Pengayaan Praktik Hands-On dengan Perkakas Manajemen Proyek Modern: Di samping pengerjaan dokumen analitis, mahasiswa akan sangat diuntungkan jika diperkenalkan secara langsung dengan simulasi perangkat lunak manajemen proyek berbasis cloud standar industri, seperti Jira Software, Microsoft Project, Monday.com, atau ClickUp. Hal ini akan mempersempit kesenjangan antara teori PMBOK dan eksekusi operasional di dunia kerja nyata.")
    add_bullet_p(doc, "Integrasi Studi Kasus Kegagalan Proyek TI Nyata (Post-Mortem Analysis): Menambahkan sesi bedah kasus terhadap proyek perangkat lunak skala nasional atau internasional yang mengalami kegagalan fatal (seperti kasus Healthcare.gov atau sistem ERP korporat besar) guna mendiskusikan 'lessons learned' terkait akar penyebab kegagalan manajemen ruang lingkup dan komunikasi.")
    add_bullet_p(doc, "Simulasi Negosiasi Stakeholder dan Manajemen Konflik Interaktif: Menyelenggarakan sesi 'role-play' interaktif di kelas di mana mahasiswa memerankan peran Project Manager, Sponsor/Klien yang menuntut perubahan mendadak, dan Lead Developer yang menolak penambahan fitur, untuk melatih kemampuan komunikasi dan negosiasi interpersonal mahasiswa secara nyata.")

    # =========================================================================
    # REFLEKSI 5
    # =========================================================================
    add_heading_2(doc, "5. Jika Anda harus memilih antara pendekatan prediktif, Agile, atau Hybrid untuk proyek Anda di masa depan, pendekatan mana yang akan Anda pilih dan mengapa?")
    add_body_p(doc,
        "Jika dihadapkan pada pilihan metodologi untuk mengelola proyek teknologi informasi di masa depan, saya secara tegas akan "
        "memilih **Pendekatan Hibrida (Hybrid Approach)**.")

    add_body_p(doc,
        "Justifikasi rasional dan pertimbangan strategis atas pilihan ini mencakup:")
    add_bullet_p(doc, "Realitas Industri Tidak Pernah Hitam-Putih: Di dunia industri nyata, hampir tidak ada proyek teknologi informasi modern yang dapat diselesaikan dengan sempurna hanya menggunakan 100% Waterfall kaku atau 100% Agile tanpa aturan. Proyek selalu memiliki batasan eksternal yang bersifat prediktif (seperti batasan anggaran tetap dari klien, tanggal peluncuran yang terikat regulasi pemerintah, dan kontrak hukum formal), berdampingan dengan kebutuhan eksekusi teknis yang membutuhkan kelincahan iteratif dan adaptasi cepat.")
    add_bullet_p(doc, "Prinsip 'Best of Both Worlds' (Disiplin Bertemu Agilitas): Pendekatan Hybrid memungkinkan saya sebagai manajer proyek untuk mengambil disiplin tata kelola dari Waterfall (kejelasan arsitektur awal, mitigasi risiko strategis, mitigasi kepatuhan hukum) sekaligus merangkul kelincahan dari Agile (siklus sprint 2 mingguan, demo produk berkala, umpan balik pengguna yang cepat, moral tim yang tinggi melalui self-organization).")
    add_bullet_p(doc, "Fleksibilitas Tailoring Berdasarkan Model Kompleksitas Stacey: Pendekatan Hybrid memberikan kebebasan bagi manajer proyek untuk mengatur 'slider' metodologi—menjadi lebih prediktif ketika menghadapi komponen infrastruktur fisik yang deterministik, atau bergeser menjadi lebih adaptif ketika mengeksplorasi antarmuka pengguna (UI/UX) dan algoritma kecerdasan buatan.")
    add_bullet_p(doc, "Kesimpulan Filosofis: Menjadi profesional manajemen proyek yang unggul bukan tentang mempertahankan kefanatikan terhadap satu nama metodologi, melainkan tentang kemampuan bijaksana untuk memilih perkakas yang tepat bagi masalah yang tepat guna memberikan nilai bisnis terbaik bagi organisasi dan masyarakat.")
