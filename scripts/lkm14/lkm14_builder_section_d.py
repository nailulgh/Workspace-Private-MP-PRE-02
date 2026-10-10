# -*- coding: utf-8 -*-
"""
Section D Builder: Daftar Pustaka Khusus Pertemuan 14
Pertemuan 14: Pendekatan Agile/Hybrid & Review Perkuliahan
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_d(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_body_p = helpers['add_body_p']

    add_heading_1(doc, "D. DAFTAR PUSTAKA KHUSUS PERTEMUAN 14")

    add_body_p(doc,
        "Daftar pustaka ini memuat seluruh literatur rujukan akademis, standar manajemen proyek internasional PMI, "
        "buku teks kerangka kerja Agile dan Scrum, serta teori perancangan metodologi hibrida yang digunakan dalam penyusunan "
        "Lembar Kerja Mahasiswa (LKM) Pertemuan 14 mengikuti format baku APA 7th Edition (American Psychological Association):")

    def add_hanging_indent_p(document, text):
        p = document.add_paragraph()
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.15
        run = p.add_run(text)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(12)
        return p

    add_hanging_indent_p(doc,
        "Agile Alliance, & Project Management Institute. (2017). Agile Practice Guide. Newtown Square, PA: Project Management Institute, Inc.")

    add_hanging_indent_p(doc,
        "Boehm, B., & Turner, R. (2004). Balancing Agility and Discipline: A Guide for the Perplexed. Boston, MA: Addison-Wesley Professional.")

    add_hanging_indent_p(doc,
        "Highsmith, J. (2009). Agile Project Management: Creating Innovative Products (2nd ed.). Boston, MA: Addison-Wesley Professional.")

    add_hanging_indent_p(doc,
        "Kerzner, H. (2017). Project Management: A Systems Approach to Planning, Scheduling, and Controlling (12th ed.). Hoboken, NJ: John Wiley & Sons.")

    add_hanging_indent_p(doc,
        "Larson, E. W., & Gray, C. F. (2021). Project Management: The Managerial Process (8th ed.). New York, NY: McGraw-Hill Education.")

    add_hanging_indent_p(doc,
        "Project Management Institute. (2017). A Guide to the Project Management Body of Knowledge (PMBOK® Guide) (6th ed.). Newtown Square, PA: Project Management Institute, Inc.")

    add_hanging_indent_p(doc,
        "Project Management Institute. (2021). A Guide to the Project Management Body of Knowledge (PMBOK® Guide) (7th ed.). Newtown Square, PA: Project Management Institute, Inc.")

    add_hanging_indent_p(doc,
        "Schwaber, K., & Sutherland, J. (2020). The Scrum Guide: The Definitive Guide to Scrum: The Rules of the Game. Scrum.org. https://scrumguides.org/")

    add_hanging_indent_p(doc,
        "Schwalbe, K. (2019). Information Technology Project Management (9th ed.). Boston, MA: Cengage Learning.")
