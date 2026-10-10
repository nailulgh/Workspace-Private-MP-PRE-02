# -*- coding: utf-8 -*-
"""
Section B Builder: Tugas Latihan Individu (Network Diagram dan Jalur Kritis)
Pertemuan 5: Manajemen Jadwal Proyek (Bagian I)
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

import os
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_b(doc, helpers, workspace_root):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "B. TUGAS LATIHAN INDIVIDU: NETWORK DIAGRAM DAN JALUR KRITIS")

    # =========================================================================
    # DATA SOAL
    # =========================================================================
    add_heading_2(doc, "1. Deskripsi Soal dan Data Aktivitas Proyek")
    
    add_body_p(doc,
        "Diberikan sekumpulan data aktivitas proyek rekayasa perangkat lunak yang memuat nama aktivitas, aktivitas pendahulu "
        "(predecessor), serta estimasi durasi deterministik dalam satuan hari kerja kalender sebagai berikut:")

    table_data_soal = [
        ["A", "-", "3"],
        ["B", "A", "4"],
        ["C", "A", "5"],
        ["D", "B, C", "2"],
        ["E", "D", "3"]
    ]
    create_table(doc, [4.0, 4.0, 4.0],
                 ["Aktivitas", "Predecessor", "Durasi (hari)"],
                 table_data_soal,
                 align_cols=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])

    add_body_p(doc,
        "Berdasarkan data aktivitas di atas, mahasiswa diinstruksikan untuk menyelesaikan empat tugas analisis jaringan jadwal:")
    add_bullet_p(doc, "Instruksi 1:", "Membuat network diagram menggunakan Precedence Diagramming Method (PDM) secara rapi dan jelas lengkap dengan node aktivitas.")
    add_bullet_p(doc, "Instruksi 2:", "Menentukan jalur kritis dari proyek tersebut dengan mengidentifikasi semua kemungkinan jalur beserta durasinya.")
    add_bullet_p(doc, "Instruksi 3:", "Menentukan total durasi penyelesaian proyek berdasarkan perhitungan jalur kritis.")
    add_bullet_p(doc, "Instruksi 4:", "Melakukan perhitungan maju (forward pass) dan perhitungan mundur (backward pass) untuk menentukan Early Start (ES), Early Finish (EF), Late Start (LS), Late Finish (LF), serta Float/Slack setiap aktivitas.")

    # =========================================================================
    # INSTRUKSI 1: NETWORK DIAGRAM PDM
    # =========================================================================
    add_heading_2(doc, "2. Penyusunan Network Diagram Menggunakan Precedence Diagramming Method (PDM)")
    
    add_body_p(doc,
        "Precedence Diagramming Method (PDM) / Activity-on-Node (AON) memodelkan aktivitas sebagai simpul kotak (node) dan relasi "
        "ketergantungan sebagai garis panah berarah. Setiap simpul node dirancang mengacu pada standar format 7 kompartemen baku "
        "PMBOK Guide (PMI, 2017) dan Larson & Gray (2021) dengan struktur internal sebagai berikut:")

    table_node_anatomy = [
        ["Early Start (ES)", "Durasi Aktivitas (D)", "Early Finish (EF)"],
        ["Kode & Nama Aktivitas / Identitas Pekerjaan", "Kode & Nama Aktivitas / Identitas Pekerjaan", "Kode & Nama Aktivitas / Identitas Pekerjaan"],
        ["Late Start (LS)", "Total Float (TF)", "Late Finish (LF)"]
    ]
    create_table(doc, [5.0, 5.0, 5.0],
                 ["Kompartemen Kiri", "Kompartemen Tengah", "Kompartemen Kanan"],
                 table_node_anatomy,
                 align_cols=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])

    add_body_p(doc,
        "Visualisasi grafis lengkap dari network diagram proyek yang mengintegrasikan seluruh node, anak panah relasi ketergantungan Finish-to-Start (FS), "
        "serta penandaan tegas lintasan jalur kritis (highlight merah) disajikan pada Gambar 5.1 di bawah ini:")

    # Insert Image
    img_path = os.path.join(workspace_root, "scripts", "lkm5", "pdm_network_diagram.png")
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Cm(15.2))

        p_caption = doc.add_paragraph()
        p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_caption.paragraph_format.space_before = Pt(2)
        p_caption.paragraph_format.space_after = Pt(10)
        r_cap = p_caption.add_run("Gambar 5.1 Precedence Diagramming Method (PDM) Network Diagram Proyek Latihan CPM")
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(80, 80, 80)

    add_body_p(doc,
        "Penjelasan struktur ketergantungan logis antarsimpul node pada Gambar 5.1:")
    add_num_p(doc, "1.", "Aktivitas A (Durasi 3 hari):",
        "Merupakan simpul awal proyek tanpa pendahulu (Predecessor: -). Aktivitas A menjadi pendahulu bersama bagi dua cabang aktivitas selanjutnya, "
        "yaitu Aktivitas B dan Aktivitas C (hubungan Finish-to-Start / FS).")
    add_num_p(doc, "2.", "Aktivitas B (Durasi 4 hari):",
        "Berada pada cabang atas dengan pendahulu tunggal Aktivitas A. Setelah Aktivitas B selesai, ia mengalirkan ketergantungan menuju Aktivitas D.")
    add_num_p(doc, "3.", "Aktivitas C (Durasi 5 hari):",
        "Berada pada cabang bawah dengan pendahulu tunggal Aktivitas A. Karena durasinya lebih panjang (5 hari berbanding 4 hari milik B), "
        "Aktivitas C menjadi penentu waktu tunggu (bottleneck) bagi simpul konvergen berikutnya.")
    add_num_p(doc, "4.", "Aktivitas D (Durasi 2 hari):",
        "Merupakan simpul konvergen / penggabungan (merge activity) yang memiliki dua pendahulu, yaitu Aktivitas B dan Aktivitas C. "
        "Berdasarkan logika Finish-to-Start, Aktivitas D baru dapat dimulai jika KEDUA pendahulunya (B dan C) telah selesai secara tuntas.")
    add_num_p(doc, "5.", "Aktivitas E (Durasi 3 hari):",
        "Merupakan simpul akhir proyek (terminal activity) yang memiliki pendahulu tunggal Aktivitas D. Selesainya Aktivitas E menandai tuntasnya "
        "seluruh pekerjaan proyek.")

    # =========================================================================
    # INSTRUKSI 2 & 3: JALUR KRITIS & TOTAL DURASI
    # =========================================================================
    add_heading_2(doc, "3. Identifikasi Seluruh Jalur Proyek, Penentuan Jalur Kritis, dan Total Durasi Proyek")
    
    add_body_p(doc,
        "Sebuah jalur (path) dalam network diagram didefinisikan sebagai urutan aktivitas yang terhubung secara kontinu dari titik awal proyek "
        "hingga titik akhir penyelesaian proyek. Berdasarkan konfigurasi diagram PDM pada Gambar 5.1, terdapat 2 (dua) kemungkinan jalur:")

    add_bullet_p(doc, "Jalur 1 (Path 1):",
        "Rangkaian aktivitas: A → B → D → E\n"
        "– Perhitungan Durasi: Durasi(A) + Durasi(B) + Durasi(D) + Durasi(E)\n"
        "– Substitusi Angka: 3 hari + 4 hari + 2 hari + 3 hari = 12 Hari Kalender.")

    add_bullet_p(doc, "Jalur 2 (Path 2):",
        "Rangkaian aktivitas: A → C → D → E\n"
        "– Perhitungan Durasi: Durasi(A) + Durasi(C) + Durasi(D) + Durasi(E)\n"
        "– Substitusi Angka: 3 hari + 5 hari + 2 hari + 3 hari = 13 Hari Kalender.")

    add_body_p(doc,
        "Ringkasan evaluasi komparatif kedua jalur disajikan dalam tabel berikut:")

    table_data_paths = [
        ["Jalur 1", "A → B → D → E", "3 + 4 + 2 + 3", "12 Hari", "Non-Kritis (Slack = 1 Hari)"],
        ["Jalur 2", "A → C → D → E", "3 + 5 + 2 + 3", "13 Hari", "JALUR KRITIS (Total Float = 0)"]
    ]
    create_table(doc, [2.5, 4.0, 3.5, 2.5, 4.5],
                 ["Identitas Jalur", "Rangkaian Aktivitas", "Formula Penjumlahan", "Total Durasi", "Status Lintasan"],
                 table_data_paths,
                 align_cols=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])

    add_heading_3(doc, "Kesimpulan Instruksi 2 & Instruksi 3:")
    add_num_p(doc, "a.", "Penetapan Jalur Kritis (Critical Path):",
        "Jalur kritis proyek adalah Jalur 2, yaitu rangkaian aktivitas A → C → D → E. Hal ini didasarkan pada definisi matematis "
        "bahwa jalur kritis adalah jalur terpanjang dalam network diagram yang menentukan waktu tercepat penyelesaian proyek. "
        "Seluruh aktivitas pada jalur ini (A, C, D, dan E) merupakan aktivitas kritis yang tidak memiliki kelonggaran waktu sama sekali.")

    add_num_p(doc, "b.", "Penetapan Total Durasi Proyek:",
        "Total durasi penyelesaian proyek berdasarkan perhitungan jalur kritis adalah 13 Hari Kalender. Proyek tidak dapat diselesaikan "
        "dalam waktu 12 hari karena Aktivitas C memerlukan 5 hari kerja, yang secara fisik menahan dimulainya Aktivitas D hingga hari ke-8.")

    # =========================================================================
    # INSTRUKSI 4: FORWARD PASS & BACKWARD PASS
    # =========================================================================
    add_heading_2(doc, "4. Perhitungan Maju (Forward Pass) dan Perhitungan Mundur (Backward Pass) Lengkap")
    
    add_body_p(doc,
        "Untuk menentukan parameter waktu komprehensif pada setiap simpul aktivitas, dilakukan analisis dua arah menggunakan metode "
        "Critical Path Method (CPM) dengan konvensi standar Hari ke-0 (Day 0 Convention) yang dianut secara universal dalam perangkat lunak "
        "manajemen proyek modern (PMI, 2017; Larson & Gray, 2021):")

    add_heading_3(doc, "A. Perhitungan Maju (Forward Pass) — Menentukan ES dan EF:")
    add_body_p(doc,
        "Forward pass dihitung mulai dari simpul awal menuju simpul akhir proyek untuk mencari tanggal mulai paling awal (Early Start / ES) "
        "dan tanggal selesai paling awal (Early Finish / EF). Formula matematis yang berlaku:")
    add_subbullet_p(doc, "Rumus Durasi:", "EF = ES + Durasi")
    add_subbullet_p(doc, "Rumus Simpul Awal:", "ES_awal = 0")
    add_subbullet_p(doc, "Rumus Simpul Penggabungan (Merge):", "ES_successor = max(EF seluruh predecessor)")

    add_body_p(doc, "Langkah-langkah komputasi maju per aktivitas:")
    add_num_p(doc, "1.", "Aktivitas A (Durasi = 3 hari):",
        "– Karena A tidak memiliki predecessor, maka ES(A) = 0.\n"
        "– EF(A) = ES(A) + Durasi(A) = 0 + 3 = 3.")

    add_num_p(doc, "2.", "Aktivitas B (Durasi = 4 hari):",
        "– Predecessor tunggal adalah A, maka ES(B) = EF(A) = 3.\n"
        "– EF(B) = ES(B) + Durasi(B) = 3 + 4 = 7.")

    add_num_p(doc, "3.", "Aktivitas C (Durasi = 5 hari):",
        "– Predecessor tunggal adalah A, maka ES(C) = EF(A) = 3.\n"
        "– EF(C) = ES(C) + Durasi(C) = 3 + 5 = 8.")

    add_num_p(doc, "4.", "Aktivitas D (Durasi = 2 hari):",
        "– D memiliki dua predecessor (B dan C). Berdasarkan aturan komputasi maju simpul penggabungan, ES(D) ditentukan oleh nilai EF terbesar:\n"
        "  ES(D) = max(EF(B), EF(C)) = max(7, 8) = 8.\n"
        "– EF(D) = ES(D) + Durasi(D) = 8 + 2 = 10.")

    add_num_p(doc, "5.", "Aktivitas E (Durasi = 3 hari):",
        "– Predecessor tunggal adalah D, maka ES(E) = EF(D) = 10.\n"
        "– EF(E) = ES(E) + Durasi(E) = 10 + 3 = 13.\n"
        "– Karena E adalah aktivitas terakhir, maka Durasi Total Proyek = EF(E) = 13 hari.")

    add_heading_3(doc, "B. Perhitungan Mundur (Backward Pass) — Menentukan LS dan LF:")
    add_body_p(doc,
        "Backward pass dihitung mulai dari simpul akhir menuju simpul awal proyek untuk mencari batas waktu selesai paling lambat "
        "(Late Finish / LF) dan batas waktu mulai paling lambat (Late Start / LS) tanpa menyebabkan keterlambatan proyek secara keseluruhan. "
        "Formula matematis yang berlaku:")
    add_subbullet_p(doc, "Rumus Simpul Akhir:", "LF_akhir = EF_proyek = 13")
    add_subbullet_p(doc, "Rumus Pengurangan:", "LS = LF - Durasi")
    add_subbullet_p(doc, "Rumus Simpul Percabangan (Burst):", "LF_predecessor = min(LS seluruh successor)")

    add_body_p(doc, "Langkah-langkah komputasi mundur per aktivitas:")
    add_num_p(doc, "1.", "Aktivitas E (Durasi = 3 hari):",
        "– Sebagai aktivitas terminal akhir, LF(E) = EF(E) = 13.\n"
        "– LS(E) = LF(E) - Durasi(E) = 13 - 3 = 10.")

    add_num_p(doc, "2.", "Aktivitas D (Durasi = 2 hari):",
        "– Successor tunggal adalah E, maka LF(D) = LS(E) = 10.\n"
        "– LS(D) = LF(D) - Durasi(D) = 10 - 2 = 8.")

    add_num_p(doc, "3.", "Aktivitas C (Durasi = 5 hari):",
        "– Successor tunggal adalah D, maka LF(C) = LS(D) = 8.\n"
        "– LS(C) = LF(C) - Durasi(C) = 8 - 5 = 3.")

    add_num_p(doc, "4.", "Aktivitas B (Durasi = 4 hari):",
        "– Successor tunggal adalah D, maka LF(B) = LS(D) = 8.\n"
        "– LS(B) = LF(B) - Durasi(B) = 8 - 4 = 4.")

    add_num_p(doc, "5.", "Aktivitas A (Durasi = 3 hari):",
        "– A merupakan simpul percabangan (burst activity) yang memiliki dua successor (B dan C). Berdasarkan aturan komputasi mundur, "
        "LF(A) ditentukan oleh nilai LS terkecil dari para successor-nya:\n"
        "  LF(A) = min(LS(B), LS(C)) = min(4, 3) = 3.\n"
        "– LS(A) = LF(A) - Durasi(A) = 3 - 3 = 0.")

    add_heading_3(doc, "C. Perhitungan Float / Slack (Total Float & Free Float):")
    add_body_p(doc,
        "Float atau slack mencerminkan fleksibilitas kelonggaran waktu yang dimiliki suatu aktivitas:")
    add_bullet_p(doc, "Total Float (TF):",
        "Kelonggaran waktu di mana suatu aktivitas dapat ditunda tanpa menunda tanggal target penyelesaian proyek secara keseluruhan.\n"
        "Formula: TF = LS - ES = LF - EF.")
    add_bullet_p(doc, "Free Float (FF):",
        "Kelonggaran waktu di mana suatu aktivitas dapat ditunda tanpa menunda tanggal Early Start (ES) dari aktivitas penerusnya (successor).\n"
        "Formula: FF = min(ES seluruh successor) - EF_aktivitas.")

    add_body_p(doc, "Perhitungan Float masing-masing aktivitas:")
    add_subbullet_p(doc, "Aktivitas A:", "TF = 0 - 0 = 0 hari; FF = min(ES_B, ES_C) - EF_A = min(3, 3) - 3 = 0 hari.")
    add_subbullet_p(doc, "Aktivitas B:", "TF = 4 - 3 = 1 hari (atau 8 - 7 = 1 hari); FF = ES_D - EF_B = 8 - 7 = 1 hari.")
    add_subbullet_p(doc, "Aktivitas C:", "TF = 3 - 3 = 0 hari; FF = ES_D - EF_C = 8 - 8 = 0 hari.")
    add_subbullet_p(doc, "Aktivitas D:", "TF = 8 - 8 = 0 hari; FF = ES_E - EF_D = 10 - 10 = 0 hari.")
    add_subbullet_p(doc, "Aktivitas E:", "TF = 10 - 10 = 0 hari; FF = EF_proyek - EF_E = 13 - 13 = 0 hari.")

    # =========================================================================
    # TABEL MASTER CPM
    # =========================================================================
    add_heading_3(doc, "D. Tabel Master Ringkasan Perhitungan CPM:")
    add_body_p(doc,
        "Rekapitulasi komprehensif seluruh parameter penjadwalan jaringan (ES, EF, LS, LF, Total Float, Free Float, dan Status Kritis) "
        "disajikan pada Tabel 5.1 berikut:")

    table_data_master_cpm = [
        ["A", "-", "3", "0", "3", "0", "3", "0", "0", "KRITIS"],
        ["B", "A", "4", "3", "7", "4", "8", "1", "1", "Non-Kritis"],
        ["C", "A", "5", "3", "8", "3", "8", "0", "0", "KRITIS"],
        ["D", "B, C", "2", "8", "10", "8", "10", "0", "0", "KRITIS"],
        ["E", "D", "3", "10", "13", "10", "13", "0", "0", "KRITIS"]
    ]
    create_table(doc, [1.8, 2.0, 1.8, 1.4, 1.4, 1.4, 1.4, 1.8, 1.8, 2.2],
                 ["Aktivitas", "Predecessor", "Durasi", "ES", "EF", "LS", "LF", "Total Float", "Free Float", "Status"],
                 table_data_master_cpm,
                 align_cols=[WD_ALIGN_PARAGRAPH.CENTER] * 10)

    add_body_p(doc,
        "Analisis Manajerial Hasil Perhitungan:",
        bold_prefix="Interpretasi Kritis bagi Project Manager: ")
    add_body_p(doc,
        "Berdasarkan Tabel 5.1, Aktivitas B merupakan satu-satunya aktivitas non-kritis dalam proyek dengan nilai Total Float = 1 hari "
        "dan Free Float = 1 hari. Hal ini berarti tim pengembang yang dialokasikan pada Aktivitas B memiliki kelonggaran untuk menunda mulai "
        "hingga hari ke-4 (LS=4) atau memperpanjang durasi hingga 5 hari tanpa mengganggu dimulainya Aktivitas D pada hari ke-8 dan tanpa "
        "menggeser tanggal selesai proyek pada hari ke-13. Sebaliknya, Aktivitas A, C, D, dan E memiliki nilai float absolut nol (TF=0), "
        "sehingga setiap penyimpangan durasi sekecil apa pun pada keempat aktivitas tersebut akan secara otomatis mengakibatkan proyek terlambat.")
