# -*- coding: utf-8 -*-
"""
Section A Builder: Tugas Tertulis Individu
Pertemuan 5: Manajemen Jadwal Proyek (Bagian I)
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_a(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "A. TUGAS TERTULIS INDIVIDU")

    # =========================================================================
    # SOAL 1
    # =========================================================================
    add_heading_2(doc, "1. Apa yang dimaksud dengan manajemen jadwal proyek dan mengapa area pengetahuan ini penting dalam keberhasilan proyek?")
    
    add_body_p(doc,
        "Manajemen Jadwal Proyek (Project Schedule Management) adalah area pengetahuan manajemen proyek yang mencakup proses-proses "
        "yang diperlukan untuk mengelola penyelesaian proyek secara tepat waktu sesuai dengan baseline waktu yang telah ditetapkan "
        "(Project Management Institute [PMI], 2017, hal. 173). Dalam disiplin rekayasa perangkat lunak dan manajemen proyek modern, "
        "jadwal bukan sekadar kalender kerja statis, melainkan sebuah model komputasi matematis dinamis yang mengintegrasikan "
        "seluruh dekomposisi aktivitas, durasi kerja, alokasi sumber daya manusia dan infrastruktur, dependensi logis antartugas, "
        "serta batasan kalender organisasi menjadi sebuah kerangka waktu eksekusi yang terukur.")

    add_body_p(doc,
        "Secara teoritis dan praktis, Manajemen Jadwal Proyek menempati peran sentral dalam menentukan keberhasilan proyek karena "
        "alasan-alasan fundamental berikut:")

    add_num_p(doc, "1.", "Pilar Utama Penyeimbang Segitiga Kendala Proyek (The Iron Triangle):",
        "Waktu adalah salah satu dari tiga kendala klasik proyek (bersama dengan Ruang Lingkup dan Biaya). Proyek yang mengalami disrupsi "
        "jadwal secara otomatis akan memicu kegagalan berantai pada dua pilar lainnya. Penambahan durasi kerja secara langsung "
        "meningkatkan beban biaya overhead, upah tenaga ahli pengembang, dan biaya operasional server (burn rate). Sebaliknya, upaya memotong "
        "waktu secara serampangan tanpa perencanaan matang kerap memaksa tim mengorbankan kualitas (quality trade-off) atau memangkas ruang "
        "lingkup penting (Larson & Gray, 2021).")

    add_num_p(doc, "2.", "Sifat Waktu yang Tidak Dapat Diperbarui (Irreversible Resource):",
        "Berbeda dengan anggaran keuangan yang masih dapat dinegosiasikan ulang melalui pendanaan darurat (contingency reserves) atau "
        "sumber daya komputasi yang dapat ditingkatkan kapasitasnya (scalable infrastructure), waktu yang telah terlewat tidak dapat "
        "ditarik kembali. Keterlambatan satu hari kalender pada aktivitas kritis memangkas margin keselamatan proyek secara permanen.")

    add_num_p(doc, "3.", "Menjaga Relevansi Nilai Bisnis dan Jendela Pasar (Window of Opportunity):",
        "Dalam industri teknologi informasi, ketepatan waktu rilis (time-to-market) merupakan pembeda utama antara keberhasilan komersial "
        "dan kegagalan produk. Keterlambatan penyelesaian modul perangkat lunak strategis—seperti sistem ERP perusahaan—dapat mengakibatkan "
        "kehilangan momentum persaingan, terkena klausul penalti kontrak keterlambatan (liquidated damages), serta merusak kepercayaan "
        "klien korporat dan pemangku kepentingan eksternal (Schwalbe, 2019).")

    add_num_p(doc, "4.", "Fondasi Sistem Peringatan Dini Proyek (Konteks Riset PRE-02):",
        "Bagi penelitian kelompok PRE-02 yang mengembangkan Early Warning System prediksi probabilitas keterlambatan, baseline jadwal "
        "yang terstruktur secara matematis merupakan prasyarat mutlak. Variabel prediktor esensial seperti Progres Aktual vs Rencana "
        "dan Sisa Waktu Deadline hanya dapat dihitung secara valid apabila proyek memiliki jadwal acuan (baseline) yang terdefinisi "
        "secara formal dan diawasi secara kuantitatif.")

    # =========================================================================
    # SOAL 2
    # =========================================================================
    add_heading_2(doc, "2. Sebutkan dan jelaskan secara singkat enam proses dalam manajemen jadwal proyek!")
    
    add_body_p(doc,
        "Berdasarkan standar PMBOK Guide 6th Edition (PMI, 2017, hal. 173-224), Manajemen Jadwal Proyek terdiri atas enam proses sistematis "
        "yang terbagi ke dalam Kelompok Proses Perencanaan (Planning Process Group) dan Kelompok Proses Pemantauan dan Pengendalian "
        "(Monitoring and Controlling Process Group). Keenam proses tersebut adalah:")

    add_num_p(doc, "1.", "Plan Schedule Management (Merencanakan Manajemen Jadwal):",
        "Proses menetapkan kebijakan, prosedur, standar, dan dokumentasi yang diperlukan untuk merencanakan, mengembangkan, mengelola, "
        "mengeksekusi, dan mengendalikan jadwal proyek. Output utama dari proses ini adalah Rencana Manajemen Jadwal (Schedule Management "
        "Plan) yang menetapkan unit waktu kerja (misal: jam atau hari kalender), tingkat akurasi estimasi, aturan pengukuran progres "
        "(seperti Earned Value rules), serta ambang batas varians jadwal yang ditoleransi.")

    add_num_p(doc, "2.", "Define Activities (Mendefinisikan Aktivitas):",
        "Proses mengidentifikasi dan mendokumentasikan aksi-aksi spesifik yang harus dilakukan guna menghasilkan deliverable proyek yang "
        "tercantum dalam paket kerja (work package) WBS. Proses ini mendekomposisi paket kerja menjadi unit-unit aktivitas yang memiliki "
        "durasi, estimasi biaya, dan kebutuhan sumber daya. Output kuncinya adalah Daftar Aktivitas (Activity List), Atribut Aktivitas "
        "(Activity Attributes), dan Daftar Milestone (Milestone List).")

    add_num_p(doc, "3.", "Sequence Activities (Mengurutkan Aktivitas):",
        "Proses mengidentifikasi dan mendokumentasikan hubungan logis dan ketergantungan (dependencies) antartugas proyek. Setiap aktivitas "
        "—kecuali aktivitas awal dan akhir—dihubungkan dengan minimal satu predecessor dan successor menggunakan metode Precedence "
        "Diagramming Method (PDM). Output utamanya adalah Diagram Jaringan Jadwal Proyek (Project Schedule Network Diagrams).")

    add_num_p(doc, "4.", "Estimate Activity Durations (Mengestimasi Durasi Aktivitas):",
        "Proses memperkirakan jumlah periode kerja kalender yang diperlukan untuk menyelesaikan masing-masing aktivitas dengan alokasi "
        "sumber daya yang direncanakan. Estimasi dilakukan menggunakan berbagai teknik estimasi durasi dan menghasilkan dokumen Estimasi "
        "Durasi Aktivitas (Activity Duration Estimates) beserta Basis Estimasi (Basis of Estimates).")

    add_num_p(doc, "5.", "Develop Schedule (Mengembangkan Jadwal):",
        "Proses menganalisis urutan aktivitas, durasi, kebutuhan sumber daya, dan batasan jadwal untuk menciptakan model jadwal terintegrasi "
        "guna eksekusi, pemantauan, dan pengendalian proyek. Proses ini menerapkan analisis jalur kritis (Critical Path Method / CPM), "
        "pemerataan sumber daya (Resource Leveling), dan kompresi jadwal. Output utamanya adalah Garis Dasar Jadwal (Schedule Baseline) "
        "dan Jadwal Proyek (Project Schedule) berupa Gantt Chart atau Milestone Chart.")

    add_num_p(doc, "6.", "Control Schedule (Mengendalikan Jadwal):",
        "Proses memantau status aktivitas proyek secara berkala untuk memperbarui progres proyek, membandingkan performa aktual terhadap "
        "Schedule Baseline, serta mengelola perubahan pada baseline jadwal melalui sistem kendali perubahan terintegrasi. Output kuncinya "
        "adalah Informasi Kinerja Pekerjaan (Work Performance Information), Perkiraan Jadwal (Schedule Forecasts), dan Permintaan Perubahan "
        "(Change Requests).")

    add_body_p(doc,
        "Ringkasan keterkaitan keenam proses manajemen jadwal proyek disajikan dalam tabel komparatif berikut:")

    table_data_p2 = [
        ["1. Plan Schedule Management", "Planning", "Project Charter, Project Management Plan", "Expert Judgment, Data Analysis, Meetings", "Schedule Management Plan"],
        ["2. Define Activities", "Planning", "Schedule Management Plan, Scope Baseline (WBS)", "Decomposition, Rolling Wave Planning", "Activity List, Activity Attributes, Milestone List"],
        ["3. Sequence Activities", "Planning", "Activity List, Scope Baseline, Assumptions Log", "Precedence Diagramming Method (PDM), Dependency Determination, Leads & Lags", "Project Schedule Network Diagrams"],
        ["4. Estimate Activity Durations", "Planning", "Activity List, Resource Requirements, Resource Calendars", "Analogous, Parametric, Three-Point Estimating, Bottom-Up, Reserve Analysis", "Duration Estimates, Basis of Estimates"],
        ["5. Develop Schedule", "Planning", "Network Diagrams, Duration Estimates, Resource Calendars", "Critical Path Method (CPM), Resource Optimization, Schedule Compression", "Schedule Baseline, Project Schedule, Schedule Data"],
        ["6. Control Schedule", "Monitoring & Controlling", "Schedule Baseline, Work Performance Data, Schedule Data", "Earned Value Analysis, Iteration Burndown Chart, Critical Path Method", "Work Performance Information, Schedule Forecasts, Change Requests"]
    ]
    create_table(doc, [4.0, 2.5, 3.5, 3.5, 3.5],
                 ["Proses PMBOK", "Process Group", "Input Kunci", "Tools & Techniques Utama", "Output Kunci"],
                 table_data_p2)

    # =========================================================================
    # SOAL 3
    # =========================================================================
    add_heading_2(doc, "3. Apa perbedaan antara aktivitas (activity) dan work package dalam WBS? Jelaskan dengan contoh!")
    
    add_body_p(doc,
        "Meskipun sering kali disalahpahami oleh praktisi pemula, Work Package (Paket Kerja) dalam WBS dan Aktivitas (Activity) dalam "
        "penjadwalan proyek adalah dua konsep yang berada pada domain dan tingkat abstraksi yang berbeda secara mendasar:")

    add_bullet_p(doc, "Work Package (Paket Kerja):",
        "Merupakan elemen hierarki terendah dari Work Breakdown Structure (WBS). Karakteristik utamanya adalah berorientasi pada hasil akhir "
        "atau luaran fisik (deliverable-oriented atau noun-oriented). Work package mendefinisikan APA (WHAT) yang harus diserahkan sebagai "
        "suatu entitas pekerjaan yang utuh, dan berfungsi sebagai unit kontrol akuntansi biaya (cost control account) serta penetapan "
        "tanggung jawab tim tingkat manajerial (Haugan, 2002; PMI, 2019).")

    add_bullet_p(doc, "Aktivitas (Activity):",
        "Merupakan komponen pekerjaan operasional yang dihasilkan dari proses dekomposisi lebih lanjut atas work package. Karakteristik "
        "utamanya adalah berorientasi pada tindakan atau proses pengerjaan (action-oriented atau verb-oriented). Aktivitas mendefinisikan "
        "BAGAIMANA (HOW) pekerjaan dilakukan, memiliki atribut durasi kalender tertentu, mengonsumsi sumber daya secara langsung, memiliki "
        "hubungan pendahulu/penerus (predecessor/successor), dan dijadwalkan secara eksplisit pada garis waktu kalender (Kerzner, 2017).")

    add_body_p(doc,
        "Perbandingan komprehensif antara Work Package dan Aktivitas diuraikan pada tabel berikut:")

    table_data_p3 = [
        ["Dimensi Konsep", "Work Package (WBS)", "Aktivitas (Penjadwalan Proyek)"],
        ["Orientasi Semantik", "Deliverable-oriented (Berorientasi hasil/benda/noun)", "Action-oriented (Berorientasi tindakan/kerja/verb)"],
        ["Tingkat Dokumen", "Level terbawah dari Scope Baseline (WBS)", "Elemen dasar dari Schedule Model (Network Diagram)"],
        ["Atribut Jadwal", "Tidak memiliki atribut predecessor, successor, dan tanggal kalender independen secara langsung", "Memiliki durasi spesifik, tanggal mulai/selesai (ES, EF, LS, LF), dan dependensi logis"],
        ["Pengukuran Kinerja", "Dievaluasi berdasarkan kelengkapan spesifikasi fisik deliverable (100% Rule)", "Dievaluasi berdasarkan persentase kemajuan durasi kerja fisik (progress tracking)"],
        ["Fungsi Manajerial", "Unit estimasi biaya dan penugasan penanggung jawab (Control Account / CAM)", "Unit eksekusi harian teknis tim dan pelacakan lintasan kritis (CPM)"]
    ]
    create_table(doc, [3.2, 6.9, 6.9],
                 ["Dimensi Konsep", "Work Package (WBS)", "Aktivitas (Penjadwalan Proyek)"],
                 table_data_p3[1:])

    add_body_p(doc,
        "Contoh Konkret pada Proyek PRE-02 (Sistem Prediksi Probabilitas Keterlambatan Proyek):",
        bold_prefix="Ilustrasi Dekomposisi Riil: ")
    add_subbullet_p(doc, "Work Package WBS (Kode 3.2):",
        "'Modul Kalibrasi Probabilitas (Calibrated Probability Module)' — Berorientasi benda/deliverable hasil verifikasi, yaitu paket modul "
        "perangkat lunak yang siap memproses skor probabilitas keluaran model machine learning.")
    add_subbullet_p(doc, "Dekomposisi Menjadi Aktivitas Operasional:",
        "Work package di atas didekomposisi oleh tim teknis menjadi serangkaian aktivitas dengan kata kerja konkret:")
    add_subbullet_p(doc, "  – Aktivitas ACT-3.2.1:", "Menulis script algoritma Platt Scaling berbasis regresi logistik (Durasi: 3 hari).")
    add_subbullet_p(doc, "  – Aktivitas ACT-3.2.2:", "Mengimplementasikan modul Isotonic Regression non-parametrik (Durasi: 4 hari, Pred: ACT-3.2.1).")
    add_subbullet_p(doc, "  – Aktivitas ACT-3.2.3:", "Menghitung nilai Expected Calibration Error (ECE) dan Brier Score pada test set (Durasi: 2 hari, Pred: ACT-3.2.2).")

    # =========================================================================
    # SOAL 4
    # =========================================================================
    add_heading_2(doc, "4. Sebutkan dan jelaskan empat jenis hubungan ketergantungan (dependency) antar aktivitas beserta contoh masing-masing!")
    
    add_body_p(doc,
        "Dalam metode Precedence Diagramming Method (PDM), urutan logis pengerjaan antartugas dimodelkan melalui empat jenis hubungan "
        "ketergantungan logis (logical relationships) sebagaimana didefinisikan dalam PMBOK Guide (PMI, 2017, hal. 190):")

    add_num_p(doc, "1.", "Finish-to-Start (FS) — Hubungan Selesai-ke-Mulai:",
        "Aktivitas penerus (successor) tidak dapat dimulai sebelum aktivitas pendahulu (predecessor) selesai secara tuntas. Ini merupakan "
        "jenis ketergantungan yang paling fundamental dan paling banyak digunakan dalam praktik rekayasa proyek (mencakup lebih dari 90% "
        "kasus penjadwalan).\n"
        "Contoh Nyata TI: Aktivitas 'Training Model Random Forest' baru dapat dimulai setelah aktivitas 'Pembersihan dan Preprocessing "
        "Dataset PMP FNE' selesai 100%. Model komputasi tidak dapat memproses matriks fitur yang masih memuat nilai hilang (missing values).")

    add_num_p(doc, "2.", "Start-to-Start (SS) — Hubungan Mulai-ke-Mulai:",
        "Aktivitas penerus (successor) tidak dapat dimulai sebelum aktivitas pendahulu (predecessor) telah resmi dimulai. Hubungan ini "
        "memungkinkan dua aktivitas berjalan secara paralel dengan penundaan awal tertentu.\n"
        "Contoh Nyata TI: Aktivitas 'Penulisan Draf Laporan Bab 3 Metodologi Eksperimen' dapat dimulai segera setelah aktivitas "
        "'Pengkodean Pipeline Machine Learning' resmi dimulai. Penulis laporan tidak perlu menunggu seluruh kode selesai dieksekusi untuk "
        "mulai menulis deskripsi arsitektur algoritma.")

    add_num_p(doc, "3.", "Finish-to-Finish (FF) — Hubungan Selesai-ke-Selesai:",
        "Aktivitas penerus (successor) tidak dapat selesai sebelum aktivitas pendahulu (predecessor) dinyatakan selesai tuntas. Keduanya "
        "bisa berjalan secara simultan, namun penguncian akhir dari aktivitas penerus terikat langsung pada penyelesaian pendahulunya.\n"
        "Contoh Nyata TI: Aktivitas 'Uji Regresi Sistem ERP (Regression Testing)' tidak dapat selesai sebelum seluruh aktivitas "
        "'Perbaikan Bug Kritis (Bug Fixing)' selesai tuntas dikerjakan oleh tim pengembang pengkodean.")

    add_num_p(doc, "4.", "Start-to-Finish (SF) — Hubungan Mulai-ke-Selesai:",
        "Aktivitas penerus (successor) tidak dapat selesai sebelum aktivitas pendahulu (predecessor) telah resmi dimulai. Ini adalah jenis "
        "hubungan yang paling jarang digunakan dalam industri perangkat lunak, namun sangat krusial pada skenario transisi operasional "
        "kritis atau serah terima shift tanpa jeda (zero-downtime cutover).\n"
        "Contoh Nyata TI: Server basis data warisan (legacy database system) tidak boleh dimatikan atau selesai beroperasi (successor "
        "selesai) sebelum server basis data kluster cloud yang baru telah berhasil booting dan mulai menerima traffic transaksi secara "
        "stabil (predecessor mulai).")

    add_body_p(doc,
        "Selain keempat hubungan logis di atas, PMBOK Guide mengklasifikasikan dependensi berdasarkan sifat keterikatannya menjadi 4 tipe:")
    add_bullet_p(doc, "Mandatory Dependencies (Hard Logic):",
        "Ketergantungan yang bersifat inheren dalam sifat fisik pekerjaan atau persyaratan kontrak yang tidak dapat ditawar (misal: "
        "koding modul harus dilakukan sebelum testing modul tersebut).")
    add_bullet_p(doc, "Discretionary Dependencies (Soft Logic / Preferred Logic):",
        "Ketergantungan yang ditetapkan berdasarkan praktik terbaik (best practices) atau preferensi tim proyek, yang dapat dimodifikasi "
        "apabila diperlukan percepatan jadwal.")
    add_bullet_p(doc, "External Dependencies:",
        "Ketergantungan antara aktivitas proyek internal dengan aktivitas pihak luar yang berada di luar kendali langsung tim proyek "
        "(misal: menunggu persetujuan API key atau sertifikasi keamanan dari vendor pihak ketiga).")
    add_bullet_p(doc, "Internal Dependencies:",
        "Ketergantungan yang sepenuhnya berada dalam kendali tim proyek internal organisasi.")

    # =========================================================================
    # SOAL 5
    # =========================================================================
    add_heading_2(doc, "5. Apa yang dimaksud dengan lead dan lag dalam penjadwalan proyek? Berikan contoh penggunaan keduanya!")
    
    add_body_p(doc,
        "Dalam pemodelan jaringan jadwal (network diagram), hubungan ketergantungan murni terkadang belum cukup fleksibel untuk "
        "merefleksikan kondisi kerja yang sebenarnya. Oleh karena itu, diterapkan parameter modifikasi waktu berupa Lead dan Lag "
        "(PMI, 2017, hal. 192-193):")

    add_bullet_p(doc, "Lead (Percepatan / Waktu Akselerasi):",
        "Merupakan jumlah waktu di mana suatu aktivitas penerus (successor) dapat dimajukan atau dimulai mendahului penyelesaian "
        "aktivitas pendahulunya (predecessor). Dalam formula matematis jadwal, lead diperlakukan sebagai nilai penyesuaian waktu negatif (-). "
        "Tujuan penggunaan lead adalah untuk menciptakan tumpang tindih (overlap) antartugas sehingga total durasi proyek dapat dipersingkat.\n"
        "– Notasi Matematis: FS – 2 hari (artinya successor dapat mulai 2 hari sebelum predecessor selesai 100%).\n"
        "– Contoh Kasus Nyata TI: Pada proyek pengembangan antarmuka pengguna (UI/UX), aktivitas 'Penyusunan Dokumentasi Panduan Pengguna' "
        "(successor) memiliki hubungan Finish-to-Start dengan lead 3 hari terhadap aktivitas 'Finalisasi Desain Antarmuka' (predecessor). "
        "Artinya, technical writer dapat mulai menyusun kerangka manual 3 hari sebelum tim desain menyelesaikan mockup terakhir.")

    add_bullet_p(doc, "Lag (Penundaan / Waktu Jeda):",
        "Merupakan jumlah waktu di mana suatu aktivitas penerus (successor) harus ditunda atau diundurkan mulainya setelah aktivitas "
        "pendahulu (predecessor) selesai. Lag bukanlah suatu aktivitas kerja mandiri dan tidak mengonsumsi tenaga manusia atau anggaran, "
        "melainkan jeda waktu pasif yang diwajibkan oleh proses fisik, teknis, atau hukum.\n"
        "– Notasi Matematis: FS + 1 hari (artinya successor baru boleh mulai 1 hari setelah predecessor tuntas).\n"
        "– Contoh Kasus Nyata TI: Pada modul migrasi basis data sistem ERP FNE, setelah aktivitas 'Migrasi Data Produksi 10 Juta Baris' "
        "selesai dilakukan (predecessor), sistem membutuhkan jeda waktu pasif (Lag = 24 jam) untuk proses automatisasi background indexing, "
        "replikasi kluster basis data lintas region, dan pembuatan cache memori sebelum aktivitas 'Pengujian Beban Skalabilitas (Stress Testing)' "
        "(successor) diizinkan untuk dimulai.")

    table_data_p5 = [
        ["Parameter", "Sifat Modifikasi", "Nilai Matematis", "Tujuan Utama", "Contoh Kasus Sistem TI"],
        ["Lead", "Percepatan / Akselerasi", "Negatif (-)", "Mempercepat jadwal melalui tumpang tindih (overlap) pekerjaan", "Memulai koding integrasi frontend 2 hari sebelum finalisasi spesifikasi API backend (FS - 2)"],
        ["Lag", "Penundaan Wajib / Jeda", "Positif (+)", "Memberikan jeda pasif non-kerja karena proses teknis/fisik", "Memberikan jeda 24 jam setelah deployment cloud sebelum audit penetrasi keamanan dimulai (FS + 1)"]
    ]
    create_table(doc, [2.5, 3.2, 2.5, 4.3, 4.5],
                 ["Parameter", "Sifat Modifikasi", "Nilai Matematis", "Tujuan Utama", "Contoh Kasus Sistem TI"],
                 table_data_p5)

    # =========================================================================
    # SOAL 6
    # =========================================================================
    add_heading_2(doc, "6. Sebutkan dan jelaskan empat teknik estimasi durasi aktivitas beserta kelebihan dan kekurangan masing-masing!")
    
    add_body_p(doc,
        "Untuk menghasilkan estimasi durasi yang realistis, manajer proyek dapat memilih satu atau kombinasi dari empat teknik estimasi "
        "utama sebagaimana dirumuskan dalam standar PMBOK Guide (PMI, 2017, hal. 195-204):")

    add_num_p(doc, "1.", "Analogous Estimating (Estimasi Berdasarkan Analogi / Top-Down):",
        "Teknik estimasi yang menggunakan parameter historis (seperti durasi aktual, anggaran, ukuran, atau kompleksitas) dari proyek masa "
        "lalu yang serupa sebagai dasar untuk memperkirakan durasi aktivitas proyek saat ini. Teknik ini mengandalkan penilaian para ahli "
        "(expert judgment).\n"
        "• Kelebihan: Sangat cepat, hemat biaya, dan dapat diterapkan pada fase inisiasi awal ketika informasi detail proyek belum tersedia.\n"
        "• Kekurangan: Memiliki tingkat akurasi paling rendah di antara semua teknik, sangat subjektif, dan rentan salah jika konteks proyek "
        "saat ini berbeda dari data historis masa lalu.")

    add_num_p(doc, "2.", "Parametric Estimating (Estimasi Parametrik):",
        "Teknik estimasi yang menggunakan hubungan algoritma atau statistik antara data historis masa lalu dengan variabel kuantitatif "
        "lainnya untuk menghitung durasi pekerjaan (misalnya: durasi = volume kerja ÷ produktivitas rata-rata per satuan waktu).\n"
        "• Kelebihan: Memiliki tingkat akurasi dan objektivitas yang tinggi jika didukung oleh basis data metrik organisasi yang valid.\n"
        "• Kekurangan: Sulit diterapkan pada pekerjaan kreatif, non-linier, atau riset yang tidak memiliki metrik kuantitatif seragam (seperti "
        "eksperimen hyperparameter tuning pada algoritma kecerdasan buatan).")

    add_num_p(doc, "3.", "Three-Point Estimating (Estimasi Tiga Titik / Distribusi Beta PERT):",
        "Teknik estimasi yang mengakomodasi risiko dan ketidakpastian dengan mendefinisikan tiga skenario durasi untuk setiap aktivitas: "
        "Durasi Optimis (t_O), Durasi Paling Mungkin (t_M), dan Durasi Pesimis (t_P). Nilai durasi ekspektasi terbobot (t_e) dihitung "
        "menggunakan formula Program Evaluation and Review Technique (PERT) berbasis distribusi Beta:\n"
        "    t_e = (t_O + 4 * t_M + t_P) / 6\n"
        "dengan varians ketidakpastian: σ² = ((t_P - t_O) / 6)².\n"
        "• Kelebihan: Mengurangi bias estimasi sepihak, memperhitungkan ketidakpastian risiko secara matematis, dan memungkinkan perhitungan "
        "probabilitas statistik penyelesaian proyek.\n"
        "• Kekurangan: Memerlukan waktu ekstra untuk memformulasikan tiga nilai estimasi per tugas serta memerlukan pemahaman statistik tim.")

    add_num_p(doc, "4.", "Bottom-Up Estimating (Estimasi Dari Bawah ke Atas):",
        "Teknik estimasi di mana pekerjaan pada work package didekomposisi ke tingkat tugas terkecil yang sangat mendetail, kemudian durasi "
        "dan kebutuhan sumber daya masing-masing tugas kecil diestimasi secara individual, lalu diagregasikan (dijumlahkan) ke atas "
        "hingga ke tingkat paket kerja dan proyek secara keseluruhan.\n"
        "• Kelebihan: Memberikan tingkat akurasi tertinggi dan tingkat kepemilikan komitmen (ownership) yang kuat dari tim pelaksana langsung.\n"
        "• Kekurangan: Memerlukan waktu perencanaan yang sangat lama, biaya administratif tinggi, dan membutuhkan WBS yang telah matang sempurna.")

    add_body_p(doc,
        "Perbandingan matriks komparatif keempat teknik estimasi durasi disajikan dalam tabel berikut:")

    table_data_p6 = [
        ["Teknik Estimasi", "Tingkat Akurasi", "Biaya & Waktu Upaya", "Kebutuhan Informasi Detail", "Kapan Paling Tepat Digunakan?"],
        ["Analogous", "Rendah (±25% s.d. ±50%)", "Sangat Rendah / Cepat", "Sangat Minimal (Data Level Makro)", "Fase inisiasi awal atau penyusunan Project Charter"],
        ["Parametric", "Tinggi (±10% s.d. ±15%)", "Sedang", "Tinggi (Metrik terukur & data historis)", "Pekerjaan repetitif terstandar (misal: pengujian 500 unit test cases)"],
        ["Three-Point (PERT)", "Tinggi (±10% s.d. ±20%)", "Sedang - Tinggi", "Sedang (Estimasi varians risiko)", "Pekerjaan riset, inovasi, dan pengembangan Machine Learning"],
        ["Bottom-Up", "Sangat Tinggi (±5% s.d. ±10%)", "Sangat Tinggi / Lambat", "Sangat Detail (WBS level aktivitas)", "Fase perencanaan matang sebelum penetapan Schedule Baseline"]
    ]
    create_table(doc, [3.2, 3.2, 3.2, 3.6, 3.8],
                 ["Teknik Estimasi", "Tingkat Akurasi", "Biaya & Waktu Upaya", "Kebutuhan Informasi Detail", "Kapan Paling Tepat Digunakan?"],
                 table_data_p6)

    # =========================================================================
    # SOAL 7
    # =========================================================================
    add_heading_2(doc, "7. Jelaskan apa yang dimaksud dengan Precedence Diagramming Method (PDM) atau Activity-on-Node (AON)! Mengapa metode ini banyak digunakan dalam praktik?")
    
    add_body_p(doc,
        "Precedence Diagramming Method (PDM)—yang dalam literatur manajemen operasi sering disebut sebagai Activity-on-Node (AON)—adalah "
        "metode penyusunan diagram jaringan jadwal proyek grafis di mana aktivitas direpresentasikan dalam bentuk kotak/simpul (node), "
        "dan hubungan ketergantungan logis antartugas dihubungkan menggunakan satu atau lebih garis panah berarah (PMI, 2017, hal. 189). "
        "Setiap simpul node memuat kompartemen data penjadwalan terstruktur yang mencakup Earliest Start (ES), Durasi (D), Earliest Finish (EF), "
        "Nama dan Kode Aktivitas, Latest Start (LS), Total Float (TF), serta Latest Finish (LF).")

    add_body_p(doc,
        "Dalam evolusi sejarah manajemen proyek, PDM telah menggantikan metode pendahulunya, yaitu Arrow Diagramming Method (ADM) atau "
        "Activity-on-Arrow (AOA). PDM menjadi metode standar universal yang diadopsi oleh praktisi global dan seluruh perangkat lunak modern "
        "(seperti Microsoft Project, Oracle Primavera P6, Jira, dan Smartsheet) karena empat keunggulan operasional utama:")

    add_num_p(doc, "1.", "Fleksibilitas Hubungan Ketergantungan Menyeluruh:",
        "Metode AOA terdahulu hanya mampu memodelkan satu jenis ketergantungan, yaitu Finish-to-Start (FS). Sebaliknya, PDM mendukung "
        "secara penuh keempat jenis hubungan logis (FS, SS, FF, dan SF) serta mengakomodasi parameter percepatan (Lead) dan penundaan "
        "(Lag) secara alami tanpa mengubah geometri diagram jaringan.")

    add_num_p(doc, "2.", "Eliminasi Kebutuhan Aktivitas Semu (Dummy Activities):",
        "Dalam diagram AOA, ketika dua aktivitas memiliki titik awal dan akhir yang sama atau ketika terdapat dependensi parsial yang rumit, "
        "diagrammer dipaksa menyisipkan 'dummy activity' (panah putus-putus berdurasi 0) untuk menghindari ambiguitas identitas. Dummy "
        "activity sangat membingungkan tim dan rentan memicu kesalahan komputasi manual. PDM sepenuhnya mengeliminasi dummy activity "
        "karena identitas aktivitas melekat secara unik pada node kotak.")

    add_num_p(doc, "3.", "Keterbacaan dan Kemudahan Modifikasi Diagram:",
        "Struktur node kotak pada PDM memungkinkan penambahan informasi atribut proyek yang kaya (seperti nama penanggung jawab / PIC, "
        "anggaran, dan float) langsung di dalam tubuh node. Ketika terjadi perubahan logika atau penambahan durasi pada fase pelaksanaan, "
        "diagram PDM jauh lebih mudah dimodifikasi dibandingkan diagram AOA.")

    add_num_p(doc, "4.", "Standarisasi Komputasi Algoritmik dalam Perangkat Lunak:",
        "Representasi struktur data graf berarah pada PDM (di mana node adalah objek pekerjaan dan tepi berarah adalah relasi logika) "
        "sangat selaras dengan algoritma komputasi ilmu komputer modern (seperti Topological Sort dan Dijkstra/Bellman-Ford). Hal ini "
        "membuat PDM menjadi fondasi baku mesin komputasi jadwal pada seluruh software enterprise.")

    # =========================================================================
    # SOAL 8
    # =========================================================================
    add_heading_2(doc, "8. Apa yang dimaksud dengan jalur kritis (critical path) dan mengapa penting untuk mengidentifikasinya dalam penjadwalan proyek?")
    
    add_body_p(doc,
        "Jalur Kritis (Critical Path) adalah urutan rangkaian aktivitas terpanjang yang menghubungkan titik awal proyek hingga titik akhir "
        "proyek dalam diagram jaringan, yang menentukan durasi kalender minimum mutlak yang diperlukan untuk menyelesaikan proyek secara "
        "keseluruhan (PMI, 2017, hal. 210; Larson & Gray, 2021, hal. 210). Secara matematis, seluruh aktivitas yang terletak pada jalur kritis "
        "disebut sebagai Aktivitas Kritis (Critical Activities) dan memiliki karakteristik khusus yaitu nilai Total Float (kelonggaran waktu) "
        "sama dengan nol (Total Float = 0, atau nilai float terkecil jika terdapat kendala tanggal target).")

    add_body_p(doc,
        "Identifikasi jalur kritis merupakan kompetensi terpenting bagi seorang manajer proyek karena implikasi strategis berikut:")

    add_num_p(doc, "1.", "Penerapan Prinsip Manajemen Berdasarkan Pengecualian (Management by Exception):",
        "Dalam proyek rekayasa perangkat lunak berskala besar yang memuat ratusan hingga ribuan aktivitas, seorang Project Manager "
        "tidak mungkin mengawasi setiap aktivitas dengan tingkat perhatian yang seragam. Jalur kritis bertindak sebagai 'filter prioritas "
        "eksekutif', di mana manajer memusatkan energi pengawasan harian pada aktivitas di jalur kritis. Keterlambatan satu hari kalender saja "
        "pada aktivitas kritis dipastikan akan langsung menggeser tanggal penyelesaian proyek secara keseluruhan.")

    add_num_p(doc, "2.", "Identifikasi Toleransi Kelonggaran Waktu (Flexibility of Non-Critical Paths):",
        "Aktivitas yang berada di luar jalur kritis memiliki Total Float positif (TF > 0), yang berarti aktivitas tersebut memiliki "
        "kelonggaran untuk ditunda atau diperpanjang durasinya tanpa memengaruhi tanggal selesai proyek. Pengetahuan ini memberi ruang "
        "manuver bagi manajer proyek untuk meminjamkan sumber daya dari tugas non-kritis guna membantu tugas kritis yang mengalami hambatan.")

    add_num_p(doc, "3.", "Landasan Efektif Strategi Kompresi Jadwal (Schedule Compression):",
        "Ketika proyek dituntut untuk selesai lebih cepat dari baseline karena tuntutan pasar atau sponsor, manajer proyek hanya boleh "
        "menerapkan teknik percepatan pada aktivitas di jalur kritis. Menerapkan Crashing (penambahan tenaga kerja berbayar) atau Fast-Tracking "
        "(memparalelkan tugas) pada aktivitas non-kritis hanya akan membuang-buang anggaran dan meningkatkan risiko pengerjaan ulang (rework) "
        "tanpa memotong durasi proyek satu hari pun!")

    add_num_p(doc, "4.", "Korelasi Langsung dengan Model Peringatan Dini PRE-02:",
        "Dalam konteks riset PRE-02, varians jadwal pada aktivitas di jalur kritis merupakan sinyal bahaya paling sensitif. Apabila aktivitas "
        "kritis terlambat, fitur prediktor Progres Aktual vs Rencana akan anjlok drastis dan model machine learning akan segera menghasilkan "
        "probabilitas keterlambatan tinggi (P(Delay) >= 0.65), sehingga alarm peringatan dini aktif sebelum kegagalan proyek menjadi permanen.")
