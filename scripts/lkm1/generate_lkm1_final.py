# -*- coding: utf-8 -*-
"""
Main Execution Script for LKM 1: PENGANTAR MANAJEMEN PROYEK
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))

def find_workspace_root(start_dir):
    curr = start_dir
    for _ in range(5):
        if os.path.exists(os.path.join(curr, "MP_Tugas_Individu")) or os.path.exists(os.path.join(curr, "uin_logo.png")):
            return curr
        parent = os.path.dirname(curr)
        if parent == curr:
            break
        curr = parent
    return os.getcwd()

WORKSPACE_ROOT = find_workspace_root(CURRENT_DIR)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(9)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body_p(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = False
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_bullet_p(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    
    r_bullet = p.add_run("• ")
    r_bullet.font.name = "Times New Roman"
    r_bullet.font.size = Pt(12)
    
    parts = text.split(":")
    if len(parts) > 1 and len(parts[0]) < 80:
        r_pre = p.add_run(parts[0] + ":")
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.bold = True
        
        r_body = p.add_run(":".join(parts[1:]))
        r_body.font.name = "Times New Roman"
        r_body.font.size = Pt(12)
        r_body.font.bold = False
    else:
        r_body = p.add_run(text)
        r_body.font.name = "Times New Roman"
        r_body.font.size = Pt(12)
        r_body.font.bold = False
    return p

def add_hanging_p(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Inches(-0.5)
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12)
    return p

def add_cover(doc):
    p0 = doc.add_paragraph()
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_after = Pt(5.8)
    p0.paragraph_format.line_spacing = 1.0
    r0 = p0.add_run("LKM 1: PENGANTAR MANAJEMEN PROYEK")
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(12)
    r0.font.bold = True
    
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.space_after = Pt(10)
    p1.paragraph_format.line_spacing = 1.0
    r1 = p1.add_run("Mata Kuliah: Teori Manajemen Proyek")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_after = Pt(10)
    
    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_after = Pt(5.8)
    p3.paragraph_format.line_spacing = 1.5
    r3 = p3.add_run("Dosen Pengampu:")
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(12)
    r3.font.bold = True
    
    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p4.paragraph_format.space_after = Pt(5.8)
    p4.paragraph_format.line_spacing = 1.5
    r4 = p4.add_run("Dr. Muhammad Ainul Yaqin, S.Si, M.Kom")
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(12)
    
    p5 = doc.add_paragraph()
    p5.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p5.paragraph_format.space_after = Pt(5.8)
    p5.paragraph_format.line_spacing = 1.5
    r5 = p5.add_run("                ")
    
    p6 = doc.add_paragraph()
    p6.paragraph_format.space_after = Pt(5.5)
    p6.paragraph_format.line_spacing = 1.5
    
    p7 = doc.add_paragraph()
    p7.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p7.paragraph_format.space_after = Pt(5.5)
    p7.paragraph_format.line_spacing = 1.5
    logo_path = os.path.join(WORKSPACE_ROOT, "uin_logo.png")
    if os.path.exists(logo_path):
        r7 = p7.add_run()
        r7.add_picture(logo_path, width=Cm(4.85))
        
    p8 = doc.add_paragraph()
    p8.paragraph_format.space_after = Pt(5.8)
    p8.paragraph_format.line_spacing = 1.5
    
    p9 = doc.add_paragraph()
    p9.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p9.paragraph_format.space_after = Pt(5.8)
    p9.paragraph_format.line_spacing = 1.5
    r9 = p9.add_run("Oleh :")
    r9.font.name = "Times New Roman"
    r9.font.size = Pt(12)
    r9.font.bold = True
    
    p10 = doc.add_paragraph()
    p10.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p10.paragraph_format.space_after = Pt(5.8)
    p10.paragraph_format.line_spacing = 1.5
    r10 = p10.add_run("Muhammad Nailul Ghufron Majid")
    r10.font.name = "Times New Roman"
    r10.font.size = Pt(12)
    
    p11 = doc.add_paragraph()
    p11.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p11.paragraph_format.space_after = Pt(5.8)
    p11.paragraph_format.line_spacing = 1.5
    r11 = p11.add_run("240605110160")
    r11.font.name = "Times New Roman"
    r11.font.size = Pt(12)
    
    for _ in range(3):
        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_after = Pt(5.8)
        p_sp.paragraph_format.line_spacing = 1.5
        
    inst_lines = [
        "PRODI TEKNIK INFORMATIKA",
        "FAKULTAS SAINS DAN TEKNOLOGI",
        "UNIVERSITAS ISLAM NEGERI MAULANA MALIK IBRAHIM",
        "MALANG",
        "2026"
    ]
    for line in inst_lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.5
        r = p.add_run(line)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.bold = True
        
    doc.add_page_break()

def generate_lkm1():
    print("Initializing document for LKM 1...")
    doc = docx.Document()
    
    for s in doc.sections:
        s.page_width = Cm(21.0)
        s.page_height = Cm(29.7)
        s.top_margin = Cm(3.0)
        s.bottom_margin = Cm(3.0)
        s.left_margin = Cm(3.0)
        s.right_margin = Cm(3.0)

    print("Building Cover Page...")
    add_cover(doc)

    print("Building Section A: Tugas Tertulis Individu...")
    add_heading_1(doc, "A. TUGAS TERTULIS INDIVIDU")

    # Soal 1
    add_heading_2(doc, "1. Apa yang dimaksud dengan proyek? Berikan minimal satu contoh proyek nyata dan jelaskan mengapa kegiatan tersebut termasuk proyek!")
    add_body_p(doc,
        "Definisi Proyek: Menurut PMBOK Guide 6th Edition (PMI, 2017), proyek adalah suatu usaha sementara (temporary endeavor) "
        "yang dilakukan untuk menciptakan produk, layanan, atau hasil yang bersifat unik (unique product, service, or result). "
        "Kata \"sementara\" menandakan bahwa proyek memiliki titik awal dan akhir yang terdefinisi dengan jelas, sedangkan kata \"unik\" "
        "menandakan bahwa luaran proyek memiliki karakteristik yang membedakannya secara substansial dari semua produk atau layanan lain yang pernah dibuat.")
    add_body_p(doc,
        "Contoh Proyek Nyata: Proyek Pengembangan Sistem Informasi / Aplikasi e-Commerce Kampus atau Proyek Riset \"Sistem Prediksi "
        "Probabilitas Keterlambatan Proyek Berbasis Machine Learning sebagai Early Warning System Manajemen Proyek\" (PRE-02).")
    add_body_p(doc, "Alasan kegiatan tersebut dikategorikan sebagai proyek meliputi tiga karakteristik fundamental:")
    add_bullet_p(doc, "Bersifat Sementara (Temporary): Memiliki rentang waktu pelaksanaan yang dibatasi secara definitif (misalnya diselesaikan dalam periode 1 semester / 6 bulan kalender).")
    add_bullet_p(doc, "Menghasilkan Luaran Unik (Unique Deliverable): Menghasilkan artefak perangkat lunak cerdas dengan model ML dan antarmuka analitik yang belum pernah diimplementasikan secara identik sebelumnya.")
    add_bullet_p(doc, "Dibatasi Sumber Daya (Constrained): Dikerjakan dalam koridor batasan ruang lingkup (scope), anggaran biaya (cost), dan alokasi tim pengembang (resources) yang terikat ketat.")

    # Soal 2
    add_heading_2(doc, "2. Apa perbedaan utama antara proyek dan kegiatan operasional rutin? Berikan contoh konkret untuk masing-masing!")
    add_body_p(doc,
        "Perbedaan fundamental antara proyek dan kegiatan operasional rutin terletak pada orientasi waktu, kebaruan luaran, "
        "serta tingkat ketidakpastian:")
    add_bullet_p(doc, "Proyek: Bersifat temporer (memiliki batas waktu selesai yang pasti), menghasilkan luaran yang unik, bertujuan melakukan perubahan atau inovasi (revolutionary change), dan memiliki tingkat ketidakpastian serta risiko yang relatif tinggi karena belum pernah dikerjakan persis sama sebelumnya.")
    add_bullet_p(doc, "Operasional Rutin: Berlangsung secara berkesinambungan (ongoing), bersifat repetitif dan baku, bertujuan memelihara stabilitas organisasi (evolutionary maintenance), serta memiliki tingkat ketidakpastian rendah karena proses kerjanya sudah terstandarisasi.")
    add_body_p(doc, "Contoh Konkret:")
    add_bullet_p(doc, "Contoh Proyek: Pembangunan dan implementasi sistem ERP kampus atau pelatihan model machine learning PRE-02 hingga siap rilis.")
    add_bullet_p(doc, "Contoh Operasional Rutin: Layanan administrasi akademik harian di biro kampus, pemeliharaan rutin server web, atau layanan helpdesk IT harian.")

    # Soal 3
    add_heading_2(doc, "3. Sebutkan dan jelaskan fase-fase dalam siklus hidup proyek (project life cycle) secara umum!")
    add_body_p(doc, "Siklus hidup proyek secara umum terdiri dari empat fase utama yang dilalui secara berurutan:")
    add_bullet_p(doc, "Fase Inisiasi (Starting the Project): Tahap konseptualisasi di mana urgensi proyek diidentifikasi, studi kelayakan dilakukan, sasaran strategis dirumuskan, dan persetujuan formal diterbitkan melalui Project Charter.")
    add_bullet_p(doc, "Fase Perencanaan (Organizing and Preparing): Tahap penyusunan rencana komprehensif (Project Management Plan) yang mencakup baseline ruang lingkup (WBS), estimasi durasi dan jadwal (CPM/Gantt Chart), anggaran biaya, serta mitigasi risiko.")
    add_bullet_p(doc, "Fase Eksekusi & Pengendalian (Carrying out the Work): Tahap realisasi pekerjaan secara fisik oleh tim pengembang untuk memproduksi deliverables, disertai monitoring berkala untuk menjaga ketercapaian baseline proyek.")
    add_bullet_p(doc, "Fase Penutupan (Closing the Project): Tahap serah terima deliverables resmi kepada pengguna/sponsor, evaluasi administratif, dokumentasi pelajaran yang dipetik (lessons learned), serta pelepasan seluruh sumber daya proyek.")

    # Soal 4
    add_heading_2(doc, "4. Pada proyek yang Anda identifikasi, fase-fase apa saja yang telah dilalui dan apa hasil keluaran (deliverables) dari setiap fasenya?")
    add_body_p(doc,
        "Berdasarkan proyek PRE-02 (Sistem Prediksi Probabilitas Keterlambatan Proyek Berbasis Machine Learning), berikut rincian fase "
        "dan deliverables yang dihasilkan:")
    add_bullet_p(doc, "Fase Inisiasi (Selesai): Deliverable berupa dokumen Project Charter yang memuat rumusan masalah klasifikasi biner vs probabilitas kontinu, penentuan tujuan proyek, anggaran awal, serta susunan pemangku kepentingan utama.")
    add_bullet_p(doc, "Fase Perencanaan (Selesai): Deliverable berupa Project Management Plan yang mencakup Work Breakdown Structure (WBS) hingga level paket kerja, Network Diagram PDM, jadwal baseline, serta estimasi biaya bottom-up.")
    add_bullet_p(doc, "Fase Eksekusi (Sedang Dilalui): Deliverable berupa dataset 7 fitur monitoring proyek yang telah dibersihkan, pipeline feature engineering, modul algoritma ML (Logistic Regression, Random Forest, XGBoost, MLP), serta antarmuka visualisasi prediksi.")
    add_bullet_p(doc, "Fase Penutupan (Tahap Akhir): Deliverable berupa laporan hasil evaluasi metrik (Brier Score, AUC-ROC, ECE), dokumen serah terima luaran riset, manual penggunaan sistem, dan repositori kode terintegrasi.")

    # Soal 5
    add_heading_2(doc, "5. Mengapa manajemen proyek penting bagi keberhasilan proyek yang Anda identifikasi? Sebutkan minimal tiga alasan!")
    add_body_p(doc,
        "Penerapan manajemen proyek yang disiplin sangat krusial bagi keberhasilan riset PRE-02 karena alasan-alasan strategis berikut:")
    add_bullet_p(doc, "Peningkatan Peluang Keberhasilan Produk: Memastikan bahwa model ML yang dikembangkan tidak sekadar akurat secara teoritis, tetapi benar-benar menjawab kebutuhan riil manajer proyek sebagai sistem peringatan dini yang andal.")
    add_bullet_p(doc, "Pengendalian Ketat Jadwal dan Biaya: Proyek AI/ML rentan mengalami pembengkakan waktu (schedule slip) akibat proses eksperimen model yang tidak berujung; manajemen proyek menetapkan milestone dan batasan waktu komputasi yang tegas.")
    add_bullet_p(doc, "Pengelolaan Risiko Ketidakpastian yang Terstruktur: Memungkinkan tim mengidentifikasi risiko teknis (seperti overfitting data atau ketidakseimbangan kelas) sejak dini dan menyiapkan strategi penanganan mitigasi.")
    add_bullet_p(doc, "Optimasi Alokasi Sumber Daya SDM: Menyelaraskan kapasitas kerja anggota tim lintas-fungsi (Data Scientist, ML Engineer, Backend Developer) sehingga tidak terjadi bottleneck maupun kelebihan beban kerja (over-allocation).")

    # Soal 6
    add_heading_2(doc, "6. Jelaskan perbedaan antara program dan portofolio! Berikan contoh masing-masing!")
    add_body_p(doc,
        "Perbedaan antara Program dan Portofolio terletak pada tingkat keterkaitan tujuan dan orientasi manajemen:")
    add_bullet_p(doc, "Program: Sekelompok proyek terkait, sub-program, dan aktivitas program yang dikelola secara terkoordinasi untuk memperoleh manfaat dan pengendalian yang tidak akan diperoleh jika dikelola secara terpisah (PMBOK Guide Bab 1.2.3). Contoh: Program Transformasi Smart Campus, yang membawahi Proyek Sistem Informasi Akademik Terpadu, Proyek Infrastruktur Wi-Fi Kecepatan Tinggi, dan Proyek Pembelajaran Daring (LMS).")
    add_bullet_p(doc, "Portofolio: Kumpulan proyek, program, portofolio sekunder, dan kegiatan operasional yang dikelompokkan bersama untuk memfasilitasi tata kelola efektif dari pekerjaan tersebut demi mencapai tujuan strategis organisasi. Contoh: Portofolio Transformasi Digital Perusahaan Teknologi, yang mengalokasikan investasi jangka panjang pada program otomatisasi rantai pasok, program AI analitik pelanggan, serta proyek modernisasi infrastruktur cloud.")

    # Soal 7
    add_heading_2(doc, "7. Sebutkan empat nilai inti dalam Kode Etik PMI dan jelaskan maknanya secara singkat!")
    add_body_p(doc,
        "Project Management Institute (PMI) menetapkan empat pilar nilai etis inti dalam PMI Code of Ethics and Professional Conduct:")
    add_bullet_p(doc, "Tanggung Jawab (Responsibility): Kewajiban etis manajer proyek untuk mengambil kepemilikan penuh atas setiap keputusan yang dibuat atau gagal dibuat, serta menerima konsekuensi profesional dari keputusan tersebut.")
    add_bullet_p(doc, "Rasa Hormat (Respect): Sikap memperlakukan orang lain, sumber daya yang dipercayakan, dan lingkungan secara adil, mendengarkan masukan tim, serta menjunjung tinggi martabat setiap pemangku kepentingan.")
    add_bullet_p(doc, "Keadilan (Fairness): Kewajiban bersikap objektif, tidak memihak, bebas dari diskriminasi dan benturan kepentingan (conflict of interest), serta menjamin kesetaraan akses informasi bagi pihak terkait.")
    add_bullet_p(doc, "Kejujuran (Honesty): Kewajiban memahami kebenaran dan senantiasa bertindak serta berkomunikasi secara jujur dan transparan, tanpa memanipulasi fakta status proyek kepada klien maupun pimpinan.")

    # Section B: Refleksi Mandiri
    print("Building Section B: Refleksi Mandiri...")
    add_heading_1(doc, "B. REFLEKSI MANDIRI")

    add_heading_2(doc, "1. Konsep atau materi apa yang paling menarik bagi Anda pada pertemuan hari ini, dan mengapa Anda tertarik pada konsep tersebut?")
    add_body_p(doc,
        "Konsep yang paling menarik bagi saya adalah diferensiasi tegas antara Proyek, Program, Portofolio, serta kegiatan Operasional Rutin. "
        "Selama ini, saya kerap menganggap seluruh pekerjaan institusional yang besar sebagai \"proyek\" tanpa menyadari batas hierarkinya. "
        "Memahami bahwa proyek memiliki Triple Constraints (Scope, Time, Cost) dan bertujuan menghasilkan perubahan unik, sementara portofolio "
        "berorientasi pada keselarasan investasi strategis jangka panjang, memberikan perspektif tata kelola yang sangat berharga bagi saya sebagai calon sarjana Teknik Informatika.")

    add_heading_2(doc, "2. Bagian materi apa yang masih terasa membingungkan atau sulit dipahami, dan apa yang akan Anda lakukan untuk memahaminya?")
    add_body_p(doc,
        "Bagian materi yang paling menantang adalah demarkasi taktis antara Program dan Portofolio dalam arsitektur perusahaan skala besar (Enterprise), "
        "terutama bagaimana memetakan dependensi lintas-proyek dan mengukur realisasi manfaat (benefit realization) secara objektif. "
        "Langkah tindak lanjut yang saya lakukan untuk memahaminya adalah:")
    add_bullet_p(doc, "Mempelajari literatur standar PMBOK Guide 6th Edition Bab 1 dan buku Project Management karya Larson & Gray (2021).")
    add_bullet_p(doc, "Menganalisis studi kasus proyek perangkat lunak skala enterprise (seperti integrasi multi-modul Super ERP).")
    add_bullet_p(doc, "Melakukan diskusi mendalam bersama rekan kelompok riset serta mengonsultasikannya langsung pada dosen pengampu.")

    add_heading_2(doc, "3. Bagaimana Anda dapat menerapkan pemahaman tentang manajemen proyek dalam kehidupan sehari-hari, baik di kampus, organisasi, maupun pekerjaan?")
    add_body_p(doc,
        "Penerapan manajemen proyek dapat saya wujudkan secara konkret pada berbagai dimensi aktivitas:")
    add_bullet_p(doc, "Di Lingkungan Kampus: Mengadopsi prinsip dekomposisi WBS dan penjadwalan berbasis milestone dalam menyusun pengerjaan Tugas Akhir/Skripsi dan tugas-tugas besar kuliah berjenjang agar terhindar dari sindrom penundaan (student syndrome).")
    add_bullet_p(doc, "Di Lingkungan Organisasi Mahasiswa: Mengelola kepanitiaan acara dengan baseline anggaran yang transparan, matriks penugasan tanggung jawab yang jelas, serta penilaian risiko sejak tahap perencanaan awal.")
    add_bullet_p(doc, "Di Lingkungan Karier / Pribadi: Merancang pengembangan perangkat lunak side-project atau rencana pengembangan diri menggunakan batasan ruang lingkup, estimasi waktu yang terukur, dan kriteria keberhasilan yang spesifik.")

    add_heading_2(doc, "4. Setelah mempelajari konsep proyek, apakah ada kegiatan yang selama ini Anda anggap biasa ternyata dapat dikategorikan sebagai proyek? Sebutkan contohnya!")
    add_body_p(doc,
        "Ya, ada. Contoh nyata adalah kepanitiaan festival kebudayaan/acara tahunan kampus dan proses renovasi rumah tempat tinggal. "
        "Alasan mengapa kegiatan-kegiatan tersebut tergolong sebagai proyek adalah:")
    add_bullet_p(doc, "Temporer (Temporary): Memiliki tanggal mulai dan tanggal penutupan yang definitif, di mana kepanitiaan dibubarkan setelah acara selesai.")
    add_bullet_p(doc, "Luaran Unik (Unique Output): Menghasilkan susunan acara atau bentuk renovasi spesifik dengan tantangan dan karakteristik yang tidak pernah berulang persis sama.")
    add_bullet_p(doc, "Dibatasi Sumber Daya: Pelaksanaannya terikat pada batasan dana anggaran yang terkumpul, kapasitas tim pelaksana, serta tenggat waktu pelaksanaan yang tidak dapat digeser sembarangan.")

    # Section C: Daftar Pustaka
    print("Building Section C: Daftar Pustaka...")
    add_heading_1(doc, "C. DAFTAR PUSTAKA KHUSUS PERTEMUAN 1")
    add_body_p(doc,
        "Daftar pustaka ini disusun berdasarkan kaidah sitasi baku APA 7th Edition (American Psychological Association) "
        "menggunakan literatur rujukan internasional primer yang diwajibkan dalam silabus perkuliahan:")

    add_hanging_p(doc, "Kerzner, H. (2017). Project Management: A Systems Approach to Planning, Scheduling, and Controlling (12th ed.). John Wiley & Sons.")
    add_hanging_p(doc, "Larson, E. W., & Gray, C. F. (2021). Project Management: The Managerial Process (8th ed.). McGraw-Hill Education.")
    add_hanging_p(doc, "Project Management Institute. (2017). A Guide to the Project Management Body of Knowledge (PMBOK® Guide) (6th ed.). Project Management Institute, Inc.")
    add_hanging_p(doc, "Schwalbe, K. (2019). Information Technology Project Management (9th ed.). Cengage Learning.")

    output_dir = os.path.join(WORKSPACE_ROOT, "MP_Tugas_Individu")
    os.makedirs(output_dir, exist_ok=True)
    output_path = os.path.join(output_dir, "LKM1_Manajemen_Proyek.docx")
    doc.save(output_path)
    print(f"\n[SUCCESS] Document LKM 1 successfully generated and saved to:\n  -> {output_path}")

if __name__ == "__main__":
    generate_lkm1()
