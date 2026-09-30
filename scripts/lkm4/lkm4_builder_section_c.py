# -*- coding: utf-8 -*-
"""
Section C Builder: Daftar Pustaka Khusus Pertemuan 4
Pertemuan 4: Manajemen Ruang Lingkup Proyek
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_c(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_body_p = helpers['add_body_p']

    add_heading_1(doc, "C. DAFTAR PUSTAKA KHUSUS PERTEMUAN 4")
    
    add_body_p(doc,
        "Daftar pustaka ini memuat seluruh literatur rujukan akademis, standar manajemen proyek internasional, "
        "dan buku teks rujukan utama yang diacu dalam penyusunan Lembar Kerja Mahasiswa (LKM) Pertemuan 4 ini:")

    references = [
        ("Haugan, G. T. (2002). ", "Effective Work Breakdown Structures", ". Vienna, VA: Management Concepts."),
        ("Kerzner, H. (2017). ", "Project Management: A Systems Approach to Planning, Scheduling, and Controlling", " (12th ed.). Hoboken, NJ: John Wiley & Sons."),
        ("Larson, E. W., & Gray, C. F. (2021). ", "Project Management: The Managerial Process", " (8th ed.). New York, NY: McGraw-Hill Education."),
        ("Project Management Institute. (2017). ", "A Guide to the Project Management Body of Knowledge (PMBOK® Guide)", " (6th ed.). Newtown Square, PA: Project Management Institute, Inc."),
        ("Project Management Institute. (2019). ", "Practice Standard for Work Breakdown Structures", " (3rd ed.). Newtown Square, PA: Project Management Institute, Inc."),
        ("Schwalbe, K. (2019). ", "Information Technology Project Management", " (9th ed.). Boston, MA: Cengage Learning."),
        ("Wiegers, K., & Beatty, J. (2013). ", "Software Requirements", " (3rd ed.). Redmond, WA: Microsoft Press.")
    ]

    for prefix, italic_title, suffix in references:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5

        r_pre = p.add_run(prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(12)
        r_pre.font.color.rgb = RGBColor(0, 0, 0)

        r_title = p.add_run(italic_title)
        r_title.font.name = "Times New Roman"
        r_title.font.size = Pt(12)
        r_title.font.italic = True
        r_title.font.color.rgb = RGBColor(0, 0, 0)

        r_suf = p.add_run(suffix)
        r_suf.font.name = "Times New Roman"
        r_suf.font.size = Pt(12)
        r_suf.font.color.rgb = RGBColor(0, 0, 0)
