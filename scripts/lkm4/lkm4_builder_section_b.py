# -*- coding: utf-8 -*-
"""
Section B Builder: Refleksi Mandiri
Pertemuan 4: Manajemen Ruang Lingkup Proyek
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_b(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "B. REFLEKSI MANDIRI")
    
    add_body_p(doc,
        "Bagian ini memuat refleksi kritis dan introspeksi personal saya terhadap penguasaan konsep Manajemen Ruang Lingkup Proyek "
        "yang dipelajari pada Pertemuan 4, serta penerapannya secara nyata dalam dinamika pengerjaan proyek kelompok kami, yaitu "
        "PRE-02 (Sistem Prediksi Probabilitas Keterlambatan Proyek Berbasis Machine Learning sebagai Early Warning System Manajemen Proyek).")

    # =========================================================================
    # REFLEKSI 1
    # =========================================================================
    add_heading_2(doc, "1. Menurut pemahaman Anda, mengapa Work Breakdown Structure (WBS) dianggap sebagai salah satu dokumen terpenting dalam manajemen proyek?")
    
    add_body_p(doc,
        "Berdasarkan pemahaman mendalam yang saya peroleh dari materi perkuliahan dan telaah literatur PMBOK Guide (PMI, 2018), "
        "saya memandang bahwa Work Breakdown Structure (WBS) adalah 'tulang punggung' (backbone) atau 'peta anatomi' dari keseluruhan "
        "sistem manajemen proyek. Proyek pada hakikatnya bermula dari ide abstrak, visi strategis bisnis, atau rumusan masalah penelitian "
        "yang kompleks. Tanpa adanya instrumen dekonstruksi yang sistematis, ide makro tersebut rentan menjadi target ilusi yang membingungkan. "
        "WBS adalah satu-satunya dokumen rekayasa yang mentransformasikan ketidakpastian makro menjadi unit-unit pekerjaan konkret, terukur, "
        "dan dapat dioperasionalkan.")

    add_body_p(doc,
        "Ada empat argumen kunci mengapa WBS menempati status dokumen paling fundamental:")

    add_num_p(doc, "1.", "Sebagai Poros Tunggal Pengintegrasi Area Pengetahuan:",
        "WBS bukan sekadar dokumen ruang lingkup, melainkan gerbang awal bagi seluruh perencanaan proyek lainnya. Kita tidak dapat menyusun "
        "jadwal (Schedule Management) tanpa mendekomposisi deliverable menjadi aktivitas dari WBS. Kita tidak dapat menghitung estimasi biaya "
        "(Cost Management) tanpa data biaya per work package. Kita tidak dapat menetapkan pembagian tugas tim (Resource Management) tanpa "
        "mengaitkan WBS ke Organization Breakdown Structure (OBS). Pendek kata, jika WBS runtuh atau cacat, maka seluruh dokumen turunan "
        "(jadwal, anggaran, matriks risiko) dipastikan akan salah total.")

    add_num_p(doc, "2.", "Penjelmaan Aturan 100% sebagai Pelindung Beban Kerja:",
        "Dalam realitas pelaksanaan proyek perangkat lunak, bahaya terbesar adalah timbulnya 'pekerjaan siluman' yang tidak pernah disadari "
        "hingga tenggat waktu mendekat. WBS memaksa tim untuk berpikir secara komprehensif dari awal guna memastikan 100% pekerjaan teridentifikasi. "
        "Sebaliknya, WBS juga bertindak sebagai benteng yang melindungi tim dari intervensi pekerjaan luar yang tidak resmi (scope creep).")

    add_num_p(doc, "3.", "Menciptakan Definisi Keberhasilan yang Tak Terbantahkan (Objective Completion):",
        "WBS membagi proyek besar menjadi potongan-potongan pencapaian kecil yang memiliki bukti deliverable fisik nyata. Hal ini menghilangkan "
        "kebiasaan buruk pengembang yang sering melaporkan kemajuan proyek secara subjektif (misal: 'pekerjaan koding sudah 90% selesai', namun "
        "kenyataannya sistem belum dapat dijalankan sama sekali). Dengan WBS, suatu tugas hanya diakui selesai jika deliverable pada paket kerja "
        "tersebut telah diserahterimakan dan lolos verifikasi kriteria penerimaan.")

    add_num_p(doc, "4.", "Alat Komunikasi Universal Lintas Disiplin:",
        "WBS menyajikan bahasa visual bersama (common language). Seorang manajer bisnis, sponsor, manajer proyek, ahli data science, "
        "hingga programmer junior dapat duduk di depan bagan WBS yang sama dan memahami secara tepat di mana posisi tugas mereka berada "
        "serta bagaimana kontribusi pekerjaan mereka menopang deliverable proyek secara keseluruhan.")

    # =========================================================================
    # REFLEKSI 2
    # =========================================================================
    add_heading_2(doc, "2. Apa tantangan terbesar yang Anda hadapi saat menyusun WBS untuk proyek kelompok, dan bagaimana Anda mengatasi tantangan tersebut?")
    
    add_body_p(doc,
        "Saat menyusun draf Work Breakdown Structure (WBS) untuk proyek kelompok kami, yaitu PRE-02 (Prediksi Probabilitas Keterlambatan Proyek "
        "Berbasis Machine Learning), tantangan terbesar yang kami hadapi berakar pada sifat inheren dari proyek data science dan riset AI "
        "yang sangat eksperimental, iteratif, dan sarat akan ketidakpastian (high exploratory uncertainty).")

    add_body_p(doc,
        "Secara spesifik, kami menghadapi dua hambatan metodologis utama:",
        bold_prefix="Deskripsi Tantangan Kritis:")

    add_bullet_p(doc, "Dilema Sudut Pandang Dekomposisi (Activity vs Deliverable Orientation):",
        "Pada tahap awal perancangan, anggota tim kami sempat terjebak menyusun WBS berbasis fungsi aktivitas (functional activity-oriented) "
        "seperti: 'Menganalisis', 'Mengoding', 'Mengetes', dan 'Mengevaluasi'. Struktur ini menyalahi kaidah PMBOK yang mewajibkan WBS berorientasi "
        "pada deliverables (noun/hasil kerja, bukan kata kerja proses). Selain itu, kami bingung bagaimana memetakan fase eksperimen machine learning: "
        "apakah harus didekomposisi berdasarkan jenis algoritma (Logistic Regression, Random Forest, dll.) ataukah berdasarkan tahapan CRISP-DM "
        "(Data Understanding, Preparation, Modeling, Evaluation)?")

    add_bullet_p(doc, "Ketidakpastian Durasi Eksperimen dan Risiko Kegagalan Model:",
        "Berbeda dari proyek rekayasa konvensional di mana alur pengerjaan bersifat deterministik linier, dalam machine learning ada kemungkinan "
        "besar model awal menghasilkan performa buruk (overfitting atau underfitting), sehingga tim harus berulang kali melakukan pembersihan fitur ulang "
        "dan tuning hiperparameter. Kami kesulitan menentukan tingkat kedalaman (granularity) paket kerja agar tidak terlalu dangkal namun juga tidak "
        "terlalu mikroskopis yang menyulitkan pengendalian.")

    add_body_p(doc,
        "Untuk mengatasi tantangan-tantangan metodologis tersebut, tim kelompok kami mengambil langkah-langkah solutif berikut:",
        bold_prefix="Strategi Mengatasi Tantangan:")

    add_num_p(doc, "1.", "Menerapkan Struktur Dekomposisi Berorientasi Deliverable Murni:",
        "Kami merombak total struktur draf WBS dengan menetapkan bahwa setiap elemen WBS Level 1 dan Level 2 wajib berupa kata benda atau "
        "entitas fisik/digital yang dihasilkan. Pada Level 1, kami menetapkan 5 Deliverables Utama: 1.0 Manajemen Proyek & Tata Kelola; "
        "2.0 Dataset & Pipeline Pra-pemrosesan; 3.0 Model Klasifikasi Probabilistik & Kalibrasi; 4.0 Antarmuka Dasbor & Layanan REST API; "
        "serta 5.0 Laporan Akademis & Diseminasi Riset.")

    add_num_p(doc, "2.", "Mengisolasi Komponen Eksperimen Berdasarkan Output Terukur:",
        "Pada cabang eksperimen model (Elemen 3.0), kami memecah dekomposisi berdasarkan paket deliverables eksperimen spesifik: Paket Kerja "
        "3.1 Model Baseline Logistic Regression, Paket Kerja 3.2 Model Ensemble Random Forest, Paket Kerja 3.3 Model Gradient Boosting, "
        "Paket Kerja 3.4 Model Neural Network Sigmoid, dan Paket Kerja 3.5 Modul Kalibrasi Probabilitas (Platt Scaling). Setiap paket kerja "
        "diikat dengan bukti output konkret, yaitu file bobot model (.pkl/.pt), skrip inferensi, dan log metrik evaluasi (Brier Score).")

    add_num_p(doc, "3.", "Menerapkan Prinsip Praktis '8/80 Hours Rule' untuk Ukuran Paket Kerja:",
        "Untuk menentukan batas terendah paket kerja, kami bersepakat mengadopsi kaidah emas industri 8/80 hours rule: tidak ada paket kerja "
        "yang durasi pengerjaannya kurang dari 8 jam (1 hari kerja mandiri), dan tidak ada paket kerja yang melebihi 80 jam (2 minggu kerja tim). "
        "Jika pekerjaan melebihi 80 jam, paket tersebut wajib didekomposisi lebih lanjut. Aturan ini sukses memberikan ukuran kerja yang ideal "
        "dan realistis untuk dipantau selama 8 pertemuan perkuliahan.")

    # =========================================================================
    # REFLEKSI 3
    # =========================================================================
    add_heading_2(doc, "3. Menurut Anda, apa yang terjadi jika sebuah proyek tidak memiliki WBS yang baik sejak awal? Berikan analisis dampaknya terhadap jadwal, biaya, dan kualitas proyek!")
    
    add_body_p(doc,
        "Jika sebuah proyek dieksekusi tanpa memiliki WBS yang terstruktur dengan baik sejak awal, maka proyek tersebut pada dasarnya dibangun "
        "di atas fondasi pasir yang rapuh. Ibarat membangun gedung bertingkat tanpa gambar cetak biru arsitektur (blueprint), tim akan bekerja "
        "berdasarkan asumsi liar masing-masing individu. Berdasarkan perspektif analisis Triple Constraints (Segitiga Besi Manajemen Proyek), "
        "ketiadaan WBS yang baik akan menimbulkan malapetaka sistemik yang saling berkaitan pada ketiga dimensi kinerja proyek:")

    add_num_p(doc, "1.", "Analisis Dampak terhadap Jadwal Proyek (Schedule Catastrophe):",
        "Ketiadaan WBS memicu kekacauan fatal pada linimasa proyek:")
    add_subbullet_p(doc, "Munculnya 'Sindrom 90% Selesai yang Semu' (The Illusion of Progress):",
        "Tanpa WBS yang memecah deliverables menjadi paket kerja nyata, manajer proyek tidak memiliki instrumen objektif untuk mengukur kemajuan harian tim. Tim akan mengklaim pekerjaan sudah hampir selesai, namun proyek macet berbulan-bulan di fase akhir karena ternyata banyak modul integrasi yang terlupakan.")
    add_subbullet_p(doc, "Kegagalan Jalur Kritis (Critical Path Failure) dan Efek Domino Keterlambatan:",
        "Ketiadaan paket kerja yang terdefinisi menyebabkan urutan ketergantungan tugas (precedence network) menjadi kacau. Keterlambatan pada satu tugas tersembunyi yang baru disadari belakangan akan seketika memicu penundaan berantai (cascading delays) yang meruntuhkan seluruh jadwal peluncuran proyek secara dramatis.")

    add_num_p(doc, "2.", "Analisis Dampak terhadap Biaya Proyek (Cost Explosion & Overrun):",
        "Ketiadaan WBS memicu pembengkakan anggaran yang tak terkendali:")
    add_subbullet_p(doc, "Estimasi Anggaran yang Tidak Berdasar (Pure Guestimation):",
        "Tanpa dekomposisi bottom-up berbasis WBS, estimasi anggaran proyek hanya berupa tebakan kasar di tingkat atas (top-down guess). Akibatnya, alokasi dana akan meleset jauh dari kebutuhan riil di lapangan.")
    add_subbullet_p(doc, "Pemborosan Finansial Akibat Pengerjaan Ulang (Massive Rework Costs):",
        "Karena batasan deliverable tidak jelas, programmer sering kali mengembangkan fitur yang keliru atau tidak sesuai kebutuhan arsitektur. Membongkar dan mengoding ulang modul yang salah menghabiskan ratusan jam kerja tenaga ahli, memicu biaya lembur tak terkendali, dan menghabiskan dana cadangan kontingensi (contingency reserve) hingga proyek terancam bangkrut sebelum selesai.")

    add_num_p(doc, "3.", "Analisis Dampak terhadap Kualitas Proyek (Quality Degradation & Defects):",
        "Ketiadaan WBS merusak standar keandalan produk akhir:")
    add_subbullet_p(doc, "Ketiadaan Kriteria Penerimaan Mutu yang Terikat (No Definition of Done):",
        "WBS yang buruk biasanya tidak disertai WBS Dictionary yang memadai. Tanpa adanya kriteria penerimaan spesifik per komponen, tim pengembang tidak mengetahui standar mutu apa yang harus dipenuhi. Kualitas produk akhir menjadi sangat heterogen, tidak konsisten, dan penuh cacat logika.")
    add_subbullet_p(doc, "Tahap Pengujian Mutu (Quality Testing) Menjadi Korban Pemotongan:",
        "Ketika proyek mengalami keterlambatan jadwal dan pembengkakan biaya akibat WBS yang buruk, manajemen biasanya panik dan mengambil jalan pintas yang fatal: memangkas waktu pengujian perangkat lunak, meniadakan kalibrasi model AI, dan melewati user acceptance testing. Hasilnya adalah perangkat lunak yang dirilis ke publik dalam kondisi cacat parah, tidak stabil, serta mencoreng reputasi tim pengembang.")

    add_body_p(doc,
        "Selain ketiga dampak pada Segitiga Besi di atas, dampak psikologis yang tak kalah merusak adalah hancurnya moril tim (team burnout), "
        "munculnya saling tuding antaranggota saat terjadi kesalahan, serta hilangnya kepercayaan (loss of trust) dari pihak sponsor proyek.")

    # =========================================================================
    # REFLEKSI 4
    # =========================================================================
    add_heading_2(doc, "4. Dari berbagai teknik pengumpulan kebutuhan yang dipelajari, teknik mana yang menurut Anda paling efektif untuk proyek kelompok Anda, dan mengapa?")
    
    add_body_p(doc,
        "Dalam konteks karakteristik unik proyek kelompok kami, yaitu PRE-02 (Sistem Prediksi Probabilitas Keterlambatan Proyek Berbasis Machine Learning), "
        "teknik pengumpulan kebutuhan yang menurut analisis kritis saya paling efektif adalah kombinasi hibrida antara "
        "Lokakarya Terfasilitasi (Facilitated Workshops / JAD) yang dipadukan secara langsung dengan Pembuatan Prototipe Awal (Prototyping / Storyboarding).")

    add_body_p(doc,
        "Pilihan ini didasarkan pada tiga alasan strategis dan kontekstual yang mendalam:",
        bold_prefix="Rasionalitas Pemilihan Teknik:")

    add_num_p(doc, "1.", "Menjembatani Kesenjangan Konseptual antara Ranah Bisnis Manajemen dan Teori Probabilitas AI:",
        "Proyek PRE-02 memiliki tantangan elisitasi yang unik: sistem ini bukan sekadar aplikasi CRUD (Create, Read, Update, Delete) biasa, "
        "melainkan sistem cerdas yang menghasilkan keluaran berupa probabilitas kontinu (angka antara 0.00 hingga 1.00) dan metrik kalibrasi statistik "
        "(seperti Brier Score). Pemangku kepentingan dari pihak manajerial proyek kerap memiliki persepsi keliru bahwa AI adalah 'bola kristal ajaib' "
        "yang dapat memprediksi masa depan secara 100% mutlak pasti (klasifikasi hitam-putih). "
        "Melalui Lokakarya Terfasilitasi (Facilitated Workshop), tim data science kami dapat duduk bersama pemangku kepentingan manajerial "
        "untuk menyamakan frekuensi pemahaman: mengedukasi konsep probabilitas terkalibrasi, mendiskusikan batasan model, dan menyepakati filosofi "
        "sistem sebagai Early Warning System (pendukung keputusan berbasis risiko, bukan pengambil keputusan mutlak).")

    add_num_p(doc, "2.", "Prototipe Visual Mereduksi Ambiguitas dan Memicu Kebutuhan Riil Pengguna:",
        "Jika kami hanya menggunakan teknik wawancara abstrak atau kuesioner teks, calon pengguna (seperti manajer proyek atau staf PMO) akan sangat "
        "kesulitan membayangkan bagaimana model machine learning ini akan mereka gunakan sehari-hari. Namun ketika kami menyajikan Prototipe Cepat "
        "(wireframe mockup interaktif antarmuka dasbor), pengguna dapat melihat secara visual: di mana kurva kalibrasi ditampilkan, bagaimana grafik "
        "radar 7 fitur kondisi kerja digambarkan, serta bagaimana lencana peringatan warna merah (alert badge) akan berkedip saat probabilitas melebihi 0.65. "
        "Melihat artefak konkret ini secara instan memicu respons kritis dan masukan kebutuhan yang sangat spesifik dan realistis dari pengguna, "
        "seperti: 'Kami butuh fitur simulasi what-if agar kami tahu apa yang terjadi jika backlog proyek dikurangi 10%'—sebuah kebutuhan emas yang "
        "mustahil muncul dari pertanyaan wawancara konvensional.")

    add_num_p(doc, "3.", "Membangun Konsensus Cepat dan Komitmen Bersama dalam Alokasi Waktu yang Terbatas:",
        "Dengan tenggat waktu perkuliahan yang padat (8 pertemuan), kami tidak memiliki kemewahan waktu untuk melakukan wawancara individual berulang kali "
        "yang rawan menghasilkan kebutuhan yang saling bertentangan. Melalui sesi workshop terfasilitasi selama 2 jam yang dihadiri bersama oleh dosen "
        "pembimbing, lead programmer, data scientist, dan representasi pengguna, seluruh perdebatan mengenai prioritas fitur (MoSCoW) dan ambang batas "
        "peringatan dini dapat diselesaikan dan disepakati di tempat secara konsensual.")

    # =========================================================================
    # REFLEKSI 5
    # =========================================================================
    add_heading_2(doc, "5. Bagaimana Anda membedakan antara pekerjaan yang termasuk dalam ruang lingkup proyek dan pekerjaan yang tidak termasuk? Berikan contoh dari proyek kelompok Anda!")
    
    add_body_p(doc,
        "Membedakan secara tegas antara pekerjaan yang termasuk dalam ruang lingkup proyek (in-scope) dan pekerjaan yang berada di luar ruang lingkup "
        "(out-of-scope) adalah esensi dari tata kelola proyek profesional. Untuk melakukan demarkasi batas kerja ini secara objektif dan ilmiah, "
        "saya dan tim menerapkan tiga kriteria pengujian sistemik (The Three-Gate Boundary Test):")

    add_num_p(doc, "1.", "Keterikatan Langsung terhadap Sasaran Strategis Project Charter (Goal Traceability Gate):",
        "Setiap usulan pekerjaan diuji apakah secara langsung mendukung tujuan utama proyek yang telah ditandatangani pada Project Charter. "
        "Jika suatu pekerjaan hanyalah fungsi sampingan yang tidak memiliki korelasi langsung dengan pemecahan masalah penelitian PRE-02 "
        "(yaitu memprediksi probabilitas keterlambatan proyek), maka pekerjaan tersebut secara otomatis diklasifikasikan sebagai out-of-scope.")

    add_num_p(doc, "2.", "Kaidah Pemetaan Aturan 100% pada Scope Baseline (WBS Mapping Gate):",
        "Pekerjaan yang sah (in-scope) wajib memiliki 'rumah resmi' di dalam hierarki WBS dan teruraikan dalam WBS Dictionary. "
        "Jika ada anggota tim yang mengajukan tugas baru yang tidak dapat dipetakan ke dalam salah satu work package sah yang telah disetujui, "
        "maka tugas tersebut dilarang untuk dikerjakan kecuali telah diajukan melalui prosedur perubahan formal (Change Request) dan disetujui.")

    add_num_p(doc, "3.", "Ketersediaan Alokasi Sumber Daya dan Realitas Batasan Akademis (Feasibility & Constraint Gate):",
        "Mengingat proyek ini dijalankan dalam kerangka mata kuliah dengan jangka waktu 8 pertemuan dan infrastruktur komputasi mandiri, "
        "pekerjaan-pekerjaan yang menuntut biaya lisensi komersial enterprise yang mahal, integrasi perangkat keras sensor IoT di lapangan, "
        "atau infrastruktur cloud multi-region yang kompleks harus dipagari secara tegas sebagai pekerjaan di luar ruang lingkup.")

    add_heading_3(doc, "Contoh Konkret Pembedaan Batas Ruang Lingkup pada Proyek PRE-02:")
    add_body_p(doc,
        "Tabel pembagian demarkasi batasan pekerjaan pada proyek kelompok PRE-02 diimplementasikan secara tegas sebagai berikut:")

    # Table 7: In-Scope vs Out-of-Scope PRE-02
    headers_t7 = ["Kategori Ruang Lingkup", "Deskripsi Batasan Pekerjaan", "Contoh Spesifik pada Proyek PRE-02", "Justifikasi / Alasan Rasional"]
    col_widths_t7 = [2.8, 4.5, 5.5, 4.5]
    data_t7 = [
        [
            "Pekerjaan Termasuk (IN-SCOPE)",
            "Seluruh aktivitas dan deliverables yang disepakati untuk dikerjakan dalam siklus proyek 8 pertemuan.",
            "1. Pengumpulan dan pra-pemrosesan data historis proyek (7 fitur).\n"
            "2. Pelatihan & komparasi 4 model ML (Logistic Regression, Random Forest, Gradient Boosting, Neural Network).\n"
            "3. Kalibrasi probabilitas dan evaluasi metrik Brier Score.\n"
            "4. Pembangunan REST API inferensi cepat berbasis FastAPI.\n"
            "5. Pembuatan dasbor web responsif untuk visualisasi early warning.\n"
            "6. Penyusunan paper ilmiah berformat IMRAD.",
            "Merupakan deliverable inti yang disyaratkan dalam Project Charter dan WBS untuk menjawab problem statement riset PRE-02 secara tuntas."
        ],
        [
            "Pekerjaan Tidak Termasuk (OUT-OF-SCOPE)",
            "Aktivitas atau fitur yang secara eksplisit dikecualikan dan tidak menjadi tanggung jawab tim.",
            "1. Pengembangan sistem manajemen tugas mandiri (seperti Jira/Trello).\n"
            "2. Penjadwalan ulang otomatis (autonomous rescheduling) terhadap aktivitas proyek di lapangan tanpa konfirmasi manusia.\n"
            "3. Integrasi real-time langsung ke database proprietary SAP/Oracle ERP enterprise yang berlisensi komersial tertutup.\n"
            "4. Pembuatan aplikasi mobile native (Android / iOS).\n"
            "5. Pengadaan dataset komersial berbayar eksternal.",
            "Mencegah timbulnya scope creep liar, menghindari pelanggaran lisensi perangkat lunak berbayar, dan menjaga fokus tim agar deliverables utama dapat diselesaikan tepat waktu sesuai kapasitas sumber daya yang tersedia."
        ]
    ]
    create_table(doc, col_widths_t7, headers_t7, data_t7, align_cols=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])

    doc.add_page_break()
