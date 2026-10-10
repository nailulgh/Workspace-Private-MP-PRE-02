# -*- coding: utf-8 -*-
"""
Section C Builder: Refleksi Mandiri
Pertemuan 13: Manajemen Pengadaan & Pemangku Kepentingan Proyek
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_c(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']

    add_heading_1(doc, "C. REFLEKSI MANDIRI")

    add_body_p(doc,
        "Bagian ini menyajikan perenungan kritis dan evaluasi personal mahasiswa terhadap konsep teoritis "
        "manajemen pengadaan dan manajemen pemangku kepentingan, serta refleksinya dalam mengawal keberhasilan "
        "proyek riset PRE-02 (Sistem Prediksi Probabilitas Keterlambatan Proyek Berbasis Machine Learning sebagai Early Warning System):")

    # =========================================================================
    # REFLEKSI 1
    # =========================================================================
    add_heading_2(doc, "1. Menurut Anda, mengapa pemangku kepentingan perlu diidentifikasi dan dikelola dengan baik sejak awal proyek, dan apa yang terjadi jika pemangku kepentingan diabaikan?")
    add_body_p(doc,
        "Identifikasi dan pelibatan pemangku kepentingan (stakeholder engagement) sejak fase inisiasi awal adalah penentu paling fundamental "
        "dari kesuksesan proyek. Keberhasilan proyek tidak hanya diukur dari terselesaikannya kode program atau pemenuhan jadwal semata, "
        "melainkan dari tercapainya kepuasan pemangku kepentingan (stakeholder satisfaction) dan penerimaan mereka terhadap hasil akhir sistem.")
    
    add_body_p(doc,
        "Jika pemangku kepentingan diabaikan atau terlambat diidentifikasi, proyek akan menghadapi konsekuensi destruktif berikut:")
    add_bullet_p(doc, "Penolakan Deliverables Akhir (Late Rejection): Ketika pemangku kepentingan kunci (seperti Dosen Pengampu atau calon pengguna) baru dimintai masukan di akhir, sering kali ditemukan bahwa asumsi yang dibangun tim tidak sesuai dengan ekspektasi mereka, memicu penolakan hasil karya dan keharusan mengulang pekerjaan (massive rework).")
    add_bullet_p(doc, "Munculnya Resistensi Aktif: Stakeholder yang merasa tidak dilibatkan atau aspirasinya diabaikan dapat berubah menjadi penentang aktif (resistant stakeholders) yang menghambat akses fasilitas, memotong dukungan anggaran, atau menciptakan opini negatif terhadap kelayakan proyek.")
    add_bullet_p(doc, "Perubahan Ruang Lingkup Liar (Scope Creep): Kebutuhan tersembunyi yang baru terungkap di tengah jalan akan memaksa perubahan desain mendadak yang merusak baseline jadwal dan biaya.")

    # =========================================================================
    # REFLEKSI 2
    # =========================================================================
    add_heading_2(doc, "2. Dari pengalaman Anda dalam proyek kelompok, kesalahan umum apa yang sering terjadi dalam mengelola pemangku kepentingan dan bagaimana cara menghindarinya?")
    add_body_p(doc,
        "Berdasarkan refleksi perjalanan pengerjaan proyek riset kelompok PRE-02, kesalahan umum yang paling sering terjadi mencakup:")
    
    add_num_p(doc, "Memperlakukan Seluruh Pemangku Kepentingan Secara Seragam (One-Size-Fits-All Approach): Tim sering kali mengirimkan pesan teknis mendalam yang sama kepada semua pihak. Misalnya mengirimkan log parameter kalibrasi model matematika kepada pihak manajemen yang sebenarnya hanya membutuhkan ringkasan status kelayakan proyek. Cara Menghindarinya: Mengklasifikasikan komunikasi berdasarkan kuadran Power/Interest Grid, di mana pihak eksekutif menerima ringkasan eksekutif dan tim teknis menerima detail implementasi.", 1)
    add_num_p(doc, "Komunikasi Satu Arah Tanpa Mendengarkan (The Illusion of Engagement): Manajer proyek hanya fokus menyampaikan apa yang telah dikerjakan tanpa secara aktif menggali umpan balik dan kekhawatiran yang dirasakan anggota tim lain. Cara Menghindarinya: Menerapkan active listening dan menyediakan sesi refleksi terbuka di akhir setiap sprint.", 2)
    add_num_p(doc, "Mengabaikan Pemangku Kepentingan Berdaya Rendah namun Berkepentingan Tinggi (Kuadran III): Menganggap enteng calon pengguna akhir karena mereka tidak memiliki kuasa formal atas nilai akademik. Cara Menghindarinya: Secara berkala melibatkan mereka dalam uji coba prototipe antarmuka early warning agar sistem benar-benar intuitif dan bermanfaat.", 3)

    # =========================================================================
    # REFLEKSI 3
    # =========================================================================
    add_heading_2(doc, "3. Menurut Anda, mengapa pengadaan proyek sering menjadi sumber risiko dan konflik dalam proyek, dan bagaimana manajer proyek dapat meminimalkan risiko tersebut?")
    add_body_p(doc,
        "Pengadaan proyek sering kali menjadi episentrum risiko dan konflik terbesar karena melibatkan dua entitas organisasi yang berbeda "
        "dengan motivasi dan tujuan bisnis yang saling bertentangan secara inheren: Pembeli (Buyer) menginginkan ruang lingkup sebesar mungkin "
        "dengan biaya serendah mungkin, sementara Penjual (Seller) berusaha meminimalkan biaya kerja mereka untuk memaksimalkan margin laba finansial.")

    add_body_p(doc,
        "Faktor-Faktor Pemicu Konflik Pengadaan:")
    add_bullet_p(doc, "Ambiguitas Pernyataan Kerja (Vague SOW): Spesifikasi kebutuhan yang tidak jelas atau multitafsir memicu sengketa apakah suatu fitur termasuk dalam kontrak atau harus dibayar ekstra sebagai change order.")
    add_bullet_p(doc, "Ketergantungan Eksternal (Loss of Direct Control): Manajer proyek kehilangan kendali langsung atas prioritas kerja tim vendor pihak ketiga, sehingga keterlambatan vendor langsung berdampak fatal pada jalur kritis proyek.")

    add_body_p(doc,
        "Strategi Manajer Proyek Meminimalkan Risiko Tersebut:\n"
        "1. Menyusun Procurement SOW yang Sangat Rinci dan Kriteria Penerimaan Terukur (Clear Acceptance Criteria).\n"
        "2. Memilih Jenis Kontrak yang Tepat dan Seimbang (misalnya FPIF dengan insentif ketepatan waktu untuk menyelaraskan kepentingan).\n"
        "3. Menetapkan Mekanisme Pengendalian Perubahan Kontrak yang Ketat (Formal Change Control System) dan Service Level Agreement (SLA) berpenalti.\n"
        "4. Menjalin Komunikasi Kolaboratif dan Pemantauan Kinerja Berkala (Vendor Performance Reviews) secara transparan.")

    # =========================================================================
    # REFLEKSI 4
    # =========================================================================
    add_heading_2(doc, "4. Bagaimana Anda akan menggunakan Power/Interest Grid untuk menentukan strategi komunikasi yang berbeda untuk berbagai pemangku kepentingan dalam proyek kelompok Anda?")
    add_body_p(doc,
        "Dalam proyek riset PRE-02, saya memanfaatkan Power/Interest Grid sebagai kompas navigasi distribusi informasi "
        "guna memastikan efisiensi bandwidth komunikasi dan tercapainya kepuasan seluruh pihak:")

    add_num_p(doc, "Untuk Kuadran I (Dosen Pengampu / Sponsor - Manage Closely): Saya menerapkan komunikasi interaktif dua arah dengan frekuensi tinggi. Memberikan laporan kemajuan pencapaian milestone mingguan, meminta masukan metodologi kalibrasi probabilitas, dan mendiskusikan hambatan kritis segera tanpa menunggu batas pengumpulan tugas akhir.", 1)
    add_num_p(doc, "Untuk Kuadran II (Ketua Jurusan / Kepala Lab - Keep Satisfied): Menerapkan komunikasi formal ringkas melalui email atau memo resmi bulanan. Memastikan kepatuhan terhadap aturan laboratorium komputasi kampus tanpa membebani mereka dengan detail perdebatan algoritma internal.", 2)
    add_num_p(doc, "Untuk Kuadran III (Calon Pengguna Manajer Proyek Eksternal - Keep Informed): Menerapkan komunikasi dorong (push) dan tarik (pull) melalui dokumentasi GitHub dan sesi demonstrasi prototipe berkala. Memberikan pembaruan visual tentang cara kerja sistem peringatan dini dan menampung masukan pengalaman pengguna (user feedback).", 3)
    add_num_p(doc, "Untuk Kuadran IV (Komunitas Pengembang Open-Source - Monitor): Menerapkan pemantauan pasif melalui penelusuran repositori pembaruan pustaka software tanpa memerlukan aktivitas komunikasi keluar aktif.", 4)

    # =========================================================================
    # REFLEKSI 5
    # =========================================================================
    add_heading_2(doc, "5. Menurut Anda, jenis kontrak apa yang paling sesuai jika proyek kelompok Anda memerlukan pengadaan barang atau jasa, dan mengapa?")
    add_body_p(doc,
        "Jika proyek kelompok PRE-02 memerlukan pengadaan barang atau jasa dari pihak ketiga, jenis kontrak yang paling tepat "
        "dibagi berdasarkan dua kategori kebutuhan spesifik:")

    add_heading_3(doc, "a. Pengadaan Layanan Infrastruktur Komputasi Cloud (Compute Cloud / GPU): Kontrak Fixed Price (Subscription FP)")
    add_body_p(doc,
        "Untuk pengadaan sewa server GPU pelatihan model Machine Learning, kontrak Fixed Price berbasis langganan bulanan "
        "(atau Fixed Price per Compute Hour dengan batas saldo tetap) adalah opsi terbaik. Karakteristik layanan cloud memiliki spesifikasi "
        "yang sangat terdefinisi (RAM, vCPU, GPU Memory) dan harga tetap per jam/bulan yang transparan. Kontrak ini memberikan kepastian "
        "biaya mutlak bagi anggaran mahasiswa dan memindahkan risiko pemeliharaan server fisik sepenuhnya kepada penyedia cloud.")

    add_heading_3(doc, "b. Pengadaan Jasa Bantuan Konsultasi Teknis Spesialis (Expert Consulting): Kontrak Time and Material (T&M) dengan Batas Not-To-Exceed")
    add_body_p(doc,
        "Jika tim membutuhkan bantuan konsultasi ahli luar untuk audit keamanan API atau konfigurasi deployment containerization, "
        "kontrak Time and Material (T&M) berdurasi pendek (misal: paket 15 jam konsultasi dengan klausul Not-To-Exceed) adalah yang paling fleksibel. "
        "Ruang lingkup konsultasi sering kali tidak dapat dipatok secara kaku di awal, dan klausul Not-To-Exceed melindungi tim dari risiko "
        "pembengkakan jam tagihan melebihi kemampuan anggaran proyek mahasiswa.")
