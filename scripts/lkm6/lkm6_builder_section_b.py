# -*- coding: utf-8 -*-
"""
Section B Builder: Tugas Individu (Perbandingan Manual vs Software)
Pertemuan 6: Manajemen Jadwal Proyek (Bagian II)
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_b(doc, helpers):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "B. TUGAS INDIVIDU: PERBANDINGAN PERHITUNGAN MANUAL VS PERANGKAT LUNAK")

    add_body_p(doc,
        "Bagian ini memuat evaluasi komparatif antara hasil perhitungan manual Critical Path Method (CPM) yang telah diselesaikan pada "
        "LKM Pertemuan 5 (untuk kasus 5 aktivitas: A, B, C, D, dan E) dengan hasil simulasi pemodelan menggunakan perangkat lunak manajemen "
        "proyek standar industri (Microsoft Project / ProjectLibre).")

    # =========================================================================
    # REKAP KASUS 5 AKTIVITAS
    # =========================================================================
    add_heading_2(doc, "1. Rekapitulasi Data Kasus 5 Aktivitas dan Hasil Komputasi")
    
    add_body_p(doc,
        "Data kasus dasar yang diujikan memuat 5 aktivitas dengan dependensi dan durasi kerja deterministik sebagai berikut: "
        "Aktivitas A (Durasi 3 hari, Predecessor: -), Aktivitas B (Durasi 4 hari, Predecessor: A), Aktivitas C (Durasi 5 hari, Predecessor: A), "
        "Aktivitas D (Durasi 2 hari, Predecessor: B, C), dan Aktivitas E (Durasi 3 hari, Predecessor: D).")

    table_data_rekap = [
        ["A", "-", "3 Hari", "0", "3", "0", "3", "0 Hari", "0 Hari", "KRITIS"],
        ["B", "A", "4 Hari", "3", "7", "4", "8", "1 Hari", "1 Hari", "Non-Kritis"],
        ["C", "A", "5 Hari", "3", "8", "3", "8", "0 Hari", "0 Hari", "KRITIS"],
        ["D", "B, C", "2 Hari", "8", "10", "8", "10", "0 Hari", "0 Hari", "KRITIS"],
        ["E", "D", "3 Hari", "10", "13", "10", "13", "0 Hari", "0 Hari", "KRITIS"]
    ]
    create_table(doc, [1.5, 1.8, 1.8, 1.3, 1.3, 1.3, 1.3, 1.8, 1.8, 2.1],
                 ["Aktivitas", "Predecessor", "Durasi", "ES", "EF", "LS", "LF", "Total Float", "Free Float", "Status"],
                 table_data_rekap,
                 align_cols=[WD_ALIGN_PARAGRAPH.CENTER] * 10)

    # =========================================================================
    # PERTANYAAN 1
    # =========================================================================
    add_heading_2(doc, "2. Analisis Perbedaan Hasil Perhitungan Manual dan Perhitungan Perangkat Lunak")
    
    add_body_p(doc,
        "Pertanyaan: Apakah terdapat perbedaan antara hasil perhitungan manual dan hasil perhitungan software untuk kasus 5 aktivitas? "
        "Jika ada, apa saja perbedaannya dan mengapa perbedaan itu terjadi?",
        bold_prefix="Rumusan Evaluasi: ")

    add_body_p(doc,
        "Berdasarkan pengujian komparatif, terdapat dua dimensi kesimpulan yang sangat menarik:")

    add_num_p(doc, "1.", "Kesamaan Mutlak pada Tingkat Logika Matematis (Pure Mathematical Equivalence):",
        "Pada tingkat abstraksi jaringan graf murni, **TIDAK ADA PERBEDAAN SAMA SEKALI** antara hasil perhitungan manual dan output perangkat "
        "lunak. Keduanya menghasilkan total durasi proyek yang persis sama, yaitu **13 hari kerja (workdays)**. Konfigurasi jalur kritis yang "
        "dihasilkan oleh software persis identik dengan lintasan manual, yaitu rangkaian **A → C → D → E** dengan nilai Total Float = 0. "
        "Aktivitas B secara konsisten diidentifikasi oleh software sebagai satu-satunya aktivitas non-kritis dengan nilai Total Float = 1 hari "
        "dan Free Float = 1 hari.")

    add_num_p(doc, "2.", "Perbedaan Nyata pada Tingkat Implementasi Penanggalan Kalender Riil:",
        "Meskipun durasi matematisnya identik, perbedaan signifikan muncul ketika angka abstrak diterjemahkan oleh perangkat lunak ke dalam "
        "garis waktu dunia nyata (real calendar timeline). Tiga penyebab utama perbedaan tersebut adalah:")
    add_subbullet_p(doc, "Penanganan Hari Libur dan Kalender Kerja (Working vs Non-Working Days):",
        "Perhitungan manual mengasumsikan waktu mengalir secara linier tanpa jeda (Day 0 hingga Day 13 secara berurutan). Sebaliknya, perangkat lunak "
        "menerapkan aturan Kalender Standar Organisasi (Standard Calendar 5 hari kerja per minggu, Senin–Jumat, 8 jam/hari). Jika proyek dimulai pada "
        "Senin, 5 Oktober 2026, maka durasi 13 hari kerja akan melompati dua akhir pekan (Sabtu-Minggu, 10–11 Oktober dan 17–18 Oktober). Akibatnya, "
        "tanggal selesai proyek yang dihitung software jatuh pada Rabu, 21 Oktober 2026—yang mencakup 17 hari kalender riil, bukan 13 hari kalender fisik!")
    add_subbullet_p(doc, "Konvensi Titik Awal Waktu (Day 0 vs Day 1 Convention):",
        "Perhitungan manual akademis menggunakan konvensi Day 0 (mulai pada titik 0 dan selesai di titik 3). Sementara itu, antarmuka perangkat lunak "
        "menampilkan tanggal kalender diskrit (mulai tanggal 5 Oktober pukul 08.00 dan selesai tanggal 7 Oktober pukul 17.00), yang sekilas menyerupai "
        "konvensi Day 1 sehingga sering membingungkan mahasiswa yang belum memahami representasi timestamp.")
    add_subbullet_p(doc, "Penetapan Default Constraints:",
        "Perhitungan manual mengasumsikan logika Finish-to-Start murni. Perangkat lunak secara default menerapkan kendala 'As Soon As Possible' (ASAP) "
        "dengan keterikatan jam kerja operasional.")

    # =========================================================================
    # PERTANYAAN 2
    # =========================================================================
    add_heading_2(doc, "3. Analisis Kelebihan dan Kekurangan Perhitungan Manual Dibandingkan Perangkat Lunak")
    
    add_body_p(doc,
        "Pertanyaan: Menurut Anda, apa kelebihan dan kekurangan perhitungan manual dibandingkan dengan penggunaan perangkat lunak?",
        bold_prefix="Rumusan Evaluasi: ")

    add_body_p(doc,
        "Evaluasi komparatif mendalam mengenai kelebihan dan kelemahan masing-masing metode disajikan dalam tabel komparatif berikut:")

    table_data_p2_comp = [
        ["Dimensi Evaluasi", "Perhitungan Manual (Forward/Backward Pass)", "Penggunaan Perangkat Lunak (MS Project / ProjectLibre)"],
        ["Nilai Pedagogis & Pemahaman Konseptual", "SANGAT TINGGI: Memaksa praktisi memahami esensi aliran logika dependensi graf, aturan max/min simpul majemuk, dan asal-usul rumus float.", "RENDAH - SEDANG: Pengguna rentan menjadi sekadar 'operator tombol' yang menerima hasil angka tanpa memahami logika matematis di baliknya."],
        ["Kecepatan Komputasi Awal", "Cepat untuk jaringan sederhana (di bawah 10 aktivitas).", "Memerlukan waktu penyiapan awal (input kalender, konfigurasi hari kerja, dan tabel aktivitas)."],
        ["Respon Pemutakhiran Jadwal (Agility)", "SANGAT LAMBAT: Jika terjadi perubahan durasi pada satu aktivitas awal, seluruh perhitungan maju dan mundur harus dihapus dan dihitung ulang secara manual.", "INSTAN & OTOMATIS: Mengubah satu angka durasi akan secara real-time memperbarui seluruh tanggal, float, dan jalur kritis di seluruh proyek."],
        ["Skalabilitas Kompleksitas Proyek", "SANGAT BURUK: Nyaris mustahil diterapkan secara andal pada proyek dengan lebih dari 20–30 aktivitas tanpa risiko kesalahan manusia (human error).", "SANGAT TINGGI: Mampu menangani ribuan aktivitas, dependensi majemuk bertingkat, dan multi-proyek enterprise secara presisi."],
        ["Integrasi Kalender & Sumber Daya", "SANGAT TERBATAS: Sangat rumit memperhitungkan libur nasional, shift kerja, cuti personel, dan alokasi beban kerja ganda secara manual.", "LENGKAP: Mendukung penyesuaian kalender spesifik per personel, visualisasi lembur, resource leveling otomatis, dan analisis Earned Value (EVM)."],
        ["Portabilitas & Ketergantungan", "MANDIRI: Hanya membutuhkan selembar kertas dan alat tulis, bebas dari ketergantungan perangkat keras, baterai, lisensi, atau crash software.", "TERGANTUNG: Memerlukan perangkat komputer, kompatibilitas sistem operasi, serta potensi biaya lisensi software komersial."]
    ]
    create_table(doc, [3.0, 7.0, 7.0],
                 ["Dimensi Evaluasi", "Perhitungan Manual (Forward/Backward Pass)", "Penggunaan Perangkat Lunak (MS Project / ProjectLibre)"],
                 table_data_p2_comp[1:])

    # =========================================================================
    # PERTANYAAN 3
    # =========================================================================
    add_heading_2(doc, "4. Situasi Di Mana Perhitungan Manual Masih Diperlukan Meskipun Perangkat Lunak Tersedia")
    
    add_body_p(doc,
        "Pertanyaan: Dalam situasi apa perhitungan manual masih diperlukan meskipun perangkat lunak tersedia?",
        bold_prefix="Rumusan Evaluasi: ")

    add_body_p(doc,
        "Meskipun perangkat lunak modern menawarkan otomatisasi yang sangat canggih, kemampuan perhitungan manual tetap menjadi kompetensi "
        "fundamental yang tak tergantikan dalam empat situasi profesional berikut:")

    add_num_p(doc, "1.", "Audit dan Validasi Integritas Jadwal (Sanity Check & Verification):",
        "Perangkat lunak adalah 'black box' bagi mereka yang tidak memahami logikanya. Sering kali terjadi kasus di mana seorang perencana "
        "jadwal salah menginput kendala (misal: secara tidak sengaja mengunci tugas dengan status 'Must Finish On' atau memasukkan lag negatif yang salah), "
        "sehingga software menghasilkan tanggal penyelesaian yang tidak masuk akal atau float bernilai negatif yang aneh. Dalam situasi ini, "
        "manajer proyek harus melakukan perhitungan manual cepat pada sub-jaringan terkait guna mengidentifikasi kesalahan konfigurasi internal software.")

    add_num_p(doc, "2.", "Pengambilan Keputusan Cepat di Lapangan dan Rapat Eksekutif (Field Decision Making):",
        "Ketika seorang manajer proyek sedang memimpin rapat negosiasi dengan klien atau meninjau progres langsung di lokasi pengembangan tanpa "
        "laptop aktif, kemampuan menghitung jalur kritis mental (mental math forward/backward pass) memungkinkan manajer untuk segera mengevaluasi "
        "dampak usulan perubahan: 'Jika fitur A kita tunda 2 hari, apakah proyek terlambat? Tidak, karena fitur A memiliki total float 3 hari.' "
        "Hal ini membangun wibawa profesional dan kepemimpinan yang tegas di hadapan pemangku kepentingan.")

    add_num_p(doc, "3.", "Pendidikan, Pelatihan, dan Standar Ujian Sertifikasi Profesional:",
        "Dalam ujian sertifikasi manajemen proyek internasional paling bergengsi seperti Project Management Professional (PMP®) atau Certified "
        "Associate in Project Management (CAPM®) dari PMI, seluruh soal analisis jaringan CPM disajikan dalam format perhitungan manual. "
        "Kandidat diwajibkan menghitung ES, EF, LS, LF, dan float dengan kertas cakar tanpa bantuan komputer.")

    add_num_p(doc, "4.", "Pengembangan Algoritma dan Perangkat Lunak Baru (Konteks Riset PRE-02):",
        "Dalam proyek riset kelompok PRE-02 yang mengembangkan pipeline eksperimen machine learning mandiri (experiment_pipeline.py), tim kami "
        "harus menuliskan rumus-rumus perhitungan varians jadwal dan dependency score secara langsung dalam kode Python murni. Kami tidak dapat "
        "mengintegrasikan kode komputasi jika kami tidak memahami rumus manual perhitungan jaringan jadwal.")

    # =========================================================================
    # PERTANYAAN 4
    # =========================================================================
    add_heading_2(doc, "5. Epistemologi Kepercayaan: Hasil Mana yang Lebih Dipercaya Jika Terjadi Perbedaan?")
    
    add_body_p(doc,
        "Pertanyaan: Jika hasil perhitungan manual dan software berbeda, mana yang lebih Anda percaya? Jelaskan alasan Anda!",
        bold_prefix="Rumusan Evaluasi: ")

    add_body_p(doc,
        "Sebagai seorang akademisi Teknik Informatika dan calon manajer proyek profesional, jawaban saya adalah: **Saya tidak akan mempercayai "
        "salah satu pihak secara buta, melainkan melakukan pelacakan akar penyebab (Root Cause Analysis) berdasarkan hierarki kepercayaan kontekstual:**")

    add_bullet_p(doc, "1. Mempercayai Logika Manual untuk Kebenaran Aritmatika Jaringan Murni:",
        "Jika perbedaan terjadi pada proyek berskala kecil atau pada durasi matematis murni (workdays), saya **lebih mempercayai perhitungan manual**. "
        "Alasannya adalah prinsip dasar ilmu komputer: *Garbage In, Garbage Out (GIGO)*. Komputer tidak pernah membuat kesalahan aritmatika penjumlahan, "
        "namun komputer sangat patuh pada pengaturan parameter yang dimasukkan manusia. Jika pengguna tanpa sadar mengaktifkan kalender 6 hari kerja, "
        "mengatur task constraint yang mengunci tanggal, atau memilih jenis task 'Fixed Duration' alih-alih 'Fixed Units', maka software akan "
        "menghasilkan angka yang bias. Perhitungan manual transparan 100% dan tidak memiliki asumsi tersembunyi.")

    add_bullet_p(doc, "2. Mempercayai Perangkat Lunak untuk Realitas Penanggalan Kalender Eksekusi:",
        "Sebaliknya, jika perbedaan tersebut menyangkut tanggal kalender penyelesaian riil (calendar finish dates), saya **jauh lebih mempercayai "
        "perangkat lunak**. Perhitungan manual manusia sangat rentan mengalami bias optimisme dan lupa memperhitungkan hari libur nasional, cuti bersama, "
        "serta akhir pekan. Manusia sering kali menganggap 13 hari kerja selesai dalam 13 hari kalender, padahal secara fisik dibutuhkan minimal 17 hari. "
        "Dalam hal akurasi kalender operasional, perangkat lunak memiliki presisi yang jauh lebih unggul.")

    add_body_p(doc,
        "Kesimpulan Sikap Profesional:",
        bold_prefix="Prinsip Verifikasi Rekayasa: ")
    add_body_p(doc,
        "Sikap ilmiah yang benar adalah **menggunakan perhitungan manual sebagai penguji logika (logic validator)** dan **menggunakan perangkat "
        "lunak sebagai mesin pelaksana kalender (execution scheduler)**. Jika muncul diskrepansi antara keduanya, hal tersebut merupakan sinyal "
        "bahwa ada parameter asumsi kalender atau relasi dependensi pada software yang belum terkonfigurasi secara tepat.")
