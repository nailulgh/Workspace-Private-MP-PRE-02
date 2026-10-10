# -*- coding: utf-8 -*-
"""
Section C Builder: Tugas Kelompok (CHECKPOINT #3: Jadwal Proyek PRE-02 dengan Software)
Pertemuan 6: Manajemen Jadwal Proyek (Bagian II)
Author: Muhammad Nailul Ghufron Majid (240605110160)
Course: Teori Manajemen Proyek - 2026
"""

import os
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def build_section_c(doc, helpers, workspace_root):
    add_heading_1 = helpers['add_heading_1']
    add_heading_2 = helpers['add_heading_2']
    add_heading_3 = helpers['add_heading_3']
    add_body_p = helpers['add_body_p']
    add_num_p = helpers['add_num_p']
    add_bullet_p = helpers['add_bullet_p']
    add_subbullet_p = helpers['add_subbullet_p']
    create_table = helpers['create_table']

    add_heading_1(doc, "C. TUGAS KELOMPOK (CHECKPOINT #3): PENYUSUNAN JADWAL PROYEK PRE-02 MENGGUNAKAN PERANGKAT LUNAK")

    add_body_p(doc,
        "Sebagai pemenuhan Checkpoint #3 Tugas Besar Manajemen Proyek, bagian ini menyajikan model jadwal komprehensif proyek kelompok berkode PRE-02 "
        "yang telah disusun dan disimulasikan menggunakan perangkat lunak manajemen proyek terstandar industri (Microsoft Project / ProjectLibre):")
    add_bullet_p(doc, "Nama Proyek Kelompok:", "Sistem Prediksi Probabilitas Keterlambatan Proyek Berbasis Machine Learning sebagai Early Warning System Manajemen Proyek (Kode: PRE-02).")
    add_bullet_p(doc, "Objek Studi Kasus:", "Portofolio Multi-Proyek Pengembangan 8 Modul Super ERP Farm Nation Enterprise (FNE) 2026.")
    add_bullet_p(doc, "Konfigurasi Kalender:", "Kalender Standar 5 Hari Kerja per Minggu (Senin–Jumat, 8 jam/hari kerja), mengabaikan hari libur akhir pekan.")
    add_bullet_p(doc, "Tanggal Mulai Proyek (Project Start Date):", "Senin, 5 Oktober 2026 pukul 08.00 WIB.")
    add_bullet_p(doc, "Cakupan Pekerjaan:", "21 Aktivitas Kerja Spesifik terdekomposisi dari WBS Level 3 & Level 4 yang mencakup seluruh siklus riset IMRAD.")

    # =========================================================================
    # TABEL JADWAL 21 AKTIVITAS
    # =========================================================================
    add_heading_2(doc, "1. Matriks Master Jadwal Proyek PRE-02 (21 Aktivitas Dekomposisi WBS)")
    
    add_body_p(doc,
        "Tabel 6.1 menyajikan data jadwal lengkap hasil kalkulasi perangkat lunak yang memuat kode WBS, kode aktivitas, nama pekerjaan, "
        "predecessor, estimasi durasi hari kerja (workdays), tanggal mulai, tanggal selesai, Early Start (ES), Early Finish (EF), "
        "Late Start (LS), Late Finish (LF), Total Float (TF), Free Float (FF), Person In Charge (PIC), serta status kekritisan lintasan:")

    table_data_21_act = [
        ["1.1.1", "ACT-01", "Formulasi Problem Statement & Spesifikasi 7 Fitur", "-", "3d", "05/10/26", "07/10/26", "0", "3", "0", "3", "0d", "0d", "Lead", "KRITIS"],
        ["1.1.2", "ACT-02", "Penyusunan Project Charter & Schedule Plan", "ACT-01", "2d", "08/10/26", "09/10/26", "3", "5", "3", "5", "0d", "0d", "Lead", "KRITIS"],
        ["2.1.1", "ACT-03", "Studi Literatur SOTA ML Delay Prediction", "ACT-01", "5d", "08/10/26", "14/10/26", "3", "8", "31", "36", "28d", "0d", "Lead", "Non-Kritis"],
        ["2.1.2", "ACT-04", "Penyusunan Matriks Literatur & Research Gap", "ACT-03", "3d", "15/10/26", "19/10/26", "8", "11", "36", "39", "28d", "28d", "Lead", "Non-Kritis"],
        ["3.1.1", "ACT-05", "Ekstraksi Data Metrik dari 8 Dokumen PMP ERP", "ACT-02", "4d", "12/10/26", "15/10/26", "5", "9", "5", "9", "0d", "0d", "Data Eng", "KRITIS"],
        ["3.1.2", "ACT-06", "Data Cleansing & Deteksi Anomali/Outlier", "ACT-05", "3d", "16/10/26", "20/10/26", "9", "12", "9", "12", "0d", "0d", "Data Eng", "KRITIS"],
        ["3.1.3", "ACT-07", "Imputasi Missing Values & Konsistensi Skema", "ACT-06", "2d", "21/10/26", "22/10/26", "12", "14", "12", "14", "0d", "0d", "Data Eng", "KRITIS"],
        ["3.1.4", "ACT-08", "Feature Engineering & Dependency Scoring", "ACT-07", "3d", "23/10/26", "27/10/26", "14", "17", "14", "17", "0d", "0d", "Data Eng", "KRITIS"],
        ["3.1.5", "ACT-09", "Normalisasi Fitur (Min-Max) & Encoding", "ACT-08", "2d", "28/10/26", "29/10/26", "17", "19", "17", "19", "0d", "0d", "Data Eng", "KRITIS"],
        ["3.1.6", "ACT-10", "Stratified Train-Test Dataset Splitting (80:20)", "ACT-09", "1d", "30/10/26", "30/10/26", "19", "20", "19", "20", "0d", "0d", "Data Eng", "KRITIS"],
        ["4.1.1", "ACT-11", "Training Baseline Logistic Regression", "ACT-10", "3d", "02/11/26", "04/11/26", "20", "23", "22", "25", "2d", "2d", "ML Eng", "Non-Kritis"],
        ["4.1.2", "ACT-12", "Training Ensemble Random Forest (predict_proba)", "ACT-10", "4d", "02/11/26", "05/11/26", "20", "24", "21", "25", "1d", "1d", "ML Eng", "Non-Kritis"],
        ["4.1.3", "ACT-13", "Training Gradient Boosting (XGBoost/LightGBM)", "ACT-10", "4d", "02/11/26", "05/11/26", "20", "24", "21", "25", "1d", "1d", "ML Eng", "Non-Kritis"],
        ["4.1.4", "ACT-14", "Arsitektur & Training Deep Learning MLP", "ACT-10", "5d", "02/11/26", "06/11/26", "20", "25", "20", "25", "0d", "0d", "ML Eng", "KRITIS"],
        ["4.2.1", "ACT-15", "Implementasi Kalibrasi Probabilitas (Platt)", "11,12,13,14", "3d", "09/11/26", "11/11/26", "25", "28", "25", "28", "0d", "0d", "ML Eng", "KRITIS"],
        ["4.2.2", "ACT-16", "Implementasi Kalibrasi Isotonic Regression", "ACT-15", "3d", "12/11/26", "16/11/26", "28", "31", "28", "31", "0d", "0d", "ML Eng", "KRITIS"],
        ["5.1.1", "ACT-17", "Komputasi Metrik Evaluasi (Brier, ROC, ECE)", "ACT-16", "3d", "17/11/26", "19/11/26", "31", "34", "31", "34", "0d", "0d", "ML Eng", "KRITIS"],
        ["5.1.2", "ACT-18", "Optimasi Ambang Batas Alert EWS (Youden J)", "ACT-17", "2d", "20/11/26", "23/11/26", "34", "36", "34", "36", "0d", "0d", "Lead", "KRITIS"],
        ["5.1.3", "ACT-19", "Desain & Pengujian Modul Interface Alert EWS", "ACT-18", "3d", "24/11/26", "26/11/26", "36", "39", "36", "39", "0d", "0d", "Dev", "KRITIS"],
        ["6.1.1", "ACT-20", "Penyusunan Draft Paper Lengkap Format IMRAD", "04, 19", "5d", "27/11/26", "03/12/26", "39", "44", "39", "44", "0d", "0d", "All", "KRITIS"],
        ["6.1.2", "ACT-21", "Finalisasi Slide Presentasi Eksekutif & Review", "ACT-20", "2d", "04/12/26", "07/12/26", "44", "46", "44", "46", "0d", "0d", "All", "KRITIS"]
    ]

    create_table(doc, [1.1, 1.2, 3.2, 1.4, 0.9, 1.3, 1.3, 0.7, 0.7, 0.7, 0.7, 0.9, 0.9, 1.1, 1.3],
                 ["WBS", "Kode", "Nama Aktivitas", "Pred", "Dur", "Start", "Finish", "ES", "EF", "LS", "LF", "TF", "FF", "PIC", "Status"],
                 table_data_21_act,
                 align_cols=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT,
                             WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
                             WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
                             WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
                             WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER])

    # =========================================================================
    # VISUALISASI GANTT CHART
    # =========================================================================
    add_heading_2(doc, "2. Visualisasi Gantt Chart Hasil Simulasi Perangkat Lunak")
    
    add_body_p(doc,
        "Gambar 6.1 menampilkan visualisasi Gantt Chart komprehensif dari 21 aktivitas proyek PRE-02 yang dihasilkan melalui perangkat lunak. "
        "Bagan ini memperlihatkan batang merah tebal sebagai aktivitas kritis yang berada pada jalur kritis, batang biru sebagai aktivitas "
        "non-kritis, serta garis putus-putus hijau yang mengindikasikan kelonggaran waktu (Total Float / Slack):")

    # Insert Image Gantt
    img_path = os.path.join(workspace_root, "scripts", "lkm6", "pre02_gantt_chart.png")
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Cm(15.5))

        p_caption = doc.add_paragraph()
        p_caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_caption.paragraph_format.space_before = Pt(2)
        p_caption.paragraph_format.space_after = Pt(10)
        r_cap = p_caption.add_run("Gambar 6.1 Gantt Chart Proyek PRE-02 (21 Aktivitas Dekomposisi WBS) dengan Sorotan Jalur Kritis")
        r_cap.font.name = "Times New Roman"
        r_cap.font.size = Pt(10)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(80, 80, 80)

    # =========================================================================
    # ANALISIS JADWAL PROYEK PRE-02
    # =========================================================================
    add_heading_2(doc, "3. Analisis Jadwal Proyek: Jalur Kritis, Total Float, dan Durasi Total")
    
    add_body_p(doc,
        "Berdasarkan hasil simulasi pemodelan jadwal pada Tabel 6.1 dan Gambar 6.1, tim manajemen proyek PRE-02 menyimpulkan parameter "
        "kunci operasional berikut:")

    add_num_p(doc, "1.", "Identifikasi Lintasan Jalur Kritis (Critical Path):",
        "Jalur kritis proyek PRE-02 adalah lintasan sekuensial yang terdiri atas 16 aktivitas berurutan:\n"
        "ACT-01 → ACT-02 → ACT-05 → ACT-06 → ACT-07 → ACT-08 → ACT-09 → ACT-10 → ACT-14 → ACT-15 → ACT-16 → ACT-17 → ACT-18 → ACT-19 → ACT-20 → ACT-21.\n"
        "Karakteristik seluruh aktivitas pada rantai ini adalah memiliki nilai Total Float = 0 hari dan Free Float = 0 hari. Keterlambatan satu "
        "hari saja pada salah satu dari 16 aktivitas ini dipastikan akan langsung menunda tanggal penyelesaian proyek.")

    add_num_p(doc, "2.", "Analisis Total Float untuk Aktivitas Non-Kritis:",
        "Terdapat 5 aktivitas yang berada di luar jalur kritis yang memiliki kelonggaran waktu (float positif) sehingga memberikan ruang manuver manajerial:\n"
        "• ACT-03 (Studi Literatur SOTA) & ACT-04 (Matriks Literatur & Research Gap): Memiliki Total Float sebesar 28 hari kerja! Hal ini dikarenakan "
        "tinjauan pustaka dapat dimulai sejak awal (hari ke-3) dan baru mutlak dibutuhkan sebagai input pada saat pengintegrasian paper (ACT-20 di hari ke-39).\n"
        "• ACT-11 (Baseline Logistic Regression): Memiliki Total Float = 2 hari dan Free Float = 2 hari (karena durasi 3 hari lebih singkat dibandingkan "
        "pelatihan model MLP ACT-14 yang memakan waktu 5 hari).\n"
        "• ACT-12 (Random Forest) & ACT-13 (Gradient Boosting): Masing-masing memiliki Total Float = 1 hari dan Free Float = 1 hari (durasi 4 hari vs 5 hari milik MLP).")

    add_num_p(doc, "3.", "Penetapan Total Durasi dan Tanggal Penyelesaian Proyek:",
        "– Total Durasi Proyek Berdasarkan Jalur Kritis: **46 Hari Kerja Efektif (Workdays)**.\n"
        "– Tanggal Mulai Proyek (Project Start): **Senin, 5 Oktober 2026 pukul 08.00 WIB**.\n"
        "– Tanggal Selesai Proyek (Project Finish): **Senin, 7 Desember 2026 pukul 17.00 WIB** (mencakup 9 minggu kalender penuh dan melompati "
        "18 hari libur akhir pekan Sabtu-Minggu).")

    add_num_p(doc, "4.", "Implikasi Strategis bagi Kepemimpinan Proyek PRE-02:",
        "Dalam dinamika tim riset kami, arsitektur deep learning MLP (ACT-14) terbukti menjadi penentu jalur kritis (bottleneck) di antara keempat "
        "model machine learning karena kompleksitas komputasinya yang paling tinggi (5 hari). Oleh karena itu, Machine Learning Engineer harus "
        "memprioritaskan pelatihan MLP terlebih dahulu sebelum melakukan tuning pada Random Forest atau Logistic Regression. Selain itu, float 28 hari "
        "pada ACT-03 dan ACT-04 memungkinkan Lead Author untuk mengalokasikan tenaganya membantu tim data pada fase ekstraksi dokumen PMP ERP (ACT-05) "
        "tanpa membahayakan jadwal penyusunan tinjauan pustaka.")
